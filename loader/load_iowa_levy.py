#!/usr/bin/env python3
"""Load Iowa school district property tax levy rates (Iowa Department of Management) into ia_levy_rate.

    DATABASE_URL=postgresql://...  python3 load_iowa_levy.py                 # find the files, load new or changed years
    python3 load_iowa_levy.py --file 2027=School_Tax_Rates_FY2027.xlsx --dry-run   # a file you downloaded by hand

The Department of Management lists its "School Tax Rates, FY ____.xlsx" files in a published Google Sheet
(dom.iowa.gov → Local Government → School Resources → School Property Tax Rate Files). Each file has one row per
district: its Department of Education number, name, and the rate per $1,000 of taxable valuation for each levy
(general, instructional support, management, voted PPEL, regular PPEL, debt service …) and the total levied.
FY2019 on (older years are .xls, which this doesn't read). Rows for parts of a district (a number the state's
annual-report data doesn't know, e.g. ALGONA (TITONKA)) are left out.

Runs in the "Iowa public data" GitHub workflow after the annual reports, so ia_district already exists.
Needs Python 3.9+, openpyxl and the psql client. Uses DATABASE_URL: never put the password in a file.
"""
import argparse, csv, hashlib, io, os, re, subprocess, sys, tempfile, urllib.request

FILE_LIST = ("https://docs.google.com/spreadsheets/d/e/2PACX-1vS78tJSKEh8IWFHwsLY3oqnyAwZt2XgJKiOftrrAyugpsg4Leh"
             "iPlEVkYnSMmdGnDNu36-QF1bNt8Lm/pub?output=csv")
FIRST_FY = 2019
UA = {"User-Agent": "HighGround public-data loader (Willow Holler; hello@willowholler.com)"}

# header text (lower case, spaces squeezed) -> column in ia_levy_rate
COLS = [
    (lambda h: "voted" in h and "ppel" in h, "voted_ppel"),
    (lambda h: "regular" in h and "ppel" in h, "regular_ppel"),
    (lambda h: h in ("total general", "general total"), "general_rate"),
    (lambda h: h.startswith("instructional support"), "instr_support"),
    (lambda h: h.startswith("management"), "management"),
    (lambda h: h.startswith("debt service"), "debt_service"),
    (lambda h: h.startswith("playground"), "playground"),
    (lambda h: h == "total rate", "total_rate"),
    (lambda h: h == "total levy", "total_levy"),
]
FIELDS = ["general_rate", "instr_support", "management", "voted_ppel", "regular_ppel", "debt_service", "playground",
          "total_rate", "total_levy"]


def fetch(url):
    try:
        with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=180) as r:
            return r.read()
    except Exception as e:
        raise SystemExit(f"Couldn't download {url}: {e!r}. Download the file by hand and run with --file FY=path.")


def find_files():
    """{fiscal_year: (file_name, url)} from the Department of Management's published file list."""
    text = fetch(FILE_LIST).decode("utf-8-sig", "replace")
    out = {}
    for row in csv.reader(io.StringIO(text)):
        if len(row) < 2:
            continue
        m = re.match(r"School Tax Rates, FY ?(\d{4})\.xlsx$", row[0].strip(), flags=re.I)
        if m and int(m.group(1)) >= FIRST_FY and row[1].startswith("http"):
            out[int(m.group(1))] = (row[0].strip(), row[1].strip())
    return out


def district_id(v):
    s = str(v or "").strip()
    if not re.fullmatch(r"\d{1,4}(\.0)?", s):
        return None
    n = int(float(s))
    return f"{n:04d}" if 0 < n < 8000 else None


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("$", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def parse(path):
    """[(de_district, name, {field: value})] from the first sheet with a 'Total Rate' header."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        for hi, r in enumerate(rows[:8]):
            heads = [re.sub(r"\s+", " ", str(x or "")).strip().lower() for x in r]
            if "total rate" not in heads:
                continue
            col = {}
            for i, h in enumerate(heads):
                for test, field in COLS:
                    if field not in col and h and test(h):
                        col[field] = i
                        break
            if "voted_ppel" not in col or "regular_ppel" not in col:
                raise SystemExit(f"{path}: found the header row but not the Voted PPEL / Regular PPEL columns: {heads}")
            out = []
            for r2 in rows[hi + 1:]:
                if not r2:
                    continue
                de = district_id(r2[0])
                if not de:
                    continue
                out.append((de, str(r2[1] or "").strip(), {f: num(r2[i]) if i < len(r2) else None for f, i in col.items()}))
            return out
    raise SystemExit(f"{path}: no sheet with a 'Total Rate' column. The file's layout may have changed.")


def psql(db, sql):
    r = subprocess.run(["psql", db, "-X", "-q", "-v", "ON_ERROR_STOP=1", "-tA"], input=sql, text=True, capture_output=True)
    if r.returncode != 0:
        raise SystemExit("psql failed:\n" + r.stderr)
    return r.stdout.strip()


def lit(s):
    return "null" if s is None else "'" + str(s).replace("'", "''") + "'"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--file", action="append", default=[], help="FY=path to a downloaded School Tax Rates workbook")
    ap.add_argument("--force", action="store_true", help="reload years even if the file hasn't changed")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    db = os.environ.get("DATABASE_URL")
    if not db and not a.dry_run:
        raise SystemExit("Set DATABASE_URL (or use --dry-run).")

    work = tempfile.mkdtemp(prefix="levy_")
    files = {}
    for x in a.file:
        fy, p = x.split("=", 1)
        files[int(fy)] = (os.path.basename(p), p)
    if not a.file:
        for fy, (name, url) in sorted(find_files().items()):
            data = fetch(url)
            if not data.startswith(b"PK"):
                raise SystemExit(f"{name}: the download wasn't a spreadsheet (Google may have asked for a confirmation). "
                                 "Download it by hand and run with --file.")
            p = os.path.join(work, f"levy_{fy}.xlsx")
            open(p, "wb").write(data)
            files[fy] = (name, p)
    if not files:
        raise SystemExit("No School Tax Rates files found in the Department of Management's list.")

    total = 0
    for fy, (name, p) in sorted(files.items()):
        digest = hashlib.sha256(open(p, "rb").read()).hexdigest()
        rows = parse(p)
        voted = sum(1 for _, _, v in rows if (v.get("voted_ppel") or 0) > 0)
        print(f"FY{fy}: {len(rows)} rows, {voted} with a voted PPEL ({name})", flush=True)
        if a.dry_run:
            for r in rows[:3]:
                print("   ", r)
            continue
        if not a.force and psql(db, f"select 1 from public.ia_load_run where kind = 'levy' and years = {lit(str(fy))} "
                                    f"and sha256 = {lit(digest)} and status = 'ok' limit 1;"):
            print(f"FY{fy}: unchanged, skipped", flush=True)
            continue
        csvp = os.path.join(work, f"levy_{fy}.csv")
        with open(csvp, "w", newline="") as o:
            w = csv.writer(o)
            for de, nm, v in rows:
                w.writerow([de, fy, nm] + ["" if v.get(f) is None else v.get(f) for f in FIELDS] + [name])
        n = psql(db, f"""create temp table _levy (de_district text, fiscal_year int, district_name text,
  general_rate numeric, instr_support numeric, management numeric, voted_ppel numeric, regular_ppel numeric,
  debt_service numeric, playground numeric, total_rate numeric, total_levy numeric, source_file text);
\\copy _levy from '{csvp}' with (format csv, null '')
begin;
delete from public.ia_levy_rate where fiscal_year = {fy};
insert into public.ia_levy_rate (de_district, fiscal_year, district_name, general_rate, instr_support, management,
  voted_ppel, regular_ppel, debt_service, playground, total_rate, total_levy, source_file)
select distinct on (l.de_district) l.de_district, l.fiscal_year, l.district_name, l.general_rate, l.instr_support,
  l.management, l.voted_ppel, l.regular_ppel, l.debt_service, l.playground, l.total_rate, l.total_levy, l.source_file
from _levy l join public.ia_district d on d.de_district = l.de_district
order by l.de_district, l.total_levy desc nulls last;
insert into public.ia_load_run (kind, source_url, sha256, rows_read, rows_loaded, years, status)
select 'levy', {lit(name)}, {lit(digest)}, {len(rows)}, count(*), {lit(str(fy))}, 'ok'
from public.ia_levy_rate where fiscal_year = {fy};
commit;
select count(*) from public.ia_levy_rate where fiscal_year = {fy};""").splitlines()[-1]
        print(f"FY{fy}: {n} districts loaded", flush=True)
        total += int(n)
    if not a.dry_run:
        print(f"Done: {total} district-years loaded.")


if __name__ == "__main__":
    main()
