#!/usr/bin/env python3
"""Load the Iowa Department of Education's Certified Annual Report (CAR) data files and certified
enrollment into Supabase. These replace the Data Hub for FY2019 onward (the Data Hub stops at FY2023).

    DATABASE_URL=postgresql://...  python3 load_iowa_car.py             # find files on the state's pages, load new ones
    python3 load_iowa_car.py --force                                     # reload every year
    python3 load_iowa_car.py --car 2025=car.xlsx --enroll 2025=ce.xlsx  # local files (no download)
    python3 load_iowa_car.py --dry-run --car 2025=car.xlsx              # parse only, print a summary

Needs Python 3.9+, openpyxl (pip install openpyxl) and the psql client.

Sources (VERIFIED 2026-10-07 against the FY2019, FY2023 and FY2025 files):
  CAR data, one workbook per year: educate.iowa.gov ... /accounting-reporting/certified-annual-report
    - Sheets ending "Data1": one row per district (column 'district' = DE number, 9 = AGWSR),
      header in row 3. Expenditure sheets are one per fund, 296 columns = ~40 functions x 7 objects;
      'RevData1' and 'BalSheetData1' have one column per fund x item (e.g. 'genproptx').
    - Rows 9000+ are AEAs, 8000s are charter organizations, blank-number rows are totals: skipped.
  Certified enrollment by district: educate.iowa.gov ... /certified-enrollment/public-schools
    - Header row 2: 'District #', 'District Name', ..., 'Certified Enrollment Row 7', ...,
      'Total Served Enrollment Row 11'. School year 2024-25 is used for FY2025.
"""
import argparse, collections, csv, hashlib, html, io, json, os, re, subprocess, sys, tempfile, urllib.parse, urllib.request

CAR_PAGE = "https://educate.iowa.gov/pk-12/operation-support/business-finance/accounting-reporting/certified-annual-report"
ENROLL_PAGE = "https://educate.iowa.gov/pk-12/data/data-collections/certified-enrollment/public-schools"
FIRST_FY = 2019          # first CAR year published as data; earlier years come from the Data Hub
UA = {"User-Agent": "HighGround public-data loader (Willow Holler)"}

STAGE_COLS = ["kind", "fiscal_year", "status", "aea", "dom_district", "de_district", "district_name",
              "column_name", "fund", "line", "amount", "per_pupil", "enrollment_category",
              "enrollment_category_number", "obj", "source"]

# ---------- funds ----------
# Expenditure sheet name (lower-case, starts with) -> fund. Order matters ('supptrust' before 'trust').
EXP_SHEETS = [("genexp", "General"), ("nonfidsch", "Non-Fiduciary Scholarship"), ("actexp", "Activity"),
              ("mgmntexp", "Management"), ("entreexp", "Entrepreneurial Education"), ("perlexp", "PERL"),
              ("supptrust", "Support Trust"), ("disastrecov", "Disaster Recovery"), ("libexp", "Library"),
              ("saveexp", "SAVE"), ("ppelexp", "PPEL"), ("capprojexp", "Other Capital Projects"),
              ("debtexp", "Debt Service"), ("permexp", "Permanent"), ("nutritionexp", "Nutrition"),
              ("lunchexp", "Nutrition"), ("othentexp", "Other Enterprise"), ("internalserv", "Internal Service"),
              ("trustexp", "Trust"), ("custodialexp", "Custodial")]
# Column prefix in RevData1 / BalSheetData1 -> fund. AEA funds and 'total' columns are skipped.
COL_FUNDS = {"gen": "General", "nonfid": "Non-Fiduciary Scholarship", "act": "Activity", "mgmnt": "Management",
             "entre": "Entrepreneurial Education", "perl": "PERL", "supptrst": "Support Trust",
             "disaster": "Disaster Recovery", "library": "Library", "save": "SAVE", "ppel": "PPEL",
             "capproj": "Other Capital Projects", "debt": "Debt Service", "perm": "Permanent",
             "lunch": "Nutrition", "othent": "Other Enterprise", "intserv": "Internal Service",
             "trst": "Trust", "agency": "Custodial", "cust": "Custodial",
             "aease": None, "aeajh": None, "total": None}
COL_PREFIXES = sorted(COL_FUNDS, key=len, reverse=True)

# ---------- expenditure functions -> the Data Hub's 14 groups (so callouts read the same) ----------
# ASSUMED for SpAd (special area administration -> General Administration) and OthSupp (other
# support services -> Business & Central Administration); the FY2023 comparison checks the rest.
FUNC = [("AttendSoc", "Student Support Services"), ("Guid", "Student Support Services"),
        ("Heal", "Student Support Services"), ("Psych", "Student Support Services"),
        ("Spch", "Student Support Services"), ("OthStud", "Student Support Services"),
        ("OT", "Student Support Services"), ("PT", "Student Support Services"), ("Vis", "Student Support Services"),
        ("ImpInst", "Instructional Staff Support Services"), ("Lib", "Instructional Staff Support Services"),
        ("Tech", "Instructional Staff Support Services"), ("Ass", "Instructional Staff Support Services"),
        ("OthInstSupp", "Instructional Staff Support Services"),
        ("Inst", "Instruction"),
        ("Bd", "General Administration"), ("Admin", "General Administration"), ("SpAd", "General Administration"),
        ("SchAd", "School/Building Administration"),
        ("BusAd", "Business & Central Administration"), ("Ware", "Business & Central Administration"),
        ("Print", "Business & Central Administration"), ("PR", "Business & Central Administration"),
        ("PInfo", "Business & Central Administration"), ("Pers", "Business & Central Administration"),
        ("ATech", "Business & Central Administration"), ("OthBA", "Business & Central Administration"),
        ("OthSupp", "Business & Central Administration"),
        ("OpMaint", "Plant Operation and Maintenance"), ("Trans", "Student Transportation"),
        ("FS", "Noninstructional Programs"), ("OthEnt", "Noninstructional Programs"),
        ("CommServ", "Noninstructional Programs"),
        ("Fac", "Facilities Acquisition and Construction"),
        ("Debt", "Debt Service (Principal, interest, fiscal charges)"),
        ("Flow", "AEA Support - Direct to AEA"),
        ("Out", "Transfers Out/Special Items/Down Adj"), ("Spec", "Transfers Out/Special Items/Down Adj"),
        ("Extra", "Transfers Out/Special Items/Down Adj"), ("LossD", "Transfers Out/Special Items/Down Adj"),
        ("Adj", "Transfers Out/Special Items/Down Adj")]
FUNC = sorted(FUNC, key=lambda x: len(x[0]), reverse=True)
OBJ = [("PurchServ", "Purchased services"), ("Supplies", "Supplies"), ("Equip", "Equipment"),
       ("Other", "Other"), ("Misc", "Other"), ("Serv", "Purchased services"), ("Sal", "Salaries"),
       ("Ben", "Benefits"), ("Oth", "Other")]

# ---------- revenue items -> the Data Hub's source names (ASSUMED groupings, checked for FY2023) ----------
REV = {"proptx": "Taxes Levied on Property", "incomesurtax": "Income Surtaxes",
       "excisetx": "Utility Replacement Excise Tax", "penalty": "Other Local Taxes", "othertx": "Other Local Taxes",
       "tuition": "Tuition\\Transportation Received", "transfees": "Tuition\\Transportation Received",
       "interest": "Earnings on Investments", "foodserv": "Nutrition Program Sales",
       "activities": "Student Activities and Sales",
       **{k: "Other Revenues from Local Sources" for k in ("commserv", "rentals", "donation", "capgains", "textsales",
                                                            "misclea", "miscsaleserv", "oprevprop", "refundpy", "othlocal")},
       "totintermed": "Revenue from Intermediary Sources",
       "staid": "State Foundation Aid", "instsupp": "Instructional Support State Aid",
       "save": "Statewide Sales, Services and Use Tax",
       **{k: "Other State Sources" for k in ("sped", "tlc", "fourpk", "lieutuit", "sbrc", "stcataid", "military",
                                              "othstlieuoftx")},
       **{k: "Federal Sources" for k in ("arra", "unresdirect", "unresindirect", "resdirect", "resindirect",
                                          "resindirectinter", "fedlieutx", "fedbehalflea")},
       **{k: "General Long-Term Debt Proceeds" for k in ("issuebonds", "proceedsloans", "proceedsleases",
                                                          "otherproceeds", "subscriptionproceeds")},
       "disposalproperty": "Proceeds of Fixed Asset Dispositions",
       **{k: "Transfers In/Special Items/Upward Adj" for k in ("interfundtransin", "upadjusts", "capcontrib",
                                                                "amortizepremium", "specitems", "extraitems",
                                                                "residequityin")}}
REV_TOTALS = {"totlocal", "totstate", "totfed", "totalotherfinansource", "totalotheritems", "totalrevandother"}
BAL = {"nonspendfundbal": "Nonspendable", "restfundbal": "Restricted", "commfundbal": "Committed",
       "assfundbal": "Assigned", "unassfundbal": "Unassigned"}
TOTAL_EQUITY = "totalfundequity"


def split_col(col):
    for p in COL_PREFIXES:
        if col.lower().startswith(p):
            return COL_FUNDS[p], col[len(p):].lower()
    return None, None


def exp_fund(sheet):
    t = sheet.lower().replace(" ", "")
    if t.startswith("aea"):
        return None
    for pre, fund in EXP_SHEETS:
        if t.startswith(pre):
            return fund
    return None


def exp_func_obj(col):
    for pre, line in FUNC:
        if col.startswith(pre):
            rest = col[len(pre):]
            for suf, obj in OBJ:
                if rest.endswith(suf):
                    return line, obj
            return line, "Other"
    return None, None


def district_id(v):
    try:
        n = int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None
    if n <= 0 or n >= 8000:          # 8000s charter organizations, 9000s AEAs
        return None
    return f"{n:04d}"


def short_name(s):
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    return re.sub(r"\s+(Comm(unity)?|Community)?\s*School District$", "", s, flags=re.I).strip() or s


def header_row(ws, max_rows=8):
    for i, r in enumerate(ws.iter_rows(min_row=1, max_row=max_rows, values_only=True), start=1):
        if r and any(isinstance(x, str) and x.strip().lower() in ("district", "district #") for x in r):
            return i, [x.strip() if isinstance(x, str) else x for x in r]
    return None, None


def parse_car(fy, path):
    """-> dict (kind, de, column_name) -> [fund, line, obj, amount, name]; plus stats."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    stats = collections.Counter()
    unmapped = collections.Counter()
    names = {}

    def add(kind, de, fund, line, obj, amt):
        key = (kind, de, f"car:{fund}:{line}" + (f":{obj}" if obj else ""))
        cur = out.setdefault(key, [fund, line, obj, 0.0])
        cur[3] += amt

    for ws in wb.worksheets:
        title = ws.title.strip()
        low = title.lower().replace(" ", "")
        if "data" not in low:
            continue
        hi, hdr = header_row(ws)
        if not hdr:
            stats["sheets_without_header"] += 1
            continue
        di = next(i for i, x in enumerate(hdr) if isinstance(x, str) and x.lower() == "district")
        ni = next((i for i, x in enumerate(hdr) if isinstance(x, str) and x.lower().startswith("agencyname")), None)
        if low.startswith("revdata"):
            mode, sheet_fund = "rev", None
        elif low.startswith("balsheetdata"):
            mode, sheet_fund = "bal", None
        else:
            sheet_fund = exp_fund(title)
            mode = "exp" if sheet_fund else None
            if not mode:
                if not low.startswith("aea"):
                    unmapped["sheet:" + title] += 1
                continue
        stats["sheets_" + mode] += 1
        for r in ws.iter_rows(min_row=hi + 1, values_only=True):
            if not r or di >= len(r):
                continue
            de = district_id(r[di])
            if not de:
                continue
            if ni is not None and r[ni]:
                names[de] = short_name(r[ni])
            for ci, col in enumerate(hdr):
                if ci in (di, ni) or not isinstance(col, str) or ci >= len(r):
                    continue
                v = r[ci]
                if not isinstance(v, (int, float)) or v == 0:
                    continue
                if mode == "exp":
                    if col.startswith("Tot"):
                        continue
                    line, obj = exp_func_obj(col)
                    if not line:
                        unmapped[f"exp:{col}"] += 1
                        continue
                    add("exp", de, sheet_fund, line, obj, float(v))
                else:
                    fund, item = split_col(col)
                    if fund is None:
                        continue                     # AEA funds, statewide totals
                    if mode == "rev":
                        if item in REV_TOTALS:
                            continue
                        line = REV.get(item)
                        if not line:
                            unmapped[f"rev:{item}"] += 1
                            continue
                        add("rev", de, fund, line, None, float(v))
                    else:
                        if item in BAL:
                            add("bal", de, fund, BAL[item], None, float(v))
                        elif item == TOTAL_EQUITY:
                            # ending fund balance as an expenditure-side line, like the Data Hub
                            add("exp", de, fund, "Ending Fund Balance", None, float(v))
    for (kind, de, cn), (fund, line, obj, amt) in out.items():
        stats[f"rows_{kind}"] += 1
    stats["districts"] = len({de for (_, de, _) in out})
    return out, names, stats, unmapped


def parse_enrollment(fy, path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    hi, hdr = header_row(ws)
    if not hdr:
        raise SystemExit(f"enrollment {fy}: no 'District #' header found")
    norm = [re.sub(r"\s+", " ", str(x or "")).lower() for x in hdr]
    di = next(i for i, x in enumerate(norm) if x.startswith("district #") or x == "district")
    ni = next(i for i, x in enumerate(norm) if x.startswith("district name"))
    ci = next(i for i, x in enumerate(norm) if x.startswith("certified enrollment"))
    si = next((i for i, x in enumerate(norm) if x.startswith("total served")), None)
    rows = []
    for r in ws.iter_rows(min_row=hi + 1, values_only=True):
        de = district_id(r[di]) if r and di < len(r) else None
        if not de or not isinstance(r[ci], (int, float)):
            continue
        rows.append([de, fy, short_name(r[ni]), r[ci], r[si] if si is not None and isinstance(r[si], (int, float)) else ""])
    return rows


# ---------- finding the files on the state's pages ----------
def fetch(url):
    req = urllib.request.Request(url, headers=UA)
    try:
        with urllib.request.urlopen(req, timeout=300) as r:
            return r.read()
    except Exception as e:
        raise SystemExit(f"Couldn't download {url}: {e!r}. If the state's site is blocking automated downloads, "
                         "download the files by hand and run with --car FY=path / --enroll FY=path.")


def find_links(page_url, pattern):
    """{fiscal_year: absolute_url} for anchors whose text matches pattern (groups: yyyy, yy)."""
    text = fetch(page_url).decode("utf-8", "replace")
    found = {}
    for href, label in re.findall(r'<a[^>]+href="([^"]+)"[^>]*>(.*?)</a>', text, flags=re.S | re.I):
        label = re.sub(r"<[^>]+>", " ", html.unescape(label))
        label = re.sub(r"\s+", " ", label).strip()
        m = re.search(pattern, label, flags=re.I)
        if m:
            fy = 2000 + int(m.group(2))
            found.setdefault(fy, urllib.parse.urljoin(page_url, html.unescape(href)))
    return found


def psql(db, sql, capture=True):
    r = subprocess.run(["psql", db, "-X", "-q", "-v", "ON_ERROR_STOP=1", "-tA", "-F", "\t"], input=sql,
                       text=True, capture_output=True)
    if r.returncode != 0:
        raise SystemExit("psql failed:\n" + r.stderr)
    return r.stdout.strip() if capture else None


def lit(s):
    return "null" if s is None else "'" + str(s).replace("'", "''") + "'"


def sha(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--car", action="append", default=[], help="FY=path to a local CAR workbook")
    ap.add_argument("--enroll", action="append", default=[], help="FY=path to a local enrollment workbook")
    ap.add_argument("--no-download", action="store_true", help="use only the local files given")
    ap.add_argument("--force", action="store_true", help="reload files even if unchanged")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    db = os.environ.get("DATABASE_URL")
    if not db and not a.dry_run:
        raise SystemExit("Set DATABASE_URL (or use --dry-run).")
    work = tempfile.mkdtemp(prefix="car_")

    car = {int(k): v for k, v in (x.split("=", 1) for x in a.car)}
    enr = {int(k): v for k, v in (x.split("=", 1) for x in a.enroll)}
    if not a.no_download and not car and not enr:
        car_links = find_links(CAR_PAGE, r"\b(\d{4})-(\d{2})\b.*\bCAR\b")
        enr_links = find_links(ENROLL_PAGE, r"\b(\d{4})-(\d{2})\b.*Certified Enrollment by District\b(?!.*AEA)")
        car_links = {fy: u for fy, u in car_links.items() if fy >= FIRST_FY}
        enr_links = {fy: u for fy, u in enr_links.items() if fy >= FIRST_FY}
        print(f"CAR files found: {sorted(car_links)}", flush=True)
        print(f"Enrollment files found: {sorted(enr_links)}", flush=True)
        if not car_links:
            raise SystemExit("No CAR data links found on the state's page; its layout may have changed.")
        for fy, u in car_links.items():
            p = os.path.join(work, f"car_{fy}.xlsx"); open(p, "wb").write(fetch(u)); car[fy] = p
        for fy, u in enr_links.items():
            p = os.path.join(work, f"enroll_{fy}.xlsx"); open(p, "wb").write(fetch(u)); enr[fy] = p

    def last_sha(kind):
        if not db or a.force:
            return None
        return psql(db, f"select sha256 from public.ia_load_run where kind = {lit(kind)} and status = 'ok' order by id desc limit 1;") or None

    failures = []
    # ---- enrollment first (per-pupil figures depend on it) ----
    enroll_rows, enroll_runs = [], []
    for fy in sorted(enr):
        h = sha(enr[fy])
        if last_sha(f"enroll:{fy}") == h:
            print(f"[enroll {fy}] unchanged; skipping"); continue
        try:
            rows = parse_enrollment(fy, enr[fy])
        except (StopIteration, SystemExit, Exception) as e:
            print(f"[enroll {fy}] FAILED to read ({e!r}); per-pupil figures for FY{fy} will be missing", flush=True)
            failures.append(f"enroll {fy}"); continue
        print(f"[enroll {fy}] {len(rows)} districts, total certified {sum(r[3] for r in rows):,.1f}", flush=True)
        enroll_rows += rows; enroll_runs.append((fy, h, len(rows)))
    if enroll_rows and not a.dry_run:
        p = os.path.join(work, "enroll_stage.csv")
        with open(p, "w", newline="") as o:
            w = csv.writer(o); w.writerow(["de_district", "fiscal_year", "district_name", "certified_enrollment", "served_enrollment"])
            w.writerows(enroll_rows)
        check = psql(db, f"""
set statement_timeout = 0;
delete from public.ia_enroll_stage;
\\copy public.ia_enroll_stage (de_district, fiscal_year, district_name, certified_enrollment, served_enrollment) from '{p}' with (format csv, header true, null '')
-- before replacing them: how do the Data Hub's own enrollment figures compare with certified?
select coalesce(s.fiscal_year::text, '-'), count(*),
       round((percentile_cont(0.5) within group (order by y.enrollment / nullif(s.certified_enrollment, 0)))::numeric, 3),
       count(*) filter (where ia_band(s.certified_enrollment) = y.enrollment_category_number)
from public.ia_enroll_stage s join public.ia_district_year y
  on y.de_district = s.de_district and y.fiscal_year = s.fiscal_year and y.status = 'Actual'
where coalesce(y.enrollment_source, 'datahub') = 'datahub' and y.enrollment > 0
group by s.fiscal_year order by 1;
""")
        for line in [l for l in check.splitlines() if l.strip()]:
            fy_, n, ratio, same = line.split("\t")
            print(f"[enroll check FY{fy_}] {n} districts: Data Hub enrollment / certified = {ratio} (median); "
                  f"size band agrees for {same} of {n}", flush=True)
        out = psql(db, "select public.ia_publish_enrollment()::text;")
        print(f"[enroll] published: {out.splitlines()[-1]}", flush=True)
        for fy, h, n in enroll_runs:
            psql(db, f"insert into public.ia_load_run (kind, source_url, sha256, rows_read, rows_loaded, years, status) "
                     f"values ({lit('enroll:'+str(fy))}, {lit(enr[fy] if a.no_download or a.enroll else ENROLL_PAGE)}, {lit(h)}, {n}, {n}, {lit(str(fy))}, 'ok');")

    # ---- CAR years ----
    for fy in sorted(car):
        h = sha(car[fy])
        if last_sha(f"car:{fy}") == h:
            print(f"[car {fy}] unchanged; skipping"); continue
        try:
            out, names, stats, unmapped = parse_car(fy, car[fy])
        except Exception as e:
            print(f"[car {fy}] FAILED to read ({e!r}); skipping this year", flush=True)
            failures.append(f"car {fy}"); continue
        print(f"[car {fy}] districts {stats['districts']}  rows exp {stats['rows_exp']}  rev {stats['rows_rev']}  "
              f"bal {stats['rows_bal']}  sheets exp {stats['sheets_exp']} rev {stats['sheets_rev']} bal {stats['sheets_bal']}", flush=True)
        if unmapped:
            print(f"[car {fy}] NOT MAPPED (skipped): {dict(unmapped.most_common(25))}", flush=True)
        if stats["districts"] < 250 or not stats["rows_exp"]:
            print(f"[car {fy}] too few districts/rows; not loading this year", flush=True)
            failures.append(f"car {fy}"); continue
        if a.dry_run:
            continue
        p = os.path.join(work, f"car_stage_{fy}.csv")
        with open(p, "w", newline="") as o:
            w = csv.writer(o); w.writerow(STAGE_COLS)
            for (kind, de, cn), (fund, line, obj, amt) in out.items():
                w.writerow([kind, fy, "Actual", "", "", de, names.get(de, de), cn, fund, line, round(amt, 2),
                            "", "", "", obj or "", "car"])
        cols = ", ".join(STAGE_COLS)
        compare = ""
        has_dh = psql(db, f"select count(*) from public.ia_fin where fiscal_year = {fy} and coalesce(source, 'datahub') = 'datahub';")
        res = psql(db, f"""
set statement_timeout = 0;
begin;
delete from public.ia_stage where source = 'car' or kind in ('exp','rev','bal');
\\copy public.ia_stage ({cols}) from '{p}' with (format csv, header true, null '')
{"select 'CMP', 'exp', * from public.ia_compare_stage('exp', " + str(fy) + ");" if has_dh not in ('', '0') else ''}
{"select 'CMP', 'rev', * from public.ia_compare_stage('rev', " + str(fy) + ");" if has_dh not in ('', '0') else ''}
select 'PUB', public.ia_publish('exp', false)::text;
select 'PUB', public.ia_publish('rev', false)::text;
select 'PUB', public.ia_publish('bal')::text;
insert into public.ia_load_run (kind, source_url, sha256, rows_read, rows_loaded, years, status, message)
values ({lit('car:'+str(fy))}, {lit(os.path.basename(car[fy]))}, {lit(h)}, {sum(stats[k] for k in ('rows_exp','rows_rev','rows_bal'))},
        {sum(stats[k] for k in ('rows_exp','rows_rev','rows_bal'))}, {lit(str(fy) + ' Actual')}, 'ok',
        {lit(json.dumps({'unmapped': dict(unmapped)}))});
commit;
""")
        for line in res.splitlines():
            parts = line.split("\t")
            if parts[0] == "CMP":
                _, kind, ln, staged, stored, diff, off = parts
                print(f"[check FY{fy} {kind}] {ln[:48]:<48} state report {float(staged or 0):>15,.0f}  Data Hub {float(stored or 0):>15,.0f}  "
                      f"diff {diff or '-':>6}%  districts >5% off: {off}", flush=True)
            elif parts[0] == "PUB":
                print(f"[car {fy}] published {parts[1][:160]}", flush=True)

    if failures:
        raise SystemExit(f"Finished with problems: {failures}. Everything else was loaded.")
    print("Done.")


if __name__ == "__main__":
    main()
