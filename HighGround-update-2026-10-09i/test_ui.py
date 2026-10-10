# HighGround shell: browser test against a simulated Supabase. Run: pip install playwright; playwright install chromium; python3 test_ui.py
import asyncio, json, re, urllib.parse, subprocess, time, os
from playwright.async_api import async_playwright
ROOT=os.path.dirname(os.path.abspath(__file__)); SB="https://test.supabase.co"; SHOTS=os.path.join(os.path.dirname(os.path.abspath(__file__)),"shots"); os.makedirs(SHOTS,exist_ok=True)
D1={"id":"d1","slug":"ironwood-valley","name":"Ironwood Valley Community School District","short_name":"Ironwood Valley","state":"IA","county":"Fictional","brand_color":"#1F4E8C","is_demo":True,"public_link_enabled":True,"state_district_id":"7777"}
D2={"id":"d2","slug":"cottonwood-ridge","name":"Cottonwood Ridge Community School District","short_name":"Cottonwood Ridge","state":"IA","county":None,"brand_color":"#7A4E2D","is_demo":True,"public_link_enabled":True}
USERS={"admin@example.test":("u-admin","Pat Admin",[("d1","admin"),("d2","viewer")],False),
       "viewer@example.test":("u-viewer","Val Viewer",[("d1","viewer")],False),
       "new@example.test":("u-new","Nia New",[],False),
       "staff@example.test":("u-staff","Sam Staff",[],True),
       "empty.staff@example.test":("u-estaff","Eve Staff",[],True),
       "mfa@example.test":("u-mfa","Mo Factor",[],False),
       "board@example.test":("u-board","Bo Board",[("d1","board")],False),
       "bm@example.test":("u-bm","Bea Manager",[("d1","business_manager")],False),
       "sup@example.test":("u-sup","Sue Super",[("d1","superintendent")],False)}
PW="correct horse battery"
TABLES={"initiative":[{"district_id":"d1","name":"Middle school HVAC","type":"capital","status":"approved","focus_area":"Facilities","tier":"must","cost_confidence":"firm","condition":"poor","id":"i1"}],
 "scenario":[{"district_id":"d1","id":"s1","name":"District baseline","is_board_version":True,"is_locked":True,"updated_at":"2026-09-23T12:00:00Z"}],
 "import_batch":[{"district_id":"d1","id":"b1","kind":"gl_monthly","period_end":"2026-08-31","file_name":"august.csv","status":"applied","uploaded_at":"2026-09-10T12:00:00Z"}],
 "priority":[],"measure":[],
 "fund_balance":[{"district_id":"d1","fund":"save","as_of":"2026-07-01","amount":2150000,"source":"manual"},{"district_id":"d1","fund":"save","as_of":"2026-06-01","amount":1,"source":"manual"}],
 "debt_obligation":[{"district_id":"d1","name":"SAVE revenue bonds, Series 2021","fund":"save","annual_payment":610000,"final_fy":2033}],
 "district_settings":[],"publication":[{"district_id":"d1","kind":"board_plan","title":"Board version","published_at":"2026-09-20T12:00:00Z","is_current":True}],
 "invitation":[{"district_id":"d1","id":"inv1","email":"bm@example.test","role":"business_manager","created_at":"2026-09-28T12:00:00Z","expires_at":"2026-10-28T12:00:00Z"}]}
IRON=json.loads(subprocess.check_output(["node","-e",
  "const C=require('./capital.js'),D=require('./demo_data.js');let i=0;"
  "process.stdout.write(JSON.stringify(C.demoRows(D['ironwood-valley'],'d1',()=>'00000000-0000-4000-8000-'+String(++i).padStart(12,'0'))))"],cwd=os.path.dirname(os.path.abspath(__file__))))
for t in ["district_settings","debt_obligation","initiative","scenario","scenario_initiative","phase","phase_funding","financing"]:
  TABLES[t]=IRON[t]
TABLES["fund_balance"]=IRON["fund_balance"]+[{"district_id":"d1","fund":"save","as_of":"2026-06-01","amount":1,"source":"manual"}]
for i,dbt in enumerate(TABLES["debt_obligation"]): dbt.setdefault("id","debt-%d"%i)
for i,fn in enumerate(TABLES["financing"]): fn.setdefault("id","fin-%d"%i)
_ph_sid=[sc["id"] for sc in IRON["scenario"] if not sc["is_board_version"]][0]
_init0=IRON["initiative"][0]["id"]
for sc in TABLES["scenario"]:
  sc["updated_at"]="2026-09-23T12:00:00Z"
  if sc["id"] in IRON["lock"]: sc["is_locked"]=True
GOLD=json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)),"golden_engine.json")))["districts"]["demo-medium"]
async def signout(p):
  await p.click("button[data-action=userMenu]"); await p.click(".umenu .menu-signout")
def fmtK(v):
  s="−" if v<0 else ""; v=abs(v)
  if v>=1e6: return s+"$%.2fM"%(v/1e6)
  if v>=1000: return s+"$%dk"%round(v/1000)
  return s+"$%d"%round(v)
def gold(scen,patch): return next(c for c in GOLD["cases"] if c["scenario"]==scen and c["patch"]==patch)["result"]
PUB={}
TABLES["publication"]=[]
TABLES["import_row"]=[]; TABLES["import_issue"]=[]
for t in ["assumption_set","outcome","measure_value","survey","survey_result","recurring_cost","project_request","report_snapshot","budget_line"]: TABLES.setdefault(t,[])
TABLES["initiative"]=TABLES["initiative"]+[{"id":"ini-ffa","district_id":"d1","name":"FFA program","type":"program","status":"approved","cost_confidence":"estimate"}]
_bp=[p for p in TABLES["phase"] if p["scenario_id"]!=_ph_sid][0]; _bp["label"]="Unit ventilators"
TABLES["recurring_cost"]=[{"id":"rc1","district_id":"d1","scenario_id":_ph_sid,"initiative_id":_init0,"kind":"supplies","fund":"ppel","first_fy":2028,"last_fy":None,"annual_amount":12000,"grows_with":"none"}]
TABLES["audit_log"]=[{"id":2,"district_id":"d1","table_name":"initiative","row_pk":"x","action":"update","actor":"u-admin","at":"2026-09-30T15:04:00Z",
   "old_row":{"name":"Gym floor","focus_area":"Facilities","updated_at":"a"},"new_row":{"name":"Gym floor","focus_area":"Activities","updated_at":"b"}},
  {"id":1,"district_id":"d1","table_name":"scenario","row_pk":"y","action":"insert","actor":"u-someone-else","at":"2026-09-30T14:00:00Z","old_row":None,"new_row":{"name":"Plan B"}}]
TABLES["access_request"]=[{"id":"req1","district_id":"d1","email":"asker@example.test","message":"New principal at the middle school","created_at":"2026-09-29T15:00:00Z","status":"pending"}]
# state annual-report data (part 16): fictional figures for the fictional demo district
IA_BENCH=[{"measure_key":"exp|General|Student Transportation","grp":"exp","fund":"General","line":"Student Transportation","unit":"per_pupil","flag":"high","peer_group":"districts your size",
           "callout":"$1,210/pupil — 1.8× the average of 41 districts your size (higher than 95% of them)"},
          {"measure_key":"bal|SAVE|ENDING","grp":"bal","fund":"SAVE","line":"ENDING","unit":"per_pupil","flag":"low","peer_group":"districts your size","callout":"$310/pupil — 62% below the average of 41 districts your size (lower than 93% of them)"},
          {"measure_key":"rev|ALL|TOTAL","grp":"rev","fund":"ALL","line":"TOTAL","unit":"per_pupil","flag":"jump","peer_group":"districts your size","callout":"Up 31% from last year; districts your size moved 4% (median)"},
          {"measure_key":"exp|ALL|Instruction","grp":"exp","fund":"ALL","line":"Instruction","unit":"per_pupil","flag":None,"peer_group":"districts your size","callout":None}]
IA_MORE={"de_district":"7777",
  "valuation":{"fiscal_year":2027,"taxable":727000000,"actual":1250000000,"growth":0.031,"growth_years":4,"ag_rollback":0.594401,"res_rollback":0.445345,"source":"Iowa Department of Management, School District Assessed & Taxable Valuations by Class, FY2027"},
  "aid_levy":{"fiscal_year":2027,"budget_enrollment":1190.2,"dcpp":8148,"regular_cost":9697750,"combined_cost":13500000,"other_formula":3802250,"vppel_max":970000,"source":"Iowa Department of Management, Aid and Levy worksheet, FY2027"},
  "unspent":{"fiscal_year":2025,"unspent":2100000,"max_authorized":16000000,"misc_income":1900000,"misc_growth":0.021,"source":"Iowa Department of Management, Unspent Authorized Budget report, FY2025"},
  "car":{"fiscal_year":2025,"source":"Iowa Department of Education, Certified Annual Report (Actual), FY2017 on",
    "ongoing":{"save":{"recurring":180000,"construction":400000,"years":3,"from_fy":2023,"to_fy":2025},"ppel":{"recurring":95000,"construction":0,"years":3,"from_fy":2023,"to_fy":2025}},
    "debt_payments":{"save":480000,"debt":295000},
    "grants":[{"fy":2019,"amount":10000},{"fy":2021,"amount":50000},{"fy":2023,"amount":30000},{"fy":2024,"amount":90000},{"fy":2025,"amount":60000}],
    "save_receipts":[{"fy":2021,"amount":1300000},{"fy":2023,"amount":1390000},{"fy":2025,"amount":1450000}],
    "general":{"salaries":7000000,"benefits":2100000,"benefits_pct":0.3,"sal_instruction":4800000,"sal_admin":700000,"sal_support":1500000,"nonstaff":3300000}},
  "vppel":{"first_fy":2022,"last_fy":2027,"from_start_of_data":False},
  "home_value":{"median_value":187400,"margin":9100,"acs_year":2024,"source":"U.S. Census Bureau, American Community Survey 2020-2024 5-year estimates, median value of owner-occupied homes (B25077)"},
  "construction_inflation":{"value":0.0494,"detail":{"years":3},"source":"U.S. Bureau of Labor Statistics, Producer Price Index: new school building construction (PCU236222236222), Aug 2023 to Aug 2026"}}
IA_PREFILL={"de_district":"7777","fiscal_year":2025,"source":"Iowa Department of Education, Certified Annual Report, FY2025 (year end June 30)","latest_enrollment":{"fiscal_year":2026,"certified":1188.4},
  "general":{"unassigned":2100000,"assigned":150000,"aea_flowthrough":410000},"balances":{"general":2400000,"save":1900000,"ppel":640000},"receipts":{"save":1450000,"ppel":520000}}
TABLES["import_batch"].append({"district_id":"d1","id":"b-reg1","kind":"check_register","period_end":"2026-08-31","file_name":"august_register.csv","status":"applied","uploaded_at":"2026-09-03T12:00:00Z","row_count":3})
TABLES["register_flag"]=[{"id":1,"district_id":"d1","batch_id":"b-reg1","rule":"vendor_name_change","severity":"concern","question":"Vendor V0012 was paid as Midwest Bus Parts before and is now Midwest Bus Parts & Supply. Was the change requested in writing?","status":"open","vendor_key":"midwest bus parts supply","detail":{}},
  {"id":2,"district_id":"d1","batch_id":"b-reg1","rule":"new_vendor","severity":"question","question":"First payment to Hawkeye Roofing LLC ($12,400.00). What was it for, and who approved adding this vendor?","status":"open","vendor_key":"hawkeye roofing","detail":{}},
  {"id":3,"district_id":"d1","batch_id":"b-reg1","rule":"weekend_date","severity":"info","question":"Check 10422 is dated Saturday, Aug 15.","status":"open","vendor_key":None,"detail":{}},
  {"id":4,"district_id":"d1","batch_id":"b-reg1","rule":"duplicate_payment","severity":"concern","question":"Invoice 88231 to Alliant Energy appears twice.","status":"explained","response":"Second check voided; see September.","resolved_at":"2026-09-05T12:00:00Z","vendor_key":"alliant energy","detail":{}}]
TABLES["register_rule"]=[{"id":1,"district_id":None,"rule":"near_threshold","enabled":True,"params":{"threshold":None,"within_pct":10}},{"id":2,"district_id":None,"rule":"split_purchase","enabled":True,"params":{"threshold":None,"days":14}}]
TABLES["vendor_note"]=[]; TABLES["register_line"]=[]
calls=[]
import base64
def tok(uid,aal="aal1"):
  pl=base64.urlsafe_b64encode(json.dumps({"sub":uid,"aal":aal}).encode()).decode().rstrip("=")
  return "tok-%s.%s.sig"%(uid,pl)
def uid_from(req):
  a=req.headers.get("authorization","")
  return a[len("Bearer tok-"):].split(".")[0] if a.startswith("Bearer tok-") else None
FN={"fail":False}
PROMPT={"text":None}
DBFAIL={"phase":False}
def user_obj(email):
  u=USERS[email]; o={"id":u[0],"email":email,"user_metadata":{"full_name":u[1]}}
  if email=="mfa@example.test": o["factors"]=[{"id":"f9","status":"verified","factor_type":"totp"}]
  return o
def email_of(uid): return next((e for e,u in USERS.items() if u[0]==uid),None)
async def handler(route):
  req=route.request; url=urllib.parse.urlparse(req.url); path=url.path; q=urllib.parse.parse_qs(url.query)
  body=None
  try: body=req.post_data
  except Exception: body="<binary>"
  calls.append((req.method,path+"?"+url.query,body,req.headers.get("prefer","")))
  if path.startswith("/storage/v1/object/public/"):   # public logos load like any image, without a key
    import base64 as _b; return await route.fulfill(status=200,content_type="image/png",body=_b.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="))
  assert req.headers.get("apikey")=="sb_publishable_TESTKEY_abcdefghijklmnop", "missing apikey"
  def ok(data,status=200): return route.fulfill(status=status,content_type="application/json",body=json.dumps(data))
  uid=uid_from(req); em=email_of(uid) if uid else None
  if path=="/auth/v1/token":
    b=json.loads(body)
    if q["grant_type"][0]=="password":
      if b["email"] in USERS and b["password"]==PW: return await ok({"access_token":tok(USERS[b["email"]][0]),"refresh_token":"r","expires_in":3600,"user":user_obj(b["email"])})
      return await ok({"error":"invalid_grant","error_description":"Invalid login credentials"},400)
  if path=="/auth/v1/factors" and req.method=="POST": return await ok({"id":"f1","type":"totp","totp":{"qr_code":"<svg xmlns='http://www.w3.org/2000/svg'/>","secret":"JBSWY3DPEHPK3PXP","uri":"otpauth://x"}})
  if path.startswith("/auth/v1/factors/") and path.endswith("/challenge"): return await ok({"id":"ch1","expires_at":9999999999})
  if path.startswith("/auth/v1/factors/") and path.endswith("/verify"):
    b=json.loads(body)
    if b.get("code")!="123456" or b.get("challenge_id")!="ch1": return await ok({"code":"mfa_verification_failed","msg":"Invalid TOTP code entered"},422)
    return await ok({"access_token":tok(uid,"aal2"),"refresh_token":"r2","expires_in":3600,"user":user_obj(em)})
  if path.startswith("/auth/v1/factors/") and req.method=="DELETE": return await ok({"id":path.split("/")[-1]})
  if path=="/functions/v1/send-invitation":
    b=json.loads(body)
    return await (ok({"error":"Email isn’t set up on the server yet."},500) if FN["fail"] else ok({"status":"sent","to":"x"}))
  if path=="/rest/v1/rpc/claim_my_access": return await route.fulfill(status=204,body="")
  if path=="/rest/v1/rpc/set_phase_progress": return await ok(json.loads(body))
  if path=="/rest/v1/rpc/request_access": return await ok("sent")
  if path=="/auth/v1/user":
    if not em: return await ok({"msg":"JWT expired"},401)
    return await ok(user_obj(em))
  if path=="/auth/v1/signup": return await ok({"id":"u-x","email":json.loads(body)["email"]})
  if path=="/auth/v1/recover": return await ok({})
  if path=="/auth/v1/logout": return await route.fulfill(status=204,body="")
  if path.startswith("/storage/v1/object/district-public/") and req.method=="POST":
    return await ok({"Key":path.split("/object/")[1]})
  if path=="/storage/v1/object/district-public" and req.method=="DELETE":
    return await ok([{"name":"x"}])
  if path.startswith("/storage/v1/object/district-files/") and req.method=="POST":
    return await ok({"Key":path.split("/object/")[1]})
  if path=="/rest/v1/rpc/copy_scenario": return await ok("copied-scenario-id")
  if path=="/rest/v1/rpc/ia_benchmark": return await ok(IA_BENCH)
  if path=="/rest/v1/rpc/ia_prefill": return await ok(IA_PREFILL)
  if path=="/rest/v1/rpc/ia_prefill_more": return await ok(IA_MORE)
  if path=="/rest/v1/rpc/ia_levy": return await ok({"de_district":"7777","fiscal_year":2027,"voted_ppel":1.34,"regular_ppel":0.33,"debt_service":2.1,"management":0.9,"total_rate":13.2,"source":"Iowa Department of Management, School Tax Rates, FY2027"})
  if path=="/rest/v1/rpc/ia_district_search": return await ok([{"de_district":"7777","name":"Ironwood Valley","aea":"267","last_fy":2025,"score":0.9}])
  if path=="/rest/v1/rpc/register_summary": return await ok({"total":21851.5,"lines":3,"by_fund":{"10":9451.5,"33":12400},"top_vendors":[{"vendor":"Hawkeye Roofing LLC","label":None,"total":12400}],"flags":{"question":1,"concern":1}})
  if path=="/rest/v1/rpc/register_check": return await ok({"batch_id":json.loads(body)["p_batch"],"open_flags":2,"by_rule":{}})
  if path=="/rest/v1/rpc/register_check_all": return await ok([])
  if path=="/rest/v1/ia_measure": return await ok([{"fiscal_year":2025}])
  if path=="/rest/v1/ia_district": return await ok([{"de_district":"7777","name":"Ironwood Valley","aea":"267"}])
  if path=="/rest/v1/register_rule" and req.method=="GET": return await ok(TABLES["register_rule"])
  if path=="/rest/v1/rpc/apply_import": return await ok({"status":"applied"})
  if path=="/rest/v1/rpc/public_publication":
    b=json.loads(body)
    if b["p_slug"]=="ironwood-valley" and "d1" in PUB: return await ok({"district":D1,"kind":"board_plan","title":PUB["d1"]["title"],"published_at":"2026-09-29T12:00:00Z","payload":PUB["d1"]["payload"]})
    return await ok(None)
  t=path.split("/")[-1]
  if not em: return await ok({"message":"permission denied for table "+t,"code":"42501"},401)
  u=USERS[em]
  if t=="platform_admin": return await ok([{"user_id":u[0]}] if u[3] else [])
  if t=="profile": return await ok([{"user_id":x[0],"email":e,"full_name":x[1],"title":None} for e,x in USERS.items() if (x[0]==u[0] or "in." in url.query)] if req.method=="GET" else [{"user_id":u[0]}])
  if t=="district":
    if req.method=="GET" and em=="empty.staff@example.test": return await ok([])
    if req.method=="POST": b=json.loads(body); b["id"]="d9"; return await ok([b],201)
    if req.method=="PATCH": return await ok([dict(D1,**json.loads(body))])
    return await ok([D1,D2])
  if t=="district_member":
    if "user_id=eq." in url.query and "select=role,district" in url.query:
      dd={"d1":D1,"d2":D2}; return await ok([{"role":r,"district":dd[d]} for d,r in u[2]])
    if req.method=="PATCH" or req.method=="DELETE": return await ok([{"user_id":"x"}])
    did=q.get("district_id",["eq."])[0][3:]
    return await ok([{"user_id":x[0],"role":r,"created_at":"2026-09-01T00:00:00Z"} for e,x in USERS.items() for d,r in x[2] if d==did])
  if t=="invitation":
    if req.method=="POST": b=json.loads(body); b.update(id="inv2",created_at="2026-09-29T00:00:00Z",expires_at="2026-10-29T00:00:00Z"); return await ok([b],201)
    if req.method=="DELETE": return await ok([{"id":"inv1"}])
  if t=="publication" and req.method=="POST":
    pb=json.loads(body); PUB[pb["district_id"]]=pb; return await ok([dict(pb,id="pub1")],201)
  if t=="phase" and req.method=="POST" and DBFAIL["phase"]:
    return await ok({"code":"PGRST204","message":"Could not find the 'label' column of 'phase' in the schema cache"},400)
  if t in TABLES and req.method=="POST":
    b=json.loads(body)
    if isinstance(b,list) and len({tuple(sorted(x.keys())) for x in b})>1:   # like the real database
      return await ok({"code":"PGRST102","message":"All object keys must match"},400)
    return await ok(b if isinstance(b,list) else [b],201)
  if t in TABLES and req.method=="DELETE":
    return await ok([{"deleted":True}])
  if t in TABLES and req.method=="PATCH":
    return await ok([json.loads(body)])
  if t=="hg_health_run" and t in TABLES: return await ok(TABLES[t])
  if t in TABLES:
    did=q.get("district_id",["eq.d1"])[0][3:]
    out=[r for r in TABLES[t] if r.get("district_id")==did]
    if "batch_id" in q:   # like the real database: amounts for one import, or a list of imports
      v=q["batch_id"][0]
      ids=v[4:-1].split(",") if v.startswith("in.(") else [v[3:]]
      out=[r for r in out if r.get("batch_id") in ids]
    return await ok(out)
  return await ok({"message":"unmocked "+t},404)

async def main():
  srv=subprocess.Popen(["python3","-m","http.server","8765"],cwd=ROOT,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL); time.sleep(0.8)
  errs=[]; R=[]
  def check(name,cond,detail=""): R.append((name,bool(cond),detail))
  async with async_playwright() as p:
    b=await p.chromium.launch()
    # 1. not configured
    pg=await b.new_page(); pg.on("pageerror",lambda e:errs.append(str(e)))
    await pg.route("**/fonts.googleapis.com/**",lambda r:r.abort())
    async def blank(route): await route.fulfill(content_type="application/javascript",body="window.HG_CONFIG={supabaseUrl:'https://YOUR-PROJECT-ID.supabase.co',publishableKey:'YOUR-PUBLISHABLE-KEY'};")
    await pg.route("**/config.js",blank); await pg.goto("http://localhost:8765/"); await pg.wait_for_timeout(400)
    check("unconfigured shows setup message", "Not connected yet" in await pg.inner_text("body"))
    await pg.close()
    ctx=await b.new_context(viewport={"width":1360,"height":900})
    async def cfg(route): await route.fulfill(content_type="application/javascript",body="window.HG_CONFIG={supabaseUrl:'%s',publishableKey:'sb_publishable_TESTKEY_abcdefghijklmnop'};"%SB)
    await ctx.route("**/config.js",cfg); await ctx.route(SB+"/**",handler); await ctx.route("**/fonts.g*/**",lambda r:r.abort())
    pg=await ctx.new_page(); pg.on("dialog",lambda dl: asyncio.ensure_future(dl.accept(PROMPT["text"] if PROMPT["text"] is not None else dl.default_value) if dl.type=="prompt" else dl.accept())); pg.on("pageerror",lambda e:errs.append(str(e))); pg.on("console",lambda m: errs.append("console: "+m.text) if m.type=="error" and "fonts" not in m.text and "ERR_FAILED" not in m.text else None)
    await pg.goto("http://localhost:8765/"); await pg.wait_for_timeout(500)
    check("no session goes to sign in", "Sign in" in await pg.inner_text("h1"))
    await pg.screenshot(path=SHOTS+"/signin.png")
    await pg.fill("input[name=email]","admin@example.test"); await pg.fill("input[name=password]","wrong"); await pg.click("button[type=submit]"); await pg.wait_for_timeout(400)
    check("wrong password message", "don't match" in (await pg.inner_text("#toasts")).replace("’","'"))
    await pg.fill("input[name=password]",PW); await pg.click("button[type=submit]"); await pg.wait_for_timeout(700)
    body=await pg.inner_text("body")
    check("access refresh runs on load", any("/rest/v1/rpc/claim_my_access" in c[1] for c in calls))
    check("lands in the district where the person has the most responsibility (admin here, not the district where they only view)", "ironwood-valley/home/today" in pg.url, pg.url)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pg.wait_for_timeout(500); body=await pg.inner_text("body")
    check("overview live counts", "District baseline" in body)
    await pg.click("button[data-action=openHelp]"); await pg.wait_for_timeout(300)
    hp=await pg.inner_text("[data-modal]")
    check("help: in-page help for this screen, with the terms it uses", hp.startswith("Home") and "Needs attention" in hp and "Terms on this screen" in hp and "Gap to close" in hp and "The full guide" in hp, hp[:300])
    await pg.click("[data-modal] button[data-action=closeModal]"); await pg.wait_for_timeout(200)
    gt=await pg.inner_text("[data-setup-guide]") if await pg.locator("[data-setup-guide]").count() else ""
    check("setup guide: steps checked against the district's data, with the next one", "Getting set up" in gt and "of 8 done" in gt and "Next:" in gt, gt[:300])
    await pg.click("[data-setup-guide] button[data-action=setupGo]"); await pg.wait_for_timeout(700)
    bn=await pg.inner_text(".setupbanner") if await pg.locator(".setupbanner").count() else ""
    check("setup guide: each step opens its real screen with where you are and the way on", "Setup, step" in bn and "Back to setup" in bn and "/overview/" not in pg.url, bn[:200]+pg.url)
    await pg.click(".setupbanner button[data-action=setupBack]"); await pg.wait_for_timeout(600)
    await pg.screenshot(path=SHOTS+"/overview.png",full_page=True)
    # every tab
    tabs=await pg.evaluate("""()=>{const out=[];document.querySelectorAll('.rail a[href^="#/d/"]').forEach(a=>out.push(a.getAttribute('href')));return out}""")
    visited=0; bad=[]
    secs=["home/today","plan/priorities","track/measures","plan/community","plan/initiatives","plan/ranking","plan/scenarios","money/summary","money/general","resources/funds","money/capital","resources/assumptions","track/initiatives","progress/measures","track/actuals","settings/uploads","share/board","share/plans","share/community","reports/exports","settings/district","settings/people","settings/account","help/guide"]
    for s in secs:
      await pg.goto("http://localhost:8765/#/d/ironwood-valley/"+s); await pg.wait_for_timeout(250)
      v=await pg.inner_text("#view"); visited+=1
      if "couldn’t load" in v or "Loading" in v: bad.append(s+": "+v[:80])
      if s in("money/capital","settings/people","settings/uploads","resources/funds"): await pg.screenshot(path=SHOTS+"/"+s.replace("/","-")+".png",full_page=True)
    check("all 24 screens render", visited==24 and not bad, "; ".join(bad))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/resources/funds"); await pg.wait_for_timeout(900); await pg.evaluate("document.querySelectorAll('details.fold').forEach(d=>d.open=true)")
    t=await pg.inner_text("#view"); check("latest balance only", "$2,150,000" in t and "$1\n" not in t)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/money/capital"); await pg.wait_for_timeout(500)
    t=await pg.inner_text("#view"); base=gold("orig",None)
    check("capital plan: board version gap matches the planner ($5.35M)", fmtK(base["gap"])=="$5.35M" and "$5.35M" in t and fmtK(base["need"]) in t, t[:300])
    check("capital plan: per-pupil and busiest year", "$11,836" in t and "FY2030" in t)
    await pg.screenshot(path=SHOTS+"/capital.png",full_page=True)
    await pg.click("input[data-lever=sf]"); await pg.wait_for_timeout(200)
    t=await pg.inner_text("#cap-results"); sf=gold("orig",{"sf":False})
    check("capital plan: SF 2472 off recomputes", fmtK(sf["gap"]) in t and "Gap to close" in t, fmtK(sf["gap"]))
    await pg.fill("input[data-lever=infl]","0.04"); await pg.dispatch_event("input[data-lever=infl]","input"); await pg.wait_for_timeout(200)
    t=await pg.inner_text("#cap-results")
    both=None
    await pg.click("button[data-action=capReset]"); await pg.wait_for_timeout(300)
    check("capital plan: reset returns to the scenario", "$5.35M" in await pg.inner_text("#cap-results"))
    sid=next(sc["id"] for sc in TABLES["scenario"] if not sc["is_board_version"])
    await pg.select_option("select[data-cap-scenario]",sid); await pg.wait_for_timeout(400)
    ph=gold("phased",None); t=await pg.inner_text("#view")
    check("capital plan: phased-bond scenario matches ($1.15M gap)", fmtK(ph["gap"])=="$1.15M" and "$1.15M" in t and "Bond" in t, fmtK(ph["gap"]))
    check("capital plan: complete, no unfinished parts left", "Still to come on this screen" not in t and "Projects by year" in t)
    TAXEXP=json.loads(subprocess.check_output(["node","-e","""
      const C=require('./capital.js'),E=require('./engine.js'),D=require('./demo_data.js'),T=require('./tax.js');let i=0;
      const R=C.demoRows(D['ironwood-valley'],'d1',()=>'00000000-0000-4000-8000-'+String(++i).padStart(12,'0'));
      const rows={district:{name:'x'},settings:R.district_settings[0],balances:R.fund_balance,debts:R.debt_obligation,scenarios:R.scenario,initiatives:R.initiative,phases:R.phase,funding:R.phase_funding,financing:R.financing,recurring:[]};
      const sc=R.scenario.find(s=>!s.is_board_version);const inp=C.buildInputs(rows,sc.id);const m=T.impact(inp.cfg,inp.levers,inp.tax);
      process.stdout.write(JSON.stringify({home:m.peak.home,fy:m.peak.fy,farm:m.peak.farm100k,rate:m.peak.rate}));"""],cwd=os.path.dirname(os.path.abspath(__file__))))
    tx=await pg.inner_text("#cap-tax")
    _fd=lambda v: ("$%.2f"%v) if v<100 else "$"+format(round(v),",")
    check("taxpayers: the phased-bond scenario's added cost for a $150,000 home and per $100,000 of farmland", ("$%.2f a year"%TAXEXP["home"]) in tx and (_fd(TAXEXP["farm"])+" per $100,000 of assessed farmland") in tx and "an acre" not in tx and ("FY%d"%TAXEXP["fy"]) in tx and ("$%.4f"%TAXEXP["rate"]) in tx, tx[:300])
    await pg.locator("#cap-tax").screenshot(path=SHOTS+"/tax.png")
    # Phase 2E: table view and filters (phased scenario on screen)
    total_phases=sum(1 for p in TABLES["phase"] if p["scenario_id"]==_ph_sid)
    await pg.click("button[data-action=capView][data-v=table]"); await pg.wait_for_timeout(300)
    rows_n=await pg.locator("table.captable tbody tr").count()
    await pg.locator("#cap-filters").scroll_into_view_if_needed(); await pg.screenshot(path=SHOTS+"/captable.png")
    check("table view: the switch shows which view is on", await pg.get_attribute("button[data-action=capView][data-v=table]","aria-pressed")=="true")
    check("table view: one row per phase, with totals", rows_n==total_phases and "phases" in await pg.inner_text("table.captable tfoot"), f"{rows_n} vs {total_phases}")
    await pg.select_option("select[data-cap-filter=tier]","must"); await pg.wait_for_timeout(300)
    must_rows=await pg.locator("table.captable tbody tr").count()
    tiers=await pg.locator("table.captable tbody tr td:nth-child(4)").all_inner_texts()
    check("filters: priority", 0<must_rows<total_phases and all(x=="Must-have" for x in tiers) and "Showing %d of %d phases"%(must_rows,total_phases) in await pg.inner_text("#cap-years"), str(must_rows))
    await pg.select_option("select[data-cap-filter=tier]",""); await pg.fill("input[data-cap-filter=q]","roof"); await pg.wait_for_timeout(300)
    names=await pg.locator("table.captable tbody tr td:nth-child(2)").all_inner_texts()
    check("filters: search as you type, without losing the search box", names and all("roof" in n.lower() for n in names) and await pg.evaluate("document.activeElement && document.activeElement.matches('[data-cap-filter=q]')"), str(names))
    async with pg.expect_download() as dl: await pg.click("button[data-action=capDownload]")
    f=await dl.value; import openpyxl as _ox, io as _io; _ws=_ox.load_workbook(_io.BytesIO(open(await f.path(),"rb").read())).active; _rows=[[c.value for c in r] for r in _ws.iter_rows()]
    check("table view: download the filtered table (Excel)", f.suggested_filename.endswith(".xlsx") and _rows[0][:4]==["FY","Initiative","Phase","Priority"] and len(_rows)==len(names)+1 and isinstance(_rows[1][0],int), str(_rows[:2]))
    await pg.click("a[data-action=capClearFilters]"); await pg.click("button[data-action=capView][data-v=cards]"); await pg.wait_for_timeout(300)
    await pg.select_option("select[data-cap-filter=fund]","ppel"); await pg.wait_for_timeout(300)
    check("filters: apply to the cards too", "Showing" in await pg.inner_text("#cap-years") and "Year totals still include everything" in await pg.inner_text("#cap-years"))
    await pg.click("a[data-action=capClearFilters]"); await pg.wait_for_timeout(200)
    check("filters: clearing resets the filter menus too", await pg.input_value("select[data-cap-filter=fund]")=="" and await pg.input_value("input[data-cap-filter=q]")=="")
    check("taxpayers: says when a bond's levy ends, past the plan", "continues past the plan’s last year, through FY2050" in tx)
    check("taxpayers: says it's the added cost, and that a bond needs a vote", "added cost only" in tx.lower() and "needs a public vote" in tx)
    # scenario work
    board_id=next(sc["id"] for sc in TABLES["scenario"] if sc["is_board_version"]); phased_id=sid
    await pg.select_option("select[data-cap-scenario]",board_id); await pg.wait_for_timeout(400)
    t=await pg.inner_text("#view")
    check("phase names: shown on the year cards", "· Unit ventilators" in t)
    check("taxpayers: no bond, no new levy, says so", "adds no property tax" in await pg.inner_text("#cap-tax"))
    check("locked scenario: read-only, admin can unlock", "This scenario is locked" in t and await pg.locator(".ylist a").count()==0
          and await pg.locator("button[data-action=unlockScenario]").count()==1 and await pg.locator("button[data-action=deleteScenario]").count()==0
          and await pg.locator("button[data-action=saveLevers]").count()==0)
    await pg.select_option("select[data-cap-scenario]",phased_id); await pg.wait_for_timeout(400)
    links=pg.locator(".ylist a[data-action=editProject]")
    check("unlocked scenario: projects open for editing", await links.count()>=15)
    await pg.click(".ylist a:has-text('Middle school HVAC replacement')"); await pg.wait_for_timeout(300)
    check("editor: opens with the project's phases", await pg.locator("[data-modal] [data-phase-row]").count()==2 and await pg.input_value("[data-modal] input[name=name]")=="Middle school HVAC replacement")
    ab=await pg.evaluate("""(()=>{const t=document.querySelector('[data-modal] [data-phase-body]').closest('.scroll').getBoundingClientRect(),b=document.querySelector('[data-modal] [data-action=addPhaseRow]').getBoundingClientRect();return b.top-t.bottom})()""")
    check("editor: “Add a phase” sits clear of the table above it", ab>=8, str(ab))
    await pg.screenshot(path=SHOTS+"/editor.png")
    first=pg.locator("[data-modal] [data-phase-row]").first
    await first.locator("input[name=pct]").first.fill("60"); n0=len(calls)
    await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(300)
    check("editor: a split that doesn't add to 100% is caught", "add to 60%" in await pg.inner_text("[data-modal] [data-form-errors]") and not any(c[0] in("POST","PATCH","DELETE") for c in calls[n0:]))
    await first.locator("input[name=pct]").first.fill("100"); await first.locator("input[name=cost]").fill("950,000")
    # fund rows: add, even split, two-fund complement, remove
    await first.locator("button[data-action=addFund]").click()
    pv=lambda: first.locator("input[name=pct]").evaluate_all("els=>els.map(e=>e.value)")
    two=await pv()
    await first.locator("button[data-action=addFund]").click(); three=await pv()
    hidden=await first.locator("button[data-action=addFund]").is_hidden()
    await first.locator("[data-src] >> nth=2").locator("button[data-action=removeFund]").click()
    await first.locator("input[name=pct]").first.fill("70"); comp=await pv()
    await first.locator("[data-src] >> nth=1").locator("button[data-action=removeFund]").click(); one=await pv()
    check("funds: add gives 50/50, then 34/33/33, capped at three", two==["50","50"] and three==["34","33","33"] and hidden, str([two,three,hidden]))
    check("funds: with two, the other makes up the rest; removing goes back to 100", comp==["70","30"] and one==["100"], str([comp,one]))
    await first.locator("input[name=label]").fill("Design and bid")
    n0=len(calls); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(800)
    new=calls[n0:]; ops=[(c[0],c[1].split("?")[0].replace("/rest/v1/","")) for c in new if c[0] in ("POST","PATCH","DELETE")]
    phs=json.loads(next(c[2] for c in new if c[0]=="POST" and c[1].startswith("/rest/v1/phase?") or (c[0]=="POST" and c[1]=="/rest/v1/phase?")))
    delq=next((c[1] for c in new if c[0]=="DELETE" and "/rest/v1/phase?" in c[1]),"")
    delq="scenario_id=eq."+phased_id if "id=in.(" in delq else delq
    check("phase names: saved with the phase", phs[0].get("label")=="Design and bid", str(phs[0]))
    # if the database is missing an update, the message is plain and nothing is removed
    DBFAIL["phase"]=True
    await pg.click(".ylist a:has-text('Middle school HVAC replacement')"); await pg.wait_for_timeout(300)
    n0=len(calls); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(600)
    tt=(await pg.inner_text("#toasts")).replace("’","'")
    check("errors: database out of date is explained plainly", "HighGround's database needs an update before this can be saved. Nothing was changed." in tt and "schema cache" not in tt, tt[-200:])
    check("errors: a failed save removes nothing", not any(c[0]=="DELETE" for c in calls[n0:]))
    DBFAIL["phase"]=False
    await pg.click("[data-modal] button[data-action=closeModal] >> nth=0"); await pg.wait_for_timeout(200)
    check("editor: save updates details and replaces this scenario's phases", ops==[("PATCH","initiative"),("POST","phase"),("POST","phase_funding"),("POST","recurring_cost"),("DELETE","phase"),("DELETE","recurring_cost")]
          and phs[0]["cost"]==950000 and ("scenario_id=eq."+phased_id) in delq, str(ops))
    # funds as choices: "whichever has room", in order
    yl=await pg.inner_text("#cap-years")
    pg_li=pg.locator(".ylist li:has-text('Playground surfacing')").first
    chips=await pg_li.locator(".chips").inner_text(); ttl=await pg_li.locator(".chips").get_attribute("title") or ""
    check("fund choices: the card shows the fund used, and the other choice", "Grants" in chips and "or PPEL" in chips and "whichever has room" in ttl, chips+" | "+ttl)
    await pg.click(".ylist a:has-text('Playground surfacing')"); await pg.wait_for_timeout(300)
    fm=pg.locator("[data-modal] [data-phase-row] select[name=fmode]").first
    check("fund choices: the editor opens in 'whichever has room' with the funds in order", await fm.is_visible() and await fm.input_value()=="either"
          and await pg.locator("[data-modal] [data-phase-row] [data-srcs].either").count()==1 and not await pg.locator("[data-modal] [data-phase-row] input[name=pct]").first.is_visible())
    await pg.click("[data-modal] button[data-action=closeModal] >> nth=0"); await pg.wait_for_timeout(200)
    await pg.click(".ylist a:has-text('Middle school HVAC replacement')"); await pg.wait_for_timeout(300)
    first=pg.locator("[data-modal] [data-phase-row]").first
    hid1=await first.locator("select[name=fmode]").is_hidden()
    await first.locator("button[data-action=addFund]").click()
    await first.locator("select[name=src] >> nth=1").select_option("ppel")
    await first.locator("select[name=fmode]").select_option("either"); await pg.wait_for_timeout(100)
    n0=len(calls); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(800)
    phs=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/phase"]
    fus=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/phase_funding"]
    check("fund choices: one fund has no mode picker; two funds set to 'whichever has room' save as choices, first choice at 100%",
          hid1 and phs and phs[0][0].get("fund_options")==["save","ppel"] and all("fund_options" in p for p in phs[0])
          and [f for f in fus[0] if f["phase_id"]==phs[0][0]["id"]]==[{"phase_id":phs[0][0]["id"],"district_id":phs[0][0]["district_id"],"fund":"save","pct":100}], str(phs)[:300]+str(fus)[:300])
    await pg.click("button[data-action=editProject][data-id='']"); await pg.wait_for_timeout(300)
    await pg.fill("[data-modal] input[name=name]","Library HVAC")
    row=pg.locator("[data-modal] [data-phase-row]").first
    await row.locator("select[name=fy]").select_option("2029"); await row.locator("input[name=cost]").fill("50000")
    n0=len(calls); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(800)
    ops=[c[1].split("?")[0].replace("/rest/v1/","") for c in calls[n0:] if c[0]=="POST"]
    newph=json.loads(next(c[2] for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/phase"))
    check("add a project: new project joins this scenario", ops==["initiative","scenario_initiative","phase","phase_funding"] and newph[0]["fy"]==2029 and newph[0]["scenario_id"]==phased_id, str(ops))
    # picker: add an existing initiative that isn't in this scenario
    await pg.click("button[data-action=editProject][data-id='']"); await pg.wait_for_timeout(300)
    check("picker: existing initiatives not in this scenario are offered", await pg.locator("[data-modal] select[data-pick-init] option", has_text="FFA program").count()==1)
    await pg.select_option("[data-modal] select[data-pick-init]","ini-ffa"); await pg.wait_for_timeout(300)
    check("picker: choosing one fills in its details", await pg.input_value("[data-modal] input[name=name]")=="FFA program" and "Add to this scenario" in await pg.inner_text("[data-modal] h2"))
    row=pg.locator("[data-modal] [data-phase-row]").first
    await row.locator("input[name=label]").fill("Shop buildout"); await row.locator("input[name=cost]").fill("150000")
    n0=len(calls); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(800)
    ops=[(c[0],c[1].split("?")[0].replace("/rest/v1/","")) for c in calls[n0:] if c[0] in ("POST","PATCH")]
    check("picker: saving adds it to the scenario without a duplicate", ("PATCH","initiative") in ops and ("POST","scenario_initiative") in ops and ("POST","initiative") not in ops, str(ops))
    # Phase 2B: yearly costs and programs without phases
    await pg.click("button[data-action=editProject][data-id='']"); await pg.wait_for_timeout(300)
    await pg.fill("[data-modal] input[name=name]","Library aide"); await pg.select_option("[data-modal] select[name=type]","staff")
    await pg.click("[data-modal] button[data-action=removePhaseRow]"); await pg.click("[data-modal] button[data-action=addYearlyRow]")
    yr=pg.locator("[data-modal] [data-yearly-row]").last
    await yr.locator("select[name=y_kind]").select_option("salary"); await yr.locator("select[name=y_fund]").select_option("general")
    await yr.locator("input[name=y_amount]").fill("32,000"); await yr.locator("select[name=y_first]").select_option("2028")
    n0=len(calls); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(800)
    ops=[c[1].split("?")[0].replace("/rest/v1/","") for c in calls[n0:] if c[0]=="POST"]
    rcp=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/recurring_cost"]
    ini=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/initiative"]
    check("yearly costs: a hire with only yearly costs, no phases", ops==["initiative","scenario_initiative","recurring_cost"] and ini[0]["type"]=="staff"
          and rcp[0][0]["annual_amount"]==32000 and rcp[0][0]["fund"]=="general" and rcp[0][0]["first_fy"]==2028 and rcp[0][0]["last_fy"] is None, str(ops)+str(rcp))
    await pg.click("[data-action=editProject][data-id='"+_init0+"'] >> nth=0"); await pg.wait_for_timeout(300)
    check("yearly costs: editor shows the scenario's existing yearly cost", await pg.locator("[data-modal] [data-yearly-row]").count()==1 and await pg.input_value("[data-modal] input[name=y_amount]")=="12,000")
    await pg.fill("[data-modal] input[name=y_amount]","0"); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(300)
    check("yearly costs: a zero amount is caught", "amount per year" in await pg.inner_text("[data-modal] [data-form-errors]"))
    await pg.click("[data-modal] button[data-action=closeModal] >> nth=0"); await pg.wait_for_timeout(200)
    t=await pg.inner_text("#cap-yearly")
    check("capital plan: yearly costs card", "Yearly costs" in t and "FY2028" in t and "$12k" in t and "PPEL" in t, t[:200])
    await pg.click("button[data-action=editFinancing][data-id='']"); await pg.wait_for_timeout(300)
    await pg.select_option("[data-modal] select[name=kind]","lease"); await pg.fill("[data-modal] input[name=amount]","300,000")
    await pg.fill("[data-modal] input[name=rate]","5"); await pg.fill("[data-modal] input[name=years]","5")
    n0=len(calls); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(700)
    fp=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/financing"]
    check("financing: lease saved, repaid from PPEL", fp and fp[0]["kind"]=="lease" and fp[0]["repay_from"]=="ppel" and fp[0]["rate"]==0.05 and fp[0]["amount"]==300000 and fp[0]["scenario_id"]==phased_id, str(fp))
    await pg.click("input[data-lever=sf]"); n0=len(calls)
    await pg.click("button[data-action=saveLevers]"); await pg.wait_for_timeout(600)
    lp=[json.loads(c[2]) for c in calls[n0:] if c[0]=="PATCH" and "/rest/v1/scenario?id=eq."+phased_id in c[1]]
    check("levers: saved to the scenario", lp and lp[0]["lever_sf2472"] is False and lp[0]["lever_ppel_growth"]==0.035, str(lp))
    n0=len(calls); await pg.click("button[data-action=copyScenario]"); await pg.wait_for_timeout(600)
    cp=[json.loads(c[2]) for c in calls[n0:] if "rpc/copy_scenario" in c[1]]
    check("copy: makes a new scenario from this one", cp and cp[0]["p_source"]==phased_id and cp[0]["p_name"].endswith("(copy)"), str(cp))
    await pg.select_option("select[data-cap-scenario]",phased_id); await pg.wait_for_timeout(400)
    n0=len(calls); await pg.click("button[data-action=makeBoard]"); await pg.wait_for_timeout(700)
    bp=[(c[1].split("id=eq.")[1], json.loads(c[2])) for c in calls[n0:] if c[0]=="PATCH" and "/rest/v1/scenario?" in c[1]]
    check("board version: old one cleared first, then the new one set", bp==[(board_id,{"is_board_version":False}),(phased_id,{"is_board_version":True})], str(bp))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/share/community"); await pg.wait_for_timeout(400)
    check("community page: nothing published yet", "Nothing is published" in await pg.inner_text("#view"))
    check("publish: holding back with nothing approved would show nothing, so Publish waits", "Nothing would show" in await pg.inner_text("#view") and await pg.is_disabled("button[data-action=publishBoard]"))
    appr=[i for i in TABLES["initiative"] if i["name"].startswith("Middle school HVAC") or i["name"].startswith("Secure entrances")]
    for i in appr: i["status"]="approved"
    await pg.reload(); await pg.wait_for_timeout(700)
    t=await pg.inner_text("#view")
    check("publish: shows how many will show and how many are held back, and the public gap", "2 initiatives will show; 14 held back" in t and "The public page’s gap is" in t, t[t.find("Publish the community page"):t.find("Publish the community page")+500])
    n0=len(calls); await pg.click("button[data-action=publishBoard]"); await pg.wait_for_timeout(700)
    pp=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and "/rest/v1/publication" in c[1]]
    check("publish: held-back proposals are left out of the public copy", pp and len(pp[0]["payload"]["projects"])==2 and pp[0]["payload"]["heldBack"]==14 and pp[0]["payload"]["holdBack"], str(pp)[:200])
    await pg.fill("textarea[data-pub-note]","The board adopts the plan in October. Thank you for the survey responses.")
    await pg.uncheck("input[data-pub-hold]"); await pg.wait_for_timeout(700)
    check("publish: the note survives ticking and unticking", await pg.input_value("textarea[data-pub-note]")=="The board adopts the plan in October. Thank you for the survey responses.")
    await pg.click("button[data-action=previewPublic]"); await pg.wait_for_timeout(600)
    pv=await pg.inner_text("[data-modal]")
    check("publish: preview shows exactly what the public will see", "Preview: what the public will see" in pv and "The board adopts the plan in October" in pv and "Capital plan:" in pv)
    await pg.click("[data-modal] button[data-action=closeModal] >> nth=0"); await pg.wait_for_timeout(200)
    for i in appr: i["status"]="proposed"
    n0=len(calls); await pg.click("button[data-action=publishBoard]"); await pg.wait_for_timeout(700)
    pp=[c for c in calls[n0:] if c[0]=="POST" and "/rest/v1/publication" in c[1]]
    pl=json.loads(pp[0][2])["payload"] if pp else {}
    check("publish: sends the board version as a frozen copy", pp and len(pl.get("projects",[]))==16 and pl["settings"]["save"]["receipts"]==1420000 and json.loads(pp[0][2])["kind"]=="board_plan" and pl["note"].startswith("The board adopts"), str(len(pl.get("projects",[]))))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/plan/initiatives"); await pg.wait_for_timeout(300)
    leftovers={}
    for route in ["home/today","plan/priorities","track/measures","plan/community","plan/initiatives","money/summary","money/general","money/capital",
                  "track/initiatives","progress/measures","track/actuals","settings/uploads","share/board","share/community","reports/exports","settings/district"]:
      await pg.goto("http://localhost:8765/#/d/ironwood-valley/"+route); await pg.wait_for_timeout(350)
      n=await pg.locator("[data-notbuilt]").count()
      if n: leftovers[route]=n
    check("no unbuilt buttons left anywhere", not leftovers, str(leftovers))
    # invite
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/people"); await pg.wait_for_timeout(400)
    await pg.fill("form[data-form=invite] input[name=email]","New.Person@Example.test"); await pg.select_option("form[data-form=invite] select","editor")
    await pg.click("form[data-form=invite] button"); await pg.wait_for_timeout(400)
    post=[c for c in calls if c[0]=="POST" and "/rest/v1/invitation" in c[1]]
    check("invite posts lowercased email and role", post and json.loads(post[-1][2])=={"district_id":"d1","email":"new.person@example.test","role":"editor"}, post[-1][2] if post else "")
    fnc=[c for c in calls if "/functions/v1/send-invitation" in c[1]]
    check("invite: emails the invitation through the server function", fnc and json.loads(fnc[-1][2])=={"invitation_id":"inv2"} and "Invitation emailed" in await pg.inner_text("#toasts"))
    FN["fail"]=True
    await pg.fill("form[data-form=invite] input[name=email]","second@example.test"); await pg.click("form[data-form=invite] button"); await pg.wait_for_timeout(500)
    check("invite: if the email can't go, says so and keeps the invitation", "Invitation saved, but not emailed" in await pg.inner_text("#toasts"))
    FN["fail"]=False
    n0=len(calls); await pg.click("button[data-action=resendInvite][data-id=inv1]"); await pg.wait_for_timeout(500)
    check("invite: send again", any("/functions/v1/send-invitation" in c[1] and json.loads(c[2])["invitation_id"]=="inv1" for c in calls[n0:]))
    t=await pg.inner_text("#view")
    check("access requests: listed for admins", "Asking for access" in t and "asker@example.test" in t and "New principal" in t)
    await pg.select_option("select[data-req-role=req1]","board"); n0=len(calls)
    await pg.click("button[data-action=approveRequest][data-id=req1]"); await pg.wait_for_timeout(600)
    ai=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and "/rest/v1/invitation" in c[1]]
    ap=[json.loads(c[2]) for c in calls[n0:] if c[0]=="PATCH" and "/rest/v1/access_request" in c[1]]
    check("access requests: approve invites them with the chosen role", ai==[{"district_id":"d1","email":"asker@example.test","role":"board"}] and ap and ap[0]["status"]=="approved", str(ai)+str(ap))
    await pg.click("text=Cancel"); await pg.wait_for_timeout(300)
    check("cancel invitation sends delete", any(c[0]=="DELETE" and "invitation?id=eq.inv1" in c[1] for c in calls))
    await pg.select_option("select[data-role-for='u-viewer']","editor"); await pg.wait_for_timeout(300)
    check("role change sends patch", any(c[0]=="PATCH" and "district_member" in c[1] and '"editor"' in (c[2] or "") for c in calls))
    # district save
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/district"); await pg.wait_for_timeout(300)
    await pg.fill("input[name=county]","Harrison-free County"); await pg.click("form[data-form=saveDistrict] button[type=submit]"); await pg.wait_for_timeout(400)
    check("district save sends patch", any(c[0]=="PATCH" and "/rest/v1/district?" in c[1] and "Harrison-free" in (c[2] or "") for c in calls))
    await pg.fill("input[name=allowed_domains]","@IronwoodValley.k12.ia.us, ivcsd.org"); await pg.select_option("select[name=domain_role]","board")
    n0=len(calls); await pg.click("form[data-form=saveDistrict] button[type=submit]"); await pg.wait_for_timeout(500)
    dp=[json.loads(c[2]) for c in calls[n0:] if c[0]=="PATCH" and "/rest/v1/district?" in c[1]]
    check("domain allow-list: saved with role", dp and dp[0]["allowed_domains"]==["IronwoodValley.k12.ia.us","ivcsd.org"] and dp[0]["domain_role"]=="board", str(dp))
    # switch district: admin is viewer in d2
    await pg.select_option("select[data-switch]","cottonwood-ridge"); await pg.wait_for_timeout(500)
    await pg.goto("http://localhost:8765/#/d/cottonwood-ridge/settings/people"); await pg.wait_for_timeout(400)
    t=await pg.inner_text("#content"); check("viewer role: staff settings don't open at all", "Invite someone" not in t and "Not part of your view" in t, t[:200])
    await pg.goto("http://localhost:8765/#/d/no-such-district/home/today"); await pg.wait_for_timeout(300)
    check("unknown district message", "don’t have access" in await pg.inner_text("body"))
    # uploads: projects (CSV), errors, Excel, balances, template
    import openpyxl, io
    TEMPLATE=subprocess.check_output(["node","-e","process.stdout.write(require('./uploads.js').projectTemplate(2027))"],cwd=os.path.dirname(os.path.abspath(__file__)))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/uploads"); await pg.wait_for_timeout(400)
    await pg.select_option("select[data-upload-kind]","projects")
    await pg.set_input_files("input[data-upload-file]",files=[{"name":"projects.csv","mimeType":"text/csv","buffer":TEMPLATE}]); await pg.wait_for_timeout(500)
    t=await pg.inner_text("#upload-review")
    check("upload: review shows what was read", "5 projects, 6 phases" in t and "midpoint" in t and "Track resurface" in t, t[:200])
    await pg.screenshot(path=SHOTS+"/upload-review.png",full_page=True)
    n0=len(calls); await pg.click("button[data-action=applyUpload]"); await pg.wait_for_timeout(900)
    new=calls[n0:]; seq=[(c[0], c[1].split("?")[0].replace("/rest/v1/","").replace("/storage/v1/object/","storage:")) for c in new if c[0] in ("POST","PATCH")]
    names=[x[1] if not x[1].startswith("storage:") else "storage" for x in seq]
    want=["storage","import_batch","import_row","import_issue","initiative","scenario","scenario_initiative","phase","phase_funding","rpc/apply_import"]
    check("upload: file kept, batch recorded, scenario built, then applied", names==want, str(names))
    body=lambda t: json.loads(next(c[2] for c in new if c[0]=="POST" and c[1].split("?")[0].endswith("/"+t)))
    ph=body("phase"); fu=body("phase_funding"); sc=body("scenario"); bt=body("import_batch")
    check("upload: 6 phases with absolute years, 8 fund rows; “SAVE or PPEL” saved as choices", len(ph)==6 and sorted(p["fy"] for p in ph)==[2027,2028,2028,2029,2030,2031] and len(fu)==8
          and sum(1 for p in ph if p.get("fund_options")==["save","ppel"])==1 and all("fund_options" in p for p in ph), str([p["fy"] for p in ph]))
    check("upload: new scenario, not the board version (one exists)", sc["is_board_version"] is False and sc["name"].startswith("Uploaded"), str(sc))
    check("upload: batch points at the stored file", bt["storage_path"].startswith("d1/imports/") and bt["status"]=="review" and bt["kind"]=="projects")
    check("upload: lands on the capital plan", "/money/capital" in pg.url)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/uploads"); await pg.wait_for_timeout(400)
    bad=b"Project,FY,Estimate,Funding source\r\nFar away,FY2040,1000,SAVE\r\n"
    await pg.select_option("select[data-upload-kind]","projects")
    await pg.set_input_files("input[data-upload-file]",files=[{"name":"bad.csv","mimeType":"text/csv","buffer":bad}]); await pg.wait_for_timeout(400)
    check("upload: errors block Apply", "outside this plan" in await pg.inner_text("#upload-review") and await pg.is_disabled("button[data-action=applyUpload]"))
    wb=openpyxl.Workbook(); ws=wb.active
    for r in [["Project","FY","Estimate","Funding source","Funding %"],["Chiller replacement","FY2029",640000,"SAVE",1],["Parking lot","2030-31","$95k","PPEL",1]]: ws.append(r)
    bio=io.BytesIO(); wb.save(bio)
    await pg.select_option("select[data-upload-kind]","projects")
    await pg.set_input_files("input[data-upload-file]",files=[{"name":"projects.xlsx","mimeType":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet","buffer":bio.getvalue()}]); await pg.wait_for_timeout(600)
    t=await pg.inner_text("#upload-review")
    check("upload: Excel files are read", "2 projects" in t and "Chiller replacement" in t and "FY2031: $95,000" in t, t[:200])
    await pg.click("button[data-action=cancelUpload]")
    await pg.select_option("select[data-upload-kind]","balances")
    balcsv=b"Fund,Iowa fund code,Balance,As-of date\r\nSAVE,33,\"$2,300,000\",6/30/2026\r\nPPEL,36,410000,\r\n"
    await pg.set_input_files("input[data-upload-file]",files=[{"name":"balances.csv","mimeType":"text/csv","buffer":balcsv}]); await pg.wait_for_timeout(400)
    check("upload: balances review with date from the file", await pg.input_value("input[data-upload-asof]")=="2026-06-30")
    n0=len(calls); await pg.click("button[data-action=applyUpload]"); await pg.wait_for_timeout(800)
    fb=[c for c in calls[n0:] if c[0]=="POST" and "/rest/v1/fund_balance?on_conflict=" in c[1]]
    fbb=json.loads(fb[0][2]) if fb else []
    check("upload: balances saved as an upload, linked to the batch", len(fbb)==4 and next(x for x in fbb if x["fund"]=="save")["amount"]==2300000 and all(x["source"]=="upload" and x["import_batch_id"] for x in fbb))
    async with pg.expect_download() as dl:
      await pg.click("a[data-action=downloadTemplate][data-kind=projects]")
    d=await dl.value
    check("upload: template downloads as Excel", d.suggested_filename=="highground-projects-template.xlsx")
    await pg.select_option("select[data-upload-kind]","projects")
    await pg.set_input_files("input[data-upload-file]",files=[{"name":"highground-projects-template.xlsx","mimeType":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet","buffer":open(await d.path(),"rb").read()}]); await pg.wait_for_timeout(500)
    t=await pg.inner_text("#upload-review")
    check("upload: the Excel template uploads back as is", "5 projects, 6 phases" in t, t[:200])
    await pg.click("button[data-action=cancelUpload]")
    # Phase 3A: monthly GL export
    GLCSV=subprocess.check_output(["node","-e","process.stdout.write(require('./uploads.js').toCSV(require('./gl.js').sampleExport()))"],cwd=os.path.dirname(os.path.abspath(__file__)))
    TABLES["gl_account"]=[]; TABLES["gl_amount"]=[]
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pg.wait_for_timeout(200)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/uploads"); await pg.wait_for_timeout(500)
    check("GL: offered first to business staff", await pg.input_value("select[data-upload-kind]")=="gl_monthly")
    await pg.set_input_files("input[data-upload-file]",files=[{"name":"gl-2026-09.csv","mimeType":"text/csv","buffer":GLCSV}]); await pg.wait_for_timeout(600)
    rv=await pg.inner_text("#upload-review")
    check("GL: every account is new the first time, each with a suggestion", "21 new accounts to check" in rv and "SAVE spending, facilities acquisition" in rv, rv[:300])
    await pg.fill("input[data-upload-asof]","2026-09-30"); await pg.dispatch_event("input[data-upload-asof]","change"); await pg.wait_for_timeout(200)
    bal=await pg.inner_text("#gl-balances")
    check("GL: balances previewed before applying (fund balance + revenue − spending)", "$1,309,086" in bal and "$788,200" in bal and "Will update" in bal and "Sep 30, 2026" in bal, bal[:400])
    await pg.locator("#upload-review").screenshot(path=SHOTS+"/gl-review.png")
    await pg.select_option("select[data-gl-map='33-0000-000-0000-101']","fund_balance"); await pg.wait_for_timeout(200)
    check("GL: changing an account updates the preview at once", "$2,510,968" in await pg.inner_text("#gl-balances"))
    await pg.select_option("select[data-gl-map='33-0000-000-0000-101']","ignore"); await pg.wait_for_timeout(200)
    n0=len(calls); await pg.click("button[data-action=applyUpload]"); await pg.wait_for_timeout(1200)
    new=calls[n0:]
    ib=[json.loads(c[2]) for c in new if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/import_batch"]
    ga=[json.loads(c[2]) for c in new if c[0]=="POST" and c[1].startswith("/rest/v1/gl_account")]
    gm=[json.loads(c[2]) for c in new if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/gl_amount"]
    lay=[json.loads(c[2]) for c in new if c[0]=="PATCH" and "/rest/v1/district_settings?" in c[1]]
    check("GL: the import is kept, dated, as a monthly GL", ib and ib[0]["kind"]=="gl_monthly" and ib[0]["period_end"]=="2026-09-30" and ib[0]["fiscal_year"]==2027, str(ib)[:200])
    check("GL: accounts saved with what they are, and not asked about again", ga and len(ga[0])==21 and "on_conflict=district_id%2Ccode" in next(c[1] for c in new if c[1].startswith("/rest/v1/gl_account"))
          and next(a for a in ga[0] if a["code"]=="33-0000-000-0000-101")["maps_to"]=="ignore" and all(not a["needs_review"] for a in ga[0]), str(ga)[:300])
    ids={a["code"]:a["id"] for a in ga[0]} if ga else {}
    check("GL: every month's amounts saved against their accounts", gm and len(gm[0])==21 and all(x["account_id"] in ids.values() for x in gm[0]) and any(x["ytd_amount"]==295236 for x in gm[0]))
    check("GL: the layout is remembered", lay and lay[0]["gl_layout"]["cols"]["ytd"]==3)
    check("GL: applied, with the balances it updated named", any("rpc/apply_import" in c[1] for c in new) and "Balances updated: SAVE, PPEL, Debt Service, General Fund" in await pg.inner_text("#toasts"))
    # month two: the same accounts and layout are remembered
    TABLES["gl_account"]=[dict(a, needs_review=False) for a in ga[0]]
    for st0 in TABLES["district_settings"]: st0["gl_layout"]=lay[0]["gl_layout"]
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pg.wait_for_timeout(200)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/uploads"); await pg.wait_for_timeout(500)
    await pg.set_input_files("input[data-upload-file]",files=[{"name":"gl-2026-10.csv","mimeType":"text/csv","buffer":GLCSV}]); await pg.wait_for_timeout(600)
    rv=await pg.inner_text("#upload-review")
    check("GL: next month, nothing new to check, same layout", "No new accounts" in rv and "21 accounts matched from earlier months" in rv and "the same layout as last time" in rv, rv[:300])
    await pg.click("button[data-action=cancelUpload]"); await pg.wait_for_timeout(200)
    # a big first month: grouped and folded, with unrecognised accounts open
    TABLES["gl_account"]=[]
    lines=[GLCSV.decode().rstrip()]+["10-0%03d-1100-100-0000-111,Salaries building %d,1.00,%d.00,,"%(i,i,1000+i) for i in range(1,60)]+["XYZ-9,Odd account,1.00,5.00,,"]
    await pg.set_input_files("input[data-upload-file]",files=[{"name":"big.csv","mimeType":"text/csv","buffer":("\r\n".join(lines)+"\r\n").encode()}]); await pg.wait_for_timeout(700)
    grp=pg.locator("details.glgroup")
    titles=await grp.locator("summary").all_inner_texts()
    opened=[await grp.nth(i).get_attribute("open") is not None for i in range(await grp.count())]
    check("GL: a long first month is grouped, unrecognised accounts first and open, the rest folded", titles and titles[0].startswith("Not recognised: 1 account") and opened[0] and not all(opened[1:]) and any(t.startswith("General Fund: spending, 6") for t in titles), str(list(zip(titles,opened)))[:300])
    await pg.click("button[data-action=cancelUpload]"); await pg.wait_for_timeout(200)
    TABLES["gl_account"]=[]; TABLES["gl_amount"]=[]
    # starting numbers
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/setup"); await pg.wait_for_timeout(500)
    v=await pg.input_value("input[name=save_receipts]"); g=await pg.input_value("input[name=ppel_growth]")
    check("setup: prefilled from the database", v=="1,420,000" and g=="3.5" and await pg.input_value("input[name=as_of]")=="2026-07-01", v+" "+g)
    await pg.screenshot(path=SHOTS+"/setup.png",full_page=True)
    await pg.fill("input[name=save_receipts]",""); n0=len(calls)
    await pg.click("form[data-form=saveSetup] button[type=submit]"); await pg.wait_for_timeout(300)
    err=await pg.inner_text("[data-setup-errors]")
    check("setup: missing SAVE receipts is caught before saving", "SAVE receipts is needed" in err and not any(c[0]=="POST" for c in calls[n0:]))
    await pg.fill("input[name=save_receipts]","$1,420,000"); await pg.fill("input[name=construction_inflation]","3")
    await pg.click("button[data-action=addDebtRow]")
    rows=pg.locator("[data-debt-body] [data-debt-row]"); last=rows.nth(await rows.count()-1)
    await last.locator("input[name=debt_name]").fill("Bus lease-purchase"); await last.locator("select[name=debt_fund]").select_option("ppel")
    await last.locator("input[name=debt_annual]").fill("45,000"); await last.locator("input[name=debt_final]").fill("2029")
    n0=len(calls); await pg.click("form[data-form=saveSetup] button[type=submit]"); await pg.wait_for_timeout(700)
    new=calls[n0:]
    st=[c for c in new if c[0]=="POST" and "/rest/v1/district_settings?on_conflict=district_id" in c[1]]
    fb=[c for c in new if c[0]=="POST" and "/rest/v1/fund_balance?on_conflict=" in c[1]]
    dp=[c for c in new if c[0]=="PATCH" and "/rest/v1/debt_obligation?" in c[1]]
    di=[c for c in new if c[0]=="POST" and c[1].startswith("/rest/v1/debt_obligation?")]
    ok1=False
    if st:
      sb=json.loads(st[0][2])[0]
      ok1=(sb["save_receipts"]==1420000 and sb["ppel_growth"]==0.035 and sb["construction_inflation"]==0.03 and sb["plan_start_fy"]==2027
           and sb["grants_yield"]==0.75 and sb["vppel_status"]=="active" and "merge-duplicates" in st[0][3])
    check("setup: saves settings with % converted and plan start derived", ok1, st[0][2][:200] if st else "no settings POST")
    check("setup: tax estimate settings saved (farmland is per $100,000, no per-acre field)", st and json.loads(st[0][2])[0].get("tax_home_value")==150000 and "ag_value_per_acre" not in json.loads(st[0][2])[0], st[0][2][-200:] if st else "")
    fbb=json.loads(fb[0][2]) if fb else []
    check("setup: saves all four balances for the date", len(fbb)==4 and all(x["as_of"]=="2026-07-01" for x in fbb) and next(x for x in fbb if x["fund"]=="save")["amount"]==2150000)
    check("setup: updates the existing debt and adds the new one", len(dp)==1 and len(di)==1 and json.loads(di[0][2])["fund"]=="ppel" and json.loads(di[0][2])["annual_payment"]==45000)
    await pg.goto("http://localhost:8765/#/d/cottonwood-ridge/settings/setup"); await pg.wait_for_timeout(500)
    check("setup: a viewer can't open Starting numbers", await pg.locator("form[data-form=saveSetup]").count()==0 and "Not part of your view" in await pg.inner_text("#content"))
    # two-step sign-in: turning it on
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/account"); await pg.wait_for_timeout(400)
    await pg.click("button[data-action=mfaOn]"); await pg.wait_for_timeout(400)
    check("two-step: setup shows a QR code and the key", await pg.locator("[data-modal] img[alt^='QR code']").count()==1 and "JBSWY3DPEHPK3PXP" in await pg.inner_text("[data-modal]"))
    await pg.fill("[data-modal] input[name=code]","000000"); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(400)
    check("two-step: a wrong code is explained", "That code didn’t work" in await pg.inner_text("#toasts") and await pg.locator("[data-modal]").count()==1)
    await pg.fill("[data-modal] input[name=code]","123 456"); n0=len(calls); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(600)
    seq=[c[1].split("?")[0] for c in calls[n0:] if "/factors/" in c[1]]
    check("two-step: challenge, then verify, then on", seq==["/auth/v1/factors/f1/challenge","/auth/v1/factors/f1/verify"] and "Two-step sign-in is on" in await pg.inner_text("#toasts"), str(seq))
    # Phase 2A: compare scenarios
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/plan/scenarios"); await pg.wait_for_timeout(600)
    t=await pg.inner_text("#cmp-table")
    check("compare: both scenarios side by side with their gaps", "$5.35M" in t and "$1.15M" in t and "smallest" in t, t[:200])
    check("compare: added tax row", "Added tax, example home" in t and "/yr" in t)
    check("compare: yearly costs row", "Yearly costs (first year they start)" in t and "FY2028" in t and "$12k capital" in t, t[t.find("Yearly"):t.find("Yearly")+120])
    check("compare: what each asks of the community", "general-obligation bond of $4.20M in FY2030, which needs a public vote" in t and "not yet paid for" in t)
    w=await pg.inner_text("#cmp-why")
    check("compare: why the gap differs, in plain words", "lowers the gap by $4.20M" in w and "Financing adds" in w and "no effect on its own" in w, w[:300])
    await pg.screenshot(path=SHOTS+"/compare.png",full_page=True)
    boxes=pg.locator("[data-cmp-pick]")
    await boxes.nth(1).uncheck(); await pg.wait_for_timeout(300)
    check("compare: unticking a scenario removes its column", "$1.15M" not in await pg.inner_text("#cmp-table"))
    await boxes.nth(1).check(); await pg.wait_for_timeout(300)
    await pg.select_option("select[data-why=a]", index=1); await pg.wait_for_timeout(300)
    w=await pg.inner_text("#cmp-why")
    check("compare: either direction can be explained", "raises the gap by $4.20M" in w, w[:200])
    await pg.click("a[data-action=openScenario] >> nth=1"); await pg.wait_for_timeout(600)
    check("compare: Open goes to that scenario on the capital plan", "/money/capital" in pg.url)
    # Phase 2B: All initiatives
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/plan/initiatives"); await pg.wait_for_timeout(700)
    capname=await pg.evaluate("(()=>{const s=document.querySelector('[data-top-scenario]');return s?s.options[s.selectedIndex].text:''})()")
    check("one scenario everywhere: Initiatives follows the scenario just opened on the capital plan (the picker at the top shows it)", capname and "board version" not in capname, capname)
    bv=await pg.evaluate("[...document.querySelector('[data-top-scenario]').options].find(o=>o.text.includes('board version')).value")
    await pg.select_option("select[data-top-scenario]", bv); await pg.wait_for_timeout(800)
    t=await pg.inner_text("#view")
    check("initiatives: status pipeline with counts", "All 17" in t.replace("\n"," ") and "Proposed 16" in t.replace("\n"," ") and "Approved 1" in t.replace("\n"," "), t[:200])
    check("initiatives: one-time and yearly costs from the board version", "One-time cost" in t and "$1.75M" in t and "District baseline, the board version" in t, t[:400])
    check("initiatives: which plans include each one", "Not in a plan yet" in t and "Board version · +1 other" in t)
    await pg.select_option("select[data-ini-sid]", _ph_sid); await pg.wait_for_timeout(400)
    t2=await pg.inner_text("#view")
    check("initiatives: costs from any scenario, including yearly costs from the capital plan", "$12k" in t2 and "Costs are from Addition phased" in t2, t2[-300:])
    await pg.select_option("select[data-ini-sid]", label="District baseline (board version)"); await pg.wait_for_timeout(400)
    tv=await pg.inner_text("#view")
    check("initiatives: one priority, in the funding line's words", "Must-have" in tv and "Strategic" in tv and "\tHigh\t" not in tv)
    await pg.click("button[data-action=iniStatus][data-v=approved]"); await pg.wait_for_timeout(300)
    t=await pg.inner_text("#view")
    check("initiatives: filter by status, and approved-but-not-planned is flagged", "FFA program" in t and "Middle school HVAC" not in t and "Approved, but not in the board version yet" in t)
    await pg.click("a[data-action=editInitiative]:has-text('FFA program')"); await pg.wait_for_timeout(300)
    m=await pg.inner_text("[data-modal]")
    check("decisions editor: same details, costs optional", "Not in a plan yet" in m and "Apply to scenario" in m and await pg.locator("[data-modal] [data-phase-row]").count()==0)
    locked=await pg.locator("[data-modal] select[data-ed-scenario] option[disabled]").count()
    check("decisions editor: locked scenarios listed but can't be picked", locked==1 and "(locked)" in await pg.locator("[data-modal] select[data-ed-scenario] option[disabled]").inner_text())
    await pg.select_option("[data-modal] select[data-ed-scenario]", _ph_sid); await pg.wait_for_timeout(300)
    row=pg.locator("[data-modal] [data-phase-row]").first
    await row.locator("input[name=label]").fill("Chapter start-up fees"); await row.locator("input[name=cost]").fill("5000")
    await row.locator("button[data-action=addFund]").click()
    await pg.click("[data-modal] button[data-action=addYearlyRow]")
    yr=pg.locator("[data-modal] [data-yearly-row]").last
    await yr.locator("input[name=y_amount]").fill("65,000"); await yr.locator("select[name=y_first]").select_option("2028")
    n0=len(calls); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(900)
    ops=[(c[0],c[1].split("?")[0].replace("/rest/v1/","")) for c in calls[n0:] if c[0] in ("POST","PATCH","DELETE")]
    phs=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/phase"]
    fus=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/phase_funding"]
    pat=[json.loads(c[2]) for c in calls[n0:] if c[0]=="PATCH" and "/rest/v1/initiative?" in c[1]]
    check("decisions editor: Apply to scenario writes the costs there", ("POST","scenario_initiative") in ops and phs and phs[0][0]["scenario_id"]==_ph_sid
          and phs[0][0]["label"]=="Chapter start-up fees" and len(fus[0])==2 and ("POST","recurring_cost") in ops, str(ops))
    check("decisions editor: details saved too, including owner", pat and "owner_name" in pat[0] and pat[0]["name"]=="FFA program", str(pat)[:200])
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/plan/initiatives"); await pg.wait_for_timeout(500)
    await pg.click("button[data-action=iniStatus][data-v='']"); await pg.wait_for_timeout(300)
    await pg.click("button[data-action=editInitiative][data-id='']"); await pg.wait_for_timeout(300)
    await pg.fill("[data-modal] input[name=name]","FFA program"); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(300)
    check("initiatives: a duplicate name is refused", "already an initiative called" in await pg.inner_text("[data-modal] [data-form-errors]"))
    await pg.fill("[data-modal] input[name=name]","Robotics club"); await pg.select_option("[data-modal] select[name=type]","program")
    await pg.select_option("[data-modal] select[name=status]","analysis"); await pg.fill("[data-modal] input[name=owner_name]","Ag teacher")
    dflt=await pg.evaluate("document.querySelector('[data-modal] select[data-ed-scenario]').value")
    phrows=await pg.locator("[data-modal] [data-phase-body] tr, [data-modal] [data-phase-body] [data-phase-row]").count()
    check("initiatives: a new initiative can take costs and phases straight away (a scenario that can change is chosen for it)", dflt not in ("", None) and phrows>=1, f"{dflt} {phrows}")
    await pg.select_option("[data-modal] select[data-ed-scenario]",""); await pg.wait_for_timeout(200)
    n0=len(calls); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(600)
    ip=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/initiative"]
    check("initiatives: add with type, status and owner, details only", ip and ip[0]["type"]=="program" and ip[0]["status"]=="analysis" and ip[0]["owner_name"]=="Ag teacher"
          and not any(c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/phase" for c in calls[n0:]), str(ip))
    await pg.click("a[data-action=editInitiative] >> nth=0"); await pg.wait_for_timeout(300)
    await pg.select_option("[data-modal] select[name=status]","approved"); n0=len(calls)
    await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(500)
    check("initiatives: edit details", any(c[0]=="PATCH" and "/rest/v1/initiative?" in c[1] and '"approved"' in (c[2] or "") for c in calls[n0:]))
    await pg.screenshot(path=SHOTS+"/initiatives.png",full_page=True)
    # Phase 2C: ranking and the funding line
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/plan/ranking"); await pg.wait_for_timeout(700)
    t=await pg.inner_text("#view")
    check("ranking: one list, must-have first, with the funding-line verdict", ("Everything fits" in t or "Everything planned for SAVE, PPEL, V-PPEL and grants fits" in t) and "#1" in t and t.find("Must-have")<t.find("Strategic")<t.find("Nice to have"), t[:300])
    check("ranking: campaign/bond and boosters called out separately, matching the gap", "$5.35M depends on a campaign or bond" in t and "$260k on boosters" in t and "Campaign or bond" in t, t[:500])
    check("ranking: priorities are the initiatives' own (no separate tier)", "Priority" in t and "from the old High/Med/Low" not in t)
    check("ranking: locked scenario can't be reordered", await pg.locator("button[data-action=rankMove]").count()==0 and "This scenario is locked" in t)
    # needed-by years, ordering by need, and suggestions (SAVE squeezed 10% so the money runs out)
    _turf=next(x for x in TABLES["initiative"] if x["name"]=="Stadium turf"); _cte=next(x for x in TABLES["initiative"] if x["name"].startswith("CTE welding"))
    _turf["need_by_fy"]=2030; _cte["need_by_fy"]=2028
    _set=TABLES["district_settings"][0]; _save0=_set["save_receipts"]; _set["save_receipts"]=round(float(_save0)*0.9,2)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/plan/initiatives"); await pg.wait_for_timeout(400)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/plan/ranking"); await pg.wait_for_timeout(900)
    t=await pg.inner_text("#view")
    check("ranking: needed-by years shown, late ones flagged", "needed by FY2030" in t and "needed by FY2030 (planned later)" in t and "planned after the year" in t, t[:800])
    sg=await pg.inner_text("[data-suggest]") if await pg.locator("[data-suggest]").count() else ""
    check("ranking: suggestions name a move and what it funds", "Move " in sg and "funds " in sg and await pg.locator("[data-suggest] button[data-action=rankSuggest]").count()>0, sg[:600])
    await pg.select_option("select[data-rank-by]","need"); await pg.wait_for_timeout(800)
    first=(await pg.locator("table.ranktable tbody tr:not(.fline) td:nth-child(3)").first.inner_text()).strip()
    check("ranking: ordered by when needed, a strategic item needed in FY2028 comes first", first.startswith("CTE welding"), first)
    await pg.select_option("select[data-rank-by]","rank"); await pg.wait_for_timeout(800)
    PROMPT["text"]="Squeezed, with a move"; n0=len(calls)
    await pg.locator("[data-suggest] button[data-action=rankSuggest]").first.click(); await pg.wait_for_timeout(900); PROMPT["text"]=None
    cp=[c for c in calls[n0:] if "/rpc/copy_scenario" in c[1]]; pt=[c for c in calls[n0:] if c[0]=="PATCH" and "/rest/v1/phase?" in c[1]]
    check("ranking: a suggestion becomes a new scenario with that initiative's unstarted phases moved", cp and json.loads(cp[0][2])["p_name"]=="Squeezed, with a move" and pt and all("scenario_id=eq.copied-scenario-id" in c[1] and "status=not.in.(done,underway)" in urllib.parse.unquote(c[1]) for c in pt), str(pt)[:400])
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/plan/initiatives"); await pg.wait_for_timeout(300)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/plan/ranking"); await pg.wait_for_timeout(800)
    await pg.select_option("select[data-rank-sid]", next(x["id"] for x in TABLES["scenario"] if x["is_board_version"])); await pg.wait_for_timeout(600)
    await pg.click("button[data-action=rankView][data-v=years]"); await pg.wait_for_timeout(900)
    t=await pg.inner_text("#view")
    check("year by year: everything fits if one initiative waits a year (the ranked list said the money runs out)", "Everything fits" in t and "Locker rooms renovation 1 year" in t and await pg.locator("svg.yrtimeline").count()==1 and "What changes" in t, t[:400])
    check("year by year: every mark explains itself on hover", await pg.locator("svg.yrtimeline rect title").count() >= 10)
    await pg.screenshot(path=SHOTS+"/year-by-year.png",full_page=True)
    PROMPT["text"]="Year by year"; n0=len(calls)
    await pg.click("button[data-action=rankSchedule]"); await pg.wait_for_timeout(900); PROMPT["text"]=None
    pt=[c for c in calls[n0:] if c[0]=="PATCH" and "/rest/v1/phase?" in c[1]]
    check("year by year: the timing becomes a new scenario", any("/rpc/copy_scenario" in c[1] for c in calls[n0:]) and len(pt)==1 and "fy=eq.2033" in pt[0][1] and json.loads(pt[0][2])=={"fy":2034}, str(pt))
    await pg.click("button[data-action=rankView][data-v=list]"); await pg.wait_for_timeout(500)
    lr=await pg.evaluate("[...document.querySelectorAll('select[data-goal-id] option')].find(o=>/below the line/.test(o.textContent)).value")
    await pg.select_option("select[data-goal-id]", lr); await pg.click("button[data-action=goalFind]"); await pg.wait_for_timeout(500)
    gr=await pg.inner_text("#goal-results")
    check("what would it take: options to fund a below-the-line initiative (a move, or a higher rank and who drops)", "Move " in gr and "Rank it #" in gr, gr[:500])
    PROMPT["text"]="Funding locker rooms"; n0=len(calls)
    await pg.locator("#goal-results button[data-action=goalApply]").first.click(); await pg.wait_for_timeout(900); PROMPT["text"]=None
    check("what would it take: an option becomes a new scenario", any("/rpc/copy_scenario" in c[1] for c in calls[n0:]) and any(c[0] in ("PATCH","POST") and ("/rest/v1/phase?" in c[1] or "scenario_initiative" in c[1]) for c in calls[n0:]))
    _set["save_receipts"]=_save0; _turf.pop("need_by_fy",None); _cte.pop("need_by_fy",None)
    await pg.screenshot(path=SHOTS+"/ranking.png",full_page=True)
    await pg.select_option("select[data-rank-sid]", _ph_sid); await pg.wait_for_timeout(600)
    check("ranking: unlocked scenario has move buttons", await pg.locator("button[data-action=rankMove]").count()>0)
    first_down=pg.locator("button[data-action=rankMove][data-d='1']").first
    moved_id=await first_down.get_attribute("data-id")
    n0=len(calls); await first_down.click(); await pg.wait_for_timeout(600)
    up=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and "/rest/v1/scenario_initiative?on_conflict=" in c[1]]
    check("ranking: moving saves the new order for the scenario", up and up[0][1]["initiative_id"]==moved_id and up[0][0]["rank"]==1 and up[0][1]["rank"]==2 and all(r["scenario_id"]==_ph_sid for r in up[0]), str(up[0][:2]) if up else "none")
    rows_before=await pg.locator("input[data-rank-to]").count()
    last=pg.locator("input[data-rank-to]").nth(rows_before-1); last_id=await last.get_attribute("data-rank-to")
    n0=len(calls); await last.fill("1"); await last.dispatch_event("change"); await pg.wait_for_timeout(700)
    up=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and "/rest/v1/scenario_initiative?on_conflict=" in c[1]]
    check("ranking: force rank across priorities: type #1 on the last item and it moves to the top", up and up[0][0]["initiative_id"]==last_id and up[0][0]["rank"]==1 and len(up[0])==rows_before, str(up[0][:2]) if up else "none")
    btn_dn=[await b.is_disabled() for b in await pg.locator("button[data-action=rankMove][data-d='1']").all()]
    check("ranking: ▲▼ aren't limited to one priority (only the last row can't move down)", btn_dn and not any(btn_dn[:-1]) and btn_dn[-1], str(btn_dn))
    n0=len(calls); await pg.select_option("select[data-rank-tier] >> nth=0", "strategic"); await pg.wait_for_timeout(600)
    check("ranking: changing a tier saves it on the initiative", any(c[0]=="PATCH" and "/rest/v1/initiative?" in c[1] and '"strategic"' in (c[2] or "") for c in calls[n0:]))
    check("ranking: nothing to defer when everything fits", "Nothing needs deferring" in await pg.inner_text("#view") and await pg.locator("button[data-action=rankScenario]").count()==0)
    # starting numbers: checks
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/setup"); await pg.wait_for_timeout(700)
    ck=pg.locator("[data-setup-checks]")
    t=await ck.inner_text()
    check("checks: shown on arrival, as tips, with nothing wrong yet", await ck.is_visible() and "Worth a second look" in t and "don’t match" not in t and "5% debt limit" in t, t[:300])
    await pg.fill("input[name=ppel_receipts]","100,000"); await pg.wait_for_timeout(150)
    check("checks: PPEL receipts that don't match rate × valuation", "PPEL receipts ($100,000) don’t match" in await ck.inner_text())
    await pg.fill("input[name=save_receipts_fy]",""); await pg.wait_for_timeout(150)
    check("checks: a blank SAVE receipts year", "Say which fiscal year the SAVE receipts are for" in await ck.inner_text())
    await pg.fill("input[name=enrollment]","300"); await pg.wait_for_timeout(150)
    check("checks: SAVE per student far from the statewide amount", "per student; SAVE is shared statewide at about $1,358" in await ck.inner_text())
    while await pg.locator("button[data-action=removeDebtRow]").count(): await pg.locator("button[data-action=removeDebtRow]").first.click()
    await pg.wait_for_timeout(150)
    await ck.screenshot(path=SHOTS+"/checks.png")
    check("checks: no existing debt is questioned", "No existing debt entered" in await ck.inner_text())
    check("checks: where-to-find hints on the fields", "certified budget (Iowa Department of Management)" in await pg.inner_text("#view"))
    # Phase 2E: assumption sets
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/resources/assumptions"); await pg.wait_for_timeout(600)
    check("assumption sets: none yet, with starters offered", "No assumption sets yet" in await pg.inner_text("#view"))
    n0=len(calls); await pg.click("button[data-action=starterSets]"); await pg.wait_for_timeout(600)
    sp=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and "/rest/v1/assumption_set" in c[1]]
    check("assumption sets: Base, Conservative and Growth made from the starting numbers", sp and [x["name"] for x in sp[0]]==["Base","Conservative","Growth"]
          and sp[0][0]["is_default"] and sp[0][1]["construction_inflation"]>sp[0][0]["construction_inflation"] and sp[0][2]["ppel_growth"]>sp[0][0]["ppel_growth"], str(sp)[:300])
    TABLES["assumption_set"]=[{"id":"as-c","district_id":"d1","name":"Conservative","is_default":False,"construction_inflation":0.05,"save_trend":-0.01,"ppel_growth":0.02,"grant_yield":0.5,"settlement_pct":0.04,"notes":"Tougher"}]
    for sc in TABLES["scenario"]:
      if sc["id"]==_ph_sid: sc["assumption_set_id"]="as-c"; sc["lever_inflation"]=None
    await pg.reload(); await pg.wait_for_timeout(900)
    t=await pg.inner_text("#view")
    check("assumption sets: listed with values and the scenarios using them", "Conservative" in t and "5%" in t and "Addition phased" in t)
    await pg.click("a[data-action=editSet]"); await pg.wait_for_timeout(300)
    await pg.fill("[data-modal] input[name=construction_inflation]","6"); n0=len(calls); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(500)
    check("assumption sets: edit saves percentages as fractions", any(c[0]=="PATCH" and "/rest/v1/assumption_set?" in c[1] and json.loads(c[2])["construction_inflation"]==0.06 for c in calls[n0:]))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/money/capital"); await pg.wait_for_timeout(500)
    await pg.select_option("select[data-cap-scenario]",_ph_sid); await pg.wait_for_timeout(500)
    check("capital plan: the scenario's set drives its levers", "Conservative" in await pg.inner_text("#cap-levers") and (await pg.inner_text("[data-lever-val=infl]")).startswith("5"), await pg.inner_text("[data-lever-val=infl]"))
    n0=len(calls); await pg.select_option("select[data-cap-set]",""); await pg.wait_for_timeout(600)
    pp=[json.loads(c[2]) for c in calls[n0:] if c[0]=="PATCH" and "/rest/v1/scenario?" in c[1]]
    check("capital plan: choosing assumptions clears saved levers so the choice applies", pp and pp[0]["assumption_set_id"] is None and pp[0]["lever_inflation"] is None and pp[0]["lever_ppel_growth"] is None, str(pp))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/plan/scenarios"); await pg.wait_for_timeout(600)
    check("compare: which assumptions each scenario uses", "Assumptions" in await pg.inner_text("#cmp-table") and "Conservative" in await pg.inner_text("#cmp-table"))
    TABLES["assumption_set"]=[]
    for sc in TABLES["scenario"]:
      if sc["id"]==_ph_sid: sc["assumption_set_id"]=None
    # Phase 3B: Progress → Initiatives
    TABLES["gl_account"]=[]; TABLES["gl_amount"]=[]
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/track/initiatives"); await pg.wait_for_timeout(700)
    t=await pg.inner_text("#view")
    check("progress: against the board version, before any ledger", "No month-end general ledger (GL) export for FY2027 yet" in t and "District baseline" in t and "Middle school HVAC replacement" in t and "$900,000" in t, t[:400])
    GL=json.loads(subprocess.check_output(["node","-e","""
      const G=require('./gl.js');const P=G.parse(G.sampleExport());
      process.stdout.write(JSON.stringify(P.lines.map((l,k)=>Object.assign({id:'gla'+k,district_id:'d1',code:l.code,description:l.description,project_code:l.parts.project||null,function_code:l.parts.function||null,initiative_id:null,needs_review:false,_ytd:l.ytd,_enc:l.encumbered,_bud:l.budget},G.suggest(l.parts)))))"""],cwd=os.path.dirname(os.path.abspath(__file__))))
    TABLES["gl_account"]=[{k:v for k,v in a.items() if not k.startswith("_")} for a in GL]
    TABLES["gl_amount"]=[{"district_id":"d1","batch_id":"glb1","account_id":a["id"],"ytd_amount":a["_ytd"],"encumbered":a["_enc"],"budget_amount":a["_bud"]} for a in GL]
    TABLES["import_batch"]=TABLES["import_batch"]+[{"id":"glb1","district_id":"d1","kind":"gl_monthly","status":"applied","period_end":"2026-09-30","fiscal_year":2027,"file_name":"gl.csv","row_count":21,"uploaded_at":"2026-10-01T10:00:00Z"}]
    await pg.reload(); await pg.wait_for_timeout(900)
    t=await pg.inner_text("#view")
    roof_init=next(i["id"] for i in TABLES["initiative"] if i["name"].startswith("High school roof"))
    roof_acct=next(a["id"] for a in GL if a["code"]=="33-0000-4700-000-1001-450")
    check("progress: link card suggests initiatives for capital spending", "Link spending to initiatives" in t and await pg.input_value("select[data-pi-link='%s']"%roof_acct)==roof_init and "through Sep 30, 2026" in t, t[:300])
    check("progress: debt payments and General Fund aren't offered for linking", await pg.locator("select[data-pi-link='%s']"%next(a["id"] for a in GL if a["code"]=="36-0000-5000-000-0000-831")).count()==0
          and await pg.locator("select[data-pi-link='%s']"%next(a["id"] for a in GL if a["code"].startswith("10-0109"))).count()==0)
    n0=len(calls); await pg.click("button[data-action=piLinkAll]"); await pg.wait_for_timeout(900)
    lk=[(c[1].split("id=eq.")[1], json.loads(c[2])) for c in calls[n0:] if c[0]=="PATCH" and "/rest/v1/gl_account?" in c[1]]
    check("progress: Link all suggested links each, as initiative spending", lk and any(i==roof_acct and b=={"initiative_id":roof_init,"maps_to":"initiative"} for i,b in lk), str(lk)[:300])
    for a in TABLES["gl_account"]:
      if a["id"]==roof_acct: a["initiative_id"]=roof_init; a["maps_to"]="initiative"
    await pg.reload(); await pg.wait_for_timeout(900)
    t=await pg.inner_text("table.pitable")
    check("progress: spending and encumbrances from the ledger, flagged when nothing was planned this year", "$262,750" in t and "$147,250" in t and "Spending, but nothing planned this year" in t, t[:500])
    await pg.click("a[data-action=piOpen][data-id='%s']"%roof_init); await pg.wait_for_timeout(500)
    await pg.locator("#view").screenshot(path=SHOTS+"/progress.png")
    first=pg.locator("tr[data-pi-phase]").first
    await first.locator("select[name=status]").select_option("done"); await first.locator("input[name=actual]").fill("")
    await first.locator("button[data-action=piSavePhase]").click(); await pg.wait_for_timeout(300)
    check("progress: done needs an actual cost", "Enter the actual cost when a phase is done" in (await pg.inner_text("#toasts")).replace("’","'").replace("'","’"))
    TABLES_ph=[p for p in TABLES["phase"] if p["initiative_id"]==roof_init and p["scenario_id"]!=_ph_sid]
    await first.locator("input[name=done]").fill("2026-09-25"); await first.locator("input[name=actual]").fill("401,500")
    n0=len(calls); await first.locator("button[data-action=piSavePhase]").click(); await pg.wait_for_timeout(600)
    sp=[json.loads(c[2]) for c in calls[n0:] if "rpc/set_phase_progress" in c[1]]
    check("progress: recorded through the progress-only function, even on the locked plan", sp and sp[0]["p_status"]=="done" and sp[0]["p_actual"]==401500 and sp[0]["p_done"]=="2026-09-25" and sp[0]["p_phase"]==sorted(TABLES_ph,key=lambda p:(p["fy"],p["seq"]))[0]["id"], str(sp))
    # Phase 3C: budget vs. actual (ledger still loaded from the progress tests); 3D adds an August import to compare
    TABLES["import_batch"]=TABLES["import_batch"]+[{"id":"glb0","district_id":"d1","kind":"gl_monthly","status":"applied","period_end":"2026-08-31","fiscal_year":2027,"file_name":"gl-aug.csv","row_count":21,"uploaded_at":"2026-09-10T10:00:00Z"}]
    TABLES["gl_amount"]=TABLES["gl_amount"]+[dict(g, batch_id="glb0", ytd_amount=(g["ytd_amount"] or 0)/2, encumbered=0) for g in TABLES["gl_amount"] if g["batch_id"]=="glb1"]
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/track/actuals"); await pg.wait_for_timeout(800)
    t=await pg.inner_text("#view")
    check("budget: every fund, General Fund first, from the ledger through the month end", "Every fund" in t and t.find("General Fund")<t.find("SAVE")<t.find("PPEL") and "Sep 30, 2026" in t and "25% of the fiscal year gone" in t, t[:300])
    check("budget: by default, the year ends at budget unless already exceeded (no false alarms)", "$12,300,000" in t and "$2,909,551" in t and "over by" not in t and "on budget" in t, t[:600])
    await pg.select_option("select[data-ba=method]","pace"); await pg.wait_for_timeout(600); t=await pg.inner_text("#view")
    check("budget: rest of the year at the budget's pace", "$12,134,551" in t and "short by $165,449" in t)
    check("budget: spending by function, with what's still available", "Instruction" in t and "$5,275,000" in t and "$880,634" in t and "Operation and maintenance of plant" in t and "$4,394,366" in t, t[t.find("Spending by function"):t.find("Spending by function")+400])
    await pg.locator("#view").screenshot(path=SHOTS+"/budget.png")
    await pg.select_option("select[data-ba=method]","straight"); await pg.wait_for_timeout(600)
    check("budget: straight-line forecast, with its caution", "$11,638,206" in await pg.inner_text("#view") and "can mislead this method" in await pg.inner_text("#view"))
    await pg.click("a[data-action=baFund][data-k=save]"); await pg.wait_for_timeout(600)
    t=await pg.inner_text("#view")
    check("budget: any fund in detail", "Facilities acquisition and construction" in t and "$1,180,640" in t, t[t.find("Spending by function")-300:t.find("Spending by function")+300])
    t=await pg.inner_text("#view")
    check("compare: two months side by side, with the change", "Compare two months" in t and "Aug 31, 2026" in t and "+$1,454,776" in t and "different fiscal years" not in t, t[t.find("Compare two months"):t.find("Compare two months")+500])
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/uploads"); await pg.wait_for_timeout(500)
    check("uploads: balances-only is now a labelled fallback", "Fund balances only (if you can’t export the ledger)" in await pg.inner_text("select[data-upload-kind]"))
    TABLES["fund_balance"]=TABLES["fund_balance"]+[{"district_id":"d1","fund":"save","as_of":"2026-09-30","amount":1309086,"source":"gl_import"},{"district_id":"d1","fund":"ppel","as_of":"2026-09-30","amount":788200,"source":"gl_import"}]
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/resources/funds"); await pg.wait_for_timeout(900); await pg.evaluate("document.querySelectorAll('details.fold').forEach(d=>d.open=true)")
    t=await pg.inner_text("#view")
    hist=t[t.find("Balances month by month"):]
    check("all funds: balances month by month, newest first, with where they came from", "Balances month by month" in t and hist.find("Sep 30, 2026")<hist.find("Jul 1, 2026") and "$1,309,086" in hist and "Monthly GL" in hist, t[t.find("Balances month by month"):t.find("Balances month by month")+300])
    TABLES["fund_balance"]=[x for x in TABLES["fund_balance"] if x.get("as_of")!="2026-09-30"]
    saved_ends={b0["id"]:b0.get("period_end") for b0 in TABLES["import_batch"]}
    for b0 in TABLES["import_batch"]:
      if b0.get("kind")=="gl_monthly": b0["period_end"]="2020-01-31" if b0["id"]=="glb1" else "2019-12-31"
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pg.wait_for_timeout(500)
    check("reminder: business staff see when the ledger is getting old", "The ledger is through Jan 31, 2020" in await pg.inner_text("#view") and "Feb 29, 2020" in await pg.inner_text("#view"))
    for b0 in TABLES["import_batch"]: b0["period_end"]=saved_ends.get(b0["id"])
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/track/actuals"); await pg.wait_for_timeout(600)
    for g in TABLES["gl_amount"]: g["budget_amount"]=None
    await pg.reload(); await pg.wait_for_timeout(800)
    check("budget: an export with no budget column says so", "has no budget column" in await pg.inner_text("#view"))
    TABLES["gl_account"]=[]; TABLES["gl_amount"]=[]; TABLES["import_batch"]=[b for b in TABLES["import_batch"] if b["id"] not in ("glb0","glb1")]
    kept_gl=[b for b in TABLES["import_batch"] if b.get("kind")=="gl_monthly"]; TABLES["import_batch"]=[b for b in TABLES["import_batch"] if b.get("kind")!="gl_monthly"]
    await pg.reload(); await pg.wait_for_timeout(600)
    check("budget: before any ledger, says what's needed", "No monthly ledger yet" in await pg.inner_text("#view"))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pg.wait_for_timeout(500)
    await pg.click("button[data-action=openHelp]"); await pg.wait_for_timeout(300)
    await pg.locator("[data-modal] a", has_text="The full guide").click(); await pg.wait_for_timeout(600)
    check("a link inside a pop-up leaves the page scrollable", await pg.evaluate("!document.documentElement.classList.contains('modal-open') && !document.querySelector('[data-modal]')"))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pg.wait_for_timeout(500)
    check("reminder: before any ledger, business staff are invited to start", "Bring in the business office’s month-end general ledger (GL) export each month" in await pg.inner_text("#view"))
    TABLES["import_batch"]=TABLES["import_batch"]+kept_gl
    # no unexpected scroll bars anywhere: a box may scroll sideways, never show up/down arrows for a few pixels
    stray_js="""[...document.querySelectorAll('#app *')].filter(el=>{const cs=getComputedStyle(el);
      if(!/(auto|scroll)/.test(cs.overflowY)||el.matches('.modal-back'))return false; return el.scrollHeight>el.clientHeight;}).map(el=>el.className||el.tagName).slice(0,5)"""
    stray={}
    for route in ["home/today","plan/initiatives","plan/ranking","plan/scenarios","money/summary","money/capital","resources/funds","resources/assumptions",
                  "track/initiatives","track/actuals","settings/uploads","reports/exports","share/community","settings/setup","settings/people","settings/activity","help/guide"]:
      await pg.goto("http://localhost:8765/#/d/ironwood-valley/"+route); await pg.wait_for_timeout(450)
      found=await pg.evaluate(stray_js)
      if found: stray[route]=found
    check("no unexpected scroll bars on any page", not stray, str(stray))
    # every module is loaded after the modules it depends on (in the browser, order matters)
    import re as _re
    _root=os.path.dirname(os.path.abspath(__file__)); idx=open(os.path.join(_root,"index.html")).read()
    order=_re.findall(r'<script src="([a-z_]+\.js)"></script>', idx)
    provides={}
    for f in order:
      m=_re.search(r"root\.(HG[A-Za-z]+) = api", open(os.path.join(_root,f)).read())
      if m: provides[m.group(1)]=f
    bad=[]
    for i,f in enumerate(order):
      src=open(os.path.join(_root,f)).read()
      for dep in set(_re.findall(r"root\.(HG[A-Za-z]+) \|\||need\('(HG[A-Za-z]+)'", src)):
        dep=[x for x in dep if x][0] if isinstance(dep,tuple) else dep
        if dep in provides and order.index(provides[dep])>i: bad.append(f"{f} needs {dep} ({provides[dep]}) loaded first")
    check("modules load in dependency order", not bad and len(provides)>=10, str(bad) or str(provides))
    # Phase 4A: monthly board report
    TABLES["report_snapshot"]=[]
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/share/board"); await pg.wait_for_timeout(600)
    check("board reports: none yet, create offered", "No board reports yet" in await pg.inner_text("#view") and await pg.locator("button[data-action=rpCreate]").count()==1)
    await pg.fill("input[data-rp-date]","2026-09-30")
    n0=len(calls); await pg.click("button[data-action=rpCreate]"); await pg.wait_for_timeout(900)
    rs=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/report_snapshot"]
    ok_rs = rs and rs[0]["kind"]=="board_monthly" and rs[0]["period_end"]=="2026-09-30" and rs[0]["title"]=="Board report, September 2026" and abs(rs[0]["payload"]["plan"]["gap"]-5350000)<1
    check("board reports: made from the district's data and saved as it was", ok_rs, str(rs)[:300])
    snap=dict(rs[0], district_id="d1", created_at="2026-10-01T15:00:00Z") if rs else None
    if snap:
      prev=json.loads(json.dumps(snap)); prev.update(id="snap-aug", period_end="2026-08-31", created_at="2026-09-05T15:00:00Z")
      prev["payload"]["periodEnd"]="2026-08-31"; prev["payload"]["plan"]["gap"]=5600000
      for k in prev["payload"]["balances"]: prev["payload"]["balances"][k]["amount"]+=50000
      TABLES["report_snapshot"]=[snap]
      await pg.reload(); await pg.wait_for_timeout(600)
      await pg.click("a[data-action=rpOpen]"); await pg.wait_for_timeout(500)
      t=await pg.inner_text("#view")
      check("board report: the first one says there's nothing to compare yet", "Board report · Sep 30, 2026" in t and "This is the first board report" in t and "District baseline" in t, t[:300])
      check("board report: the capital plan, decisions ahead, and no hand editing", "left to close" in t and "Decisions ahead" in t and "with no hand editing" in t)
      TABLES["report_snapshot"]=[snap, prev]
      await pg.reload(); await pg.wait_for_timeout(600)
      await pg.click("a[data-action=rpOpen][data-id='%s']"%snap["id"]); await pg.wait_for_timeout(500)
      t=await pg.inner_text("#view")
      check("board report: what changed since the previous report", "What changed since Aug 31, 2026" in t and "The gap to close shrank by $250k, to $5.35M." in t and "SAVE balance down $50k" in t and "▼ $250k since Aug 31, 2026" in t, t[t.find("What changed"):t.find("What changed")+400])
      await pg.screenshot(path=SHOTS+"/report.png", full_page=True)
      await pg.emulate_media(media="print")
      check("board report: printing hides the app's menus and buttons", not await pg.locator("nav.rail").is_visible() and not await pg.locator("button[data-action=rpPrint]").is_visible() and await pg.locator(".report .rhead").is_visible())
      await pg.emulate_media(media="screen")
    # Phase 4B: decision packets
    TABLES["report_snapshot"]=[]
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pg.wait_for_timeout(200)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/share/board"); await pg.wait_for_timeout(700)
    addn=next(i for i in TABLES["initiative"] if i["name"].startswith("Elementary classroom addition"))
    await pg.select_option("select[data-pk-init]", addn["id"])
    n0=len(calls); await pg.click("button[data-action=pkCreate]"); await pg.wait_for_timeout(900)
    pk=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/report_snapshot"]
    check("decision packet: made and saved as it was", pk and pk[0]["kind"]=="decision_packet" and pk[0]["payload"]["initiative"]["name"].startswith("Elementary classroom addition") and pk[0]["payload"]["scenario"]["board"], str(pk)[:200])
    if pk:
      TABLES["report_snapshot"]=[dict(pk[0], district_id="d1", created_at="2026-10-01T16:00:00Z")]
      await pg.reload(); await pg.wait_for_timeout(700)
      await pg.click("a[data-action=rpOpen][data-id='%s']"%pk[0]["id"]); await pg.wait_for_timeout(600)
      t=await pg.inner_text("#view")
      check("decision packet: costs by fund and year, effect on the gap, funding line, other scenarios", "Costs by fund and year" in t and "Campaign or bond" in t and "FY2031" in t and "Gap with it / without it" in t and "paid by a campaign or bond" in t and "In other scenarios" in t, t[:600])
      await pg.screenshot(path=SHOTS+"/packet.png", full_page=True)
    TABLES["report_snapshot"]=[]
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/plan/initiatives"); await pg.wait_for_timeout(600)
    await pg.click("a[data-action=editInitiative]:has-text('Middle school HVAC')"); await pg.wait_for_timeout(400)
    n0=len(calls); await pg.click("[data-modal] button[data-action=pkFromEditor]"); await pg.wait_for_timeout(1000)
    pk2=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/report_snapshot"]
    check("decision packet: one click from the initiative on Decisions", pk2 and pk2[0]["payload"]["initiative"]["name"]=="Middle school HVAC replacement" and "/share/board" in pg.url)
    TABLES["report_snapshot"]=[]
    # Phase 5A: priorities, outcomes, measures
    for t in ["priority","outcome","measure","measure_value"]: TABLES[t]=[]
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/plan/priorities"); await pg.wait_for_timeout(600)
    check("priorities: none yet, with an invitation to start", "No priorities yet" in await pg.inner_text("#view"))
    await pg.click("button[data-action=editPriority]"); await pg.wait_for_timeout(300)
    await pg.fill("[data-modal] input[name=name]","Safe, modern places to learn"); await pg.fill("[data-modal] textarea[name=outcomes]","Safe and secure buildings\nFacilities in good repair")
    n0=len(calls); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(700)
    pr=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/priority"]
    oc=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/outcome"]
    check("priorities: a priority with its outcomes", pr and pr[0]["name"]=="Safe, modern places to learn" and oc and [o["name"] for o in oc[0]]==["Safe and secure buildings","Facilities in good repair"] and all(o["priority_id"]==pr[0]["id"] for o in oc[0]), str(oc)[:200])
    DIRR=json.loads(subprocess.check_output(["node","-e","""
      const C=require('./capital.js'),D=require('./demo_data.js');let i=0;const R=C.demoRows(D['bridger-hollow'],'d1',()=>'dir-'+(++i));
      process.stdout.write(JSON.stringify({priority:R.priority,outcome:R.outcome,measure:R.measure,measure_value:R.measure_value.map((v,k)=>Object.assign({id:'mv'+k},v))}))"""],cwd=os.path.dirname(os.path.abspath(__file__))))
    for t in DIRR: TABLES[t]=DIRR[t]
    TABLES["priority"]=[dict(p) for p in DIRR["priority"]]
    for i0 in TABLES["initiative"]:
      if i0["name"].startswith("Secure entrances"): i0["priority_id"]=DIRR["priority"][1]["id"]
    await pg.reload(); await pg.wait_for_timeout(800)
    t=await pg.inner_text("#view")
    check("priorities: each with outcomes, measures and their status, and initiatives serving it", "Every graduate ready for what’s next" in t and "Career and technical pathways" in t and "On track" in t and "Off track" in t and "Secure entrances — three buildings" in t, t[:500])
    # Reports → Plans: one plan, three views
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/share/plans"); await pg.wait_for_timeout(900)
    t=await pg.inner_text("#view")
    check("plans: the full plan, with goals, initiatives by status, schedule and capital funds", "District improvement plan" in t and "approved initiative" in t and "Goals" in t and "Every graduate ready for what’s next" in t and "Capital funds, FY" in t and "Totals count approved work only" in t, t[:900])
    check("plans: approved only by default (no proposals listed)", "Proposed, not yet approved by the board" not in t)
    await pg.select_option("select[data-plan-items]","all"); await pg.wait_for_timeout(600)
    t=await pg.inner_text("#view")
    npend=len([i for i in TABLES["initiative"] if (i.get("status") or "proposed") in ("idea","proposed","analysis")])
    check("plans: all items shows proposals, labelled and not counted, and the schedule by year", (("Proposed, not yet approved by the board" in t) and "not counted" in t and "Schedule by year" in t and "(proposed)" in t) if npend else True, f"{npend} pending; "+t[:400])
    await pg.click("button[data-action=planView][data-v=capital]"); await pg.wait_for_timeout(600)
    t=await pg.inner_text("#view")
    check("plans: the capital improvement plan has capital work and the funding picture, no goals", "Capital improvement plan" in t and "Capital funds, FY" in t and "Coming in" in t and "\nGoals\n" not in t, t[:600])
    async with pg.expect_download() as dl: await pg.click("button[data-action=planXlsx]")
    f=await dl.value
    import openpyxl as _ox2, io as _io2
    wb=_ox2.load_workbook(_io2.BytesIO(open(await f.path(),"rb").read())).active; hdr=[c.value for c in wb[1]]
    check("plans: Excel in the district's columns", f.suggested_filename.startswith("ironwood-valley-capital-improvement-plan-") and hdr[:3]==["Status","Initiative","Anticipated completion"] and "Funding source" in hdr and wb.max_row>1, str(hdr))
    await pg.click("button[data-action=planView][data-v=csip]"); await pg.wait_for_timeout(600)
    check("plans: academic goals view explains how to tag CSIP goals", "No priorities are tagged as CSIP goals yet" in await pg.inner_text("#view"))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/plan/priorities"); await pg.wait_for_timeout(700)
    pid=TABLES["priority"][0]["id"]
    await pg.click(f"button[data-action=editPriority][data-id='{pid}']"); await pg.wait_for_timeout(300)
    await pg.check("[data-modal] input[name=csip_goal]")
    n0=len(calls); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(700)
    up=[json.loads(c[2]) for c in calls[n0:] if c[0]=="PATCH" and c[1].split("?")[0]=="/rest/v1/priority"]
    check("priorities: tagged as a state CSIP goal", up and up[0].get("csip_goal") is True, str(up)[:200])
    for p0 in TABLES["priority"]:
      if p0["id"]==pid: p0["csip_goal"]=True
    await pg.reload(); await pg.wait_for_timeout(700)
    check("priorities: the CSIP badge shows", "CSIP goal" in await pg.inner_text("#view"))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/share/plans"); await pg.wait_for_timeout(800)
    await pg.click("button[data-action=planView][data-v=csip]"); await pg.wait_for_timeout(600)
    t=await pg.inner_text("#view")
    check("plans: the academic goals view shows the CSIP goal with its measures", "Academic goals (CSIP)" in t and TABLES["priority"][0]["name"] in t and "filed with the Department of Education in CASA" in t, t[:700])
    TABLES.setdefault("report_snapshot",[])
    n0=len(calls); await pg.click("button[data-action=planSave]"); await pg.wait_for_timeout(700)
    sv=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/report_snapshot"]
    check("plans: a version is saved exactly as shown", sv and sv[0]["kind"]=="improvement_plan" and sv[0]["payload"]["plan"]["view"]=="csip" and "Academic goals (CSIP)" in sv[0]["title"], str(sv)[:300])
    await pg.wait_for_timeout(300)
    if "Saved versions" in await pg.inner_text("#view"):
      await pg.click("a[data-action=planOpen] >> nth=0"); await pg.wait_for_timeout(500)
      check("plans: a saved version opens as it was", "Back to the live plan" in await pg.inner_text("#view") and "saved" in await pg.inner_text("#view"))
      await pg.click("button[data-action=planBack]"); await pg.wait_for_timeout(500)
    await pg.click("button[data-action=planView][data-v=full]"); await pg.wait_for_timeout(500)
    await pg.fill("form[data-form=savePlanName] input[name=plan_name]","Comprehensive improvement plan")
    n0=len(calls); await pg.click("form[data-form=savePlanName] button[type=submit]"); await pg.wait_for_timeout(700)
    pn=[json.loads(c[2]) for c in calls[n0:] if c[0]=="PATCH" and c[1].split("?")[0]=="/rest/v1/district_settings"]
    check("plans: the district names its full plan", pn and pn[0].get("plan_name")=="Comprehensive improvement plan", str(pn)[:200])
    for st0 in TABLES.get("district_settings",[]): st0["plan_name"]="Comprehensive improvement plan"
    await pg.reload(); await pg.wait_for_timeout(800)
    check("plans: the full plan carries the district's name", "Comprehensive improvement plan" in await pg.inner_text("#view"))
    await pg.locator("#view").screenshot(path=SHOTS+"/plans.png")
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/track/measures"); await pg.wait_for_timeout(700)
    t=await pg.inner_text("#view")
    check("measures: start → target, latest, status, progress, owner", "7 measures:" in t and "Four-year graduation rate" in t and "88% 2024-25 → 94% 2028-29" in t.replace("\t"," ").replace("  "," ") and "lower is better" in t and "Principal, high school" in t, t[:600])
    await pg.screenshot(path=SHOTS+"/measures.png", full_page=True)
    await pg.click("a[data-action=editMeasure] >> nth=0"); await pg.wait_for_timeout(300)
    await pg.fill("[data-modal] input[name=target_period]","someday"); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(300)
    check("measures: a period HighGround can't read is explained", "isn’t a period HighGround can read" in await pg.inner_text("[data-modal] [data-form-errors]"))
    await pg.fill("[data-modal] input[name=target_period]","2029-30"); n0=len(calls); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(600)
    check("measures: edit saves", any(c[0]=="PATCH" and "/rest/v1/measure?" in c[1] and '"2029-30"' in (c[2] or "") for c in calls[n0:]))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/progress/measures"); await pg.wait_for_timeout(800)
    check("moved: the old Progress → Measures link lands on Direction → Measures", "/track/measures" in pg.url)
    grad=next(m["id"] for m in DIRR["measure"] if m["name"]=="Four-year graduation rate")
    check("measures: one place, led by how many are on track", "measures are on track or met" in await pg.inner_text("p.lead") and await pg.locator("button[data-action=recordResult]").count()>=6)
    await pg.click("button[data-action=recordResult][data-id='%s']"%grad); await pg.wait_for_timeout(300)
    check("record results: next period suggested by cadence", await pg.input_value("[data-modal] input[name=period]") in ("2026-27","2025-26","2027-28"))
    await pg.fill("[data-modal] input[name=period]","next year"); await pg.fill("[data-modal] input[name=value]","90.7"); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(300)
    check("record results: an unreadable period is explained", "isn’t a period HighGround can read" in await pg.inner_text("[data-modal] [data-form-errors]"))
    await pg.fill("[data-modal] input[name=period]","2026-27")
    n0=len(calls); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(600)
    mv=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].startswith("/rest/v1/measure_value")]
    check("record results: saved for the period (replacing any earlier one)", mv and mv[0][0]["period"]=="2026-27" and mv[0][0]["period_end"]=="2027-06-30" and mv[0][0]["value"]==90.7
          and "on_conflict=measure_id%2Cperiod" in next(c[1] for c in calls[n0:] if c[1].startswith("/rest/v1/measure_value")), str(mv))
    # Phase 5B: an automatic measure (phases finished in their planned year) fills itself in
    TABLES["measure"]=TABLES["measure"]+[{"id":"m-auto","district_id":"d1","priority_id":DIRR["priority"][2]["id"],"outcome_id":None,"name":"Capital phases finished in their planned year","unit":"%","better":"up",
      "baseline_value":80,"baseline_period":"2024-25","target_value":100,"target_period":"2027-28","owner_name":"Facilities director","cadence":"annual","source":"progress","is_public":True,"auto_metric":"phases_on_schedule"}]
    bph=[p0 for p0 in TABLES["phase"] if p0["scenario_id"]!=_ph_sid][:2]
    saved_ph=[dict(p0) for p0 in bph]
    bph[0].update(status="done", done_date="2026-09-20", actual_cost=bph[0]["cost"]); bph[1].update(status="done", done_date="2027-09-20", actual_cost=bph[1]["cost"])
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pg.wait_for_timeout(200)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/track/measures"); await pg.wait_for_timeout(700)
    ar=await pg.locator("tr", has=pg.locator("a[data-action=editMeasure][data-id='m-auto']")).inner_text()
    check("automatic measure: updates itself from finished phases, no typing", "Updates itself" in ar and await pg.locator("button[data-action=recordResult][data-id='m-auto']").count()==0 and ("FY2027" in ar or "FY2028" in ar), ar)
    for p0,sv0 in zip(bph,saved_ph): p0.clear(); p0.update(sv0)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/track/measures"); await pg.wait_for_timeout(600)
    await pg.click("button[data-action=editMeasure][data-id='']"); await pg.wait_for_timeout(300)
    await pg.select_option("[data-modal] select[data-me-auto]","save_balance"); await pg.wait_for_timeout(150)
    check("automatic measure: choosing one fills in its name, unit and cadence", await pg.input_value("[data-modal] input[name=name]")=="SAVE balance" and await pg.input_value("[data-modal] input[name=unit]")=="$" and await pg.input_value("[data-modal] select[name=cadence]")=="monthly")
    n0=len(calls); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(500)
    am=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/measure"]
    check("automatic measure: saved with its source", am and am[0]["auto_metric"]=="save_balance" and am[0]["source"]=="import", str(am))
    # Phase 5C: community survey, and the uploads for the plan
    SV=json.loads(subprocess.check_output(["node","-e","""
      const C=require('./capital.js'),D=require('./demo_data.js');let i=0;const R=C.demoRows(D['bridger-hollow'],'d1',()=>'dir-'+(++i));
      process.stdout.write(JSON.stringify({survey:R.survey,survey_result:R.survey_result.map((r,k)=>Object.assign({id:'sr'+k},r))}))"""],cwd=os.path.dirname(os.path.abspath(__file__))))
    TABLES["survey"]=SV["survey"]; TABLES["survey_result"]=SV["survey_result"]
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/plan/community"); await pg.wait_for_timeout(700)
    t=await pg.inner_text("#view")
    check("community: the survey, importance of each priority, themes by mentions", "Community survey, spring 2026" in t and "612 responses" in t and "4.7" in t and t.find("Elementary classrooms are crowded")<t.find("Fix the track"), t[:400])
    n0=len(calls); await pg.select_option("select[data-sr-link='sr8'][data-f=priority_id]", DIRR["priority"][1]["id"]); await pg.wait_for_timeout(500)
    check("community: link a theme to a priority", any(c[0]=="PATCH" and "/rest/v1/survey_result?id=eq.sr8" in c[1] and DIRR["priority"][1]["id"] in (c[2] or "") for c in calls[n0:]))
    svcsv=b"Kind,Label,Value,Mentions,Priority\r\nTheme,Longer library hours,,33,\r\nQuestion,Would support a bond for a CTE building,58,,\r\n"
    await pg.set_input_files("input[data-dir-upload=survey]", files=[{"name":"survey.csv","mimeType":"text/csv","buffer":svcsv}]); await pg.wait_for_timeout(500)
    check("community: survey results upload reviewed first", "2 survey results" in await pg.inner_text("#dir-upload"))
    n0=len(calls); await pg.click("button[data-action=dirUploadApply]"); await pg.wait_for_timeout(600)
    su=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/survey_result"]
    check("community: uploaded results added to the survey", su and len(su[0])==2 and su[0][0]["survey_id"]==SV["survey"][0]["id"] and su[0][1]["kind"]=="question", str(su)[:200])
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/share/community"); await pg.wait_for_timeout(700)
    await pg.uncheck("input[data-pub-hold]"); await pg.wait_for_timeout(700)
    check("community page: offers “What you told us” from the latest survey", "(Community survey, spring 2026)" in await pg.inner_text("#view") and await pg.is_checked("input[data-pub-survey]"))
    await pg.click("button[data-action=previewPublic]"); await pg.wait_for_timeout(600)
    top=await pg.evaluate("(()=>{const b=document.querySelector('[data-modal]'),m=b&&b.querySelector('.modal');return [b?b.scrollTop:-1,m?m.scrollTop:-1,(document.activeElement||{}).tagName]})()")
    check("public preview opens at the top, with nothing focused part-way down", top[0]==0 and top[1]==0 and top[2] not in ("INPUT","SELECT"), str(top))
    pv=await pg.inner_text("[data-modal]")
    check("community page: the preview shows what people said, linked to the plan", "What you told us" in pv and "Elementary classrooms are crowded" in pv and "141 mentions" in pv and "out of 5" in pv, pv[pv.find("What you told us"):pv.find("What you told us")+300])
    await pg.click("[data-modal] button[data-action=closeModal] >> nth=0"); await pg.wait_for_timeout(200)
    await pg.fill("textarea[data-pub-note]","The board adopts the plan in October. Thank you for the survey responses.")
    n0=len(calls); await pg.click("button[data-action=publishBoard]"); await pg.wait_for_timeout(700)
    ps=[json.loads(c[2])["payload"] for c in calls[n0:] if c[0]=="POST" and "/rest/v1/publication" in c[1]]
    check("community page: the published copy carries the survey summary, not individual answers", ps and ps[0]["survey"]["responses"]==612 and len(ps[0]["survey"]["themes"])==6 and "results" not in ps[0]["survey"], str(ps[0].get("survey") if ps else None)[:200])
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/plan/priorities"); await pg.wait_for_timeout(700)
    check("priorities: community importance from the latest survey", "Community importance: 4.7 out of 5" in await pg.inner_text("#view"))
    gcsv=b"Priority,Outcome,Measure,Unit,Better,Start,Start period,Target,Target period,Owner,Cadence\r\nStrong partnerships,Business partners,Active business partners,partners,Higher,4,2024-25,12,2027-28,Superintendent,Yearly\r\n\"Safe, modern places to learn\",Safe and secure buildings,Buildings with a secure entry,buildings,Higher,1,2024-25,3,2026-27,,Yearly\r\n"
    await pg.set_input_files("input[data-dir-upload=goals]", files=[{"name":"goals.csv","mimeType":"text/csv","buffer":gcsv}]); await pg.wait_for_timeout(500)
    n0=len(calls); await pg.click("button[data-action=dirUploadApply]"); await pg.wait_for_timeout(800)
    gp=[json.loads(c[2])["name"] for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/priority"]
    gm=[json.loads(c[2])["name"] for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/measure"]
    check("goals upload: new priorities and measures added; existing ones kept by name", gp==["Strong partnerships"] and gm==["Active business partners"], str(gp)+str(gm))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/track/measures"); await pg.wait_for_timeout(600)
    rcsv=b"Measure,Period,Result,Note\r\nFour-year graduation rate,2026-27,90.7,Preliminary\r\nNo such measure,2026-27,1,\r\n"
    await pg.set_input_files("input[data-dir-upload=results]", files=[{"name":"results.csv","mimeType":"text/csv","buffer":rcsv}]); await pg.wait_for_timeout(500)
    check("results upload: an unknown measure blocks Apply, explained", "no measure called “No such measure”" in await pg.inner_text("#dir-upload") and await pg.is_disabled("button[data-action=dirUploadApply]"))
    # Phase 5D: Overview → Today, and search
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pg.wait_for_timeout(800)
    t=await pg.inner_text("#view")
    check("overview: gap, balances, ledger, measures, decisions ahead, latest report", "Gap to close" in t and "$5.35M" in t and "SAVE balance" in t and "Measures on track" in t and "4 of 8" in t and "Decisions ahead" in t and "Latest board report" in t, t[:600])
    await pg.locator("#view").screenshot(path=SHOTS+"/today.png")
    check("overview: needs attention lists what's off track", "Needs attention" in t and "is off track" in t and "The strategic plan" in t)
    check("overview: recent changes for admins", "Recent changes" in t)
    mname=(await pg.locator("#view [data-measure-go]").first.inner_text()).strip()
    await pg.locator("#view [data-measure-go]").first.click(); await pg.wait_for_timeout(900)
    fr=await pg.locator("tr.focus-row").inner_text() if await pg.locator("tr.focus-row").count() else ""
    inview=await pg.evaluate("(()=>{const r=document.querySelector('tr.focus-row');if(!r)return false;const b=r.getBoundingClientRect();return b.top>=0&&b.bottom<=innerHeight})()")
    check("measures: an off-track measure on Home opens Track → Measures at that measure, highlighted", "/track/measures" in pg.url and mname and mname in fr and inview, f"{mname} | {fr[:80]}")
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/share/plans"); await pg.wait_for_timeout(800)
    pl=await pg.locator("#view [data-measure-go]").count()
    if pl:
      await pg.locator("#view [data-measure-go]").first.click(); await pg.wait_for_timeout(900)
    check("measures: each measure in the district plan opens its row on Track → Measures", pl>0 and "/track/measures" in pg.url and await pg.locator("tr.focus-row").count()==1, str(pl))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pg.wait_for_timeout(600)
    bsc=[x["id"] for x in TABLES.get("scenario",[]) if x.get("is_board_version")]
    inb={x["initiative_id"] for x in TABLES.get("phase",[])+TABLES.get("recurring_cost",[]) if bsc and x.get("scenario_id")==bsc[0]}
    pend=[i for i in TABLES["initiative"] if (i.get("status") or "proposed") in ("idea","proposed","analysis") and i["id"] in inb]
    check("overview: initiatives not yet approved but counted in the board version are called out", (("counted in the board version’s numbers" in t) == bool(pend)) and bool(pend), f"{len(pend)} pending in board; "+t[:900])
    await pg.click("button[data-action=openSearch]"); await pg.wait_for_timeout(500)
    await pg.fill("[data-search]","roof"); await pg.wait_for_timeout(200)
    hits=await pg.inner_text("[data-search-results]")
    check("search: finds initiatives as you type", "High school roof" in hits and "Initiative" in hits, hits)
    await pg.fill("[data-search]","graduation"); await pg.wait_for_timeout(200)
    check("search: finds measures too", "Four-year graduation rate" in await pg.inner_text("[data-search-results]"))
    await pg.click("[data-search-results] a >> nth=0"); await pg.wait_for_timeout(700)
    check("search: going to a result", "/track/measures" in pg.url)
    for t in ["priority","outcome","measure","measure_value","survey","survey_result"]: TABLES[t]=[]
    for i0 in TABLES["initiative"]: i0.pop("priority_id", None)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/help/guide"); await pg.wait_for_timeout(500)
    htabs=[x.strip() for x in await pg.locator("nav.tabs a").all_inner_texts()]
    check("help: no “What’s built” page (everything is built)", "What’s built" not in htabs and "What’s built" not in await pg.inner_text("body"), str(htabs))
    # step B: lead with the answer; Funds chart; folds open for editors; moved screens
    leads={}; gaps={}
    for route in ["home/today","plan/priorities","track/measures","plan/initiatives","plan/ranking","plan/scenarios","money/summary","money/general","track/initiatives","money/capital","track/actuals","track/registers","share/board","share/plans","share/community","settings/overview","settings/setup","plan/community"]:
      await pg.goto("http://localhost:8765/#/d/ironwood-valley/"+route); await pg.wait_for_timeout(600)
      leads[route]=await pg.locator("p.lead").first.inner_text() if await pg.locator("p.lead").count() else ""
      tight=await pg.evaluate("""(()=>{const els=[...document.querySelectorAll('#view .card, #view .grid, #view details.fold, #view .tile-card')].filter(e=>e.offsetParent&&!e.closest('[data-modal]'));const out=[];
        for(const a of els)for(const b of els){if(a===b||a.contains(b)||b.contains(a))continue;const r=a.getBoundingClientRect(),q=b.getBoundingClientRect();const g=q.top-r.bottom;
          if(g>=0&&g<8&&Math.min(r.right,q.right)-Math.max(r.left,q.left)>40)out.push((a.className+' > '+b.className).slice(0,80)+' '+g.toFixed(1))}return out.slice(0,5)})()""")
      gaps[route]=tight
    check("spacing: blocks stacked on a screen keep a real gap between them (no near-touching cards)", not any(gaps.values()), str({k:v for k,v in gaps.items() if v}))
    check("lead sentences on the main screens", all(leads[r] for r in ["home/today","plan/initiatives","plan/ranking","plan/scenarios","money/summary","money/general","track/initiatives"]), str({k:v[:60] for k,v in leads.items()}))
    check("funds: leads with the shortfall, the lowest fund and General Fund solvency", "short" in leads["money/summary"] and "is at its lowest" in leads["money/summary"] and "Grants" not in leads["money/summary"] and "General Fund solvency" in leads["money/summary"], leads["money/summary"])
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/money/summary"); await pg.wait_for_timeout(600)
    await pg.screenshot(path=SHOTS+"/funds-new.png", full_page=True)
    check("funds: a chart of each capital fund's balance by year", await pg.locator("svg[aria-label=\"Each capital fund's balance at the end of each year\"]").count()==1)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/money/capital"); await pg.wait_for_timeout(700)
    folds=await pg.evaluate("[...document.querySelectorAll('details.fold')].map(d=>d.open)")
    check("capital plan: details open for people who edit", folds and all(folds), str(folds))
    lead0=await pg.inner_text("#cap-lead")
    await pg.evaluate("(()=>{const el=document.querySelector('[data-lever=infl]'); el.value='0.08'; el.dispatchEvent(new Event('input',{bubbles:true}));})()"); await pg.wait_for_timeout(300)
    check("capital plan: the lead sentence follows the levers", await pg.inner_text("#cap-lead")!=lead0)
    await pg.click("button[data-action=capReset]"); await pg.wait_for_timeout(300)
    moved={}
    for old,new in [("resources/assumptions","settings/assumptions"),("reports/exports","settings/exports"),("resources/funds","money/summary"),("progress/measures","track/measures")]:
      await pg.goto("http://localhost:8765/#/d/ironwood-valley/"+old); await pg.wait_for_timeout(500); moved[old]=new in pg.url
    check("moved screens: old links land in their new homes", all(moved.values()), str(moved))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/district"); await pg.wait_for_timeout(500)
    st=[x.strip() for x in await pg.locator("nav.tabs a").all_inner_texts()]
    check("settings: assumption sets, uploads and exports live here now, after the Settings overview", st[0]=="Settings" and "Assumption sets" in st and "Uploads" in st and "Exports" in st, str(st))
    # the menu: Home and four areas; old area links land on their first screen; the Settings overview is cards
    old={}
    for o,n in [("overview","home/today"),("direction","plan/priorities"),("decisions","plan/initiatives"),("resources","money/capital"),("progress","track/initiatives"),("reports","share/board")]:
      await pg.goto("http://localhost:8765/#/d/ironwood-valley/"+o); await pg.wait_for_timeout(400); old[o]=n in pg.url
    check("moved areas: an old area link lands on its first screen in the new menu", all(old.values()), str(old))
    heads=[x.strip() for x in await pg.locator("nav.rail .nav:not(.nav-foot) .it .tx").all_inner_texts()]
    check("menu: Home and four areas, Plan, Money, Track and Share", heads==["Home","Plan","Money","Track","Share"], str(heads))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/overview"); await pg.wait_for_timeout(700)
    await pg.screenshot(path=SHOTS+"/settings-overview.png", full_page=True)
    sc=await pg.locator(".scard h3").all_inner_texts(); gh=await pg.locator(".sgrp .gh h2").all_inner_texts()
    check("settings overview: grouped cards, each with where it stands", gh==["Your district","Planning numbers","Records and you"] and len(sc)==8 and await pg.locator(".scard .stat2").count()==8, str(gh)+str(sc))
    check("settings overview: the setup banner shows progress", "of 8 done" in await pg.inner_text("#view"))
    await pg.click(".scard:has-text('Starting numbers')"); await pg.wait_for_timeout(500)
    check("settings overview: a card opens its setting", "/settings/setup" in pg.url)
    # Initiatives has two views, List and Ranked; Ranking is no longer its own menu item
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/plan/ranking"); await pg.wait_for_timeout(800)
    kids=[x.strip() for x in await pg.locator("nav.rail .grp[data-area=plan] .kids a").all_inner_texts()]
    curk=await pg.locator("nav.rail .grp[data-area=plan] .kids a[aria-current=page]").all_inner_texts()
    vs=[x.strip() for x in await pg.locator(".viewswitch a").all_inner_texts()]
    check("initiatives: one screen with List and Ranked views; Ranking isn't a separate menu item", "Ranking" not in kids and curk==["Initiatives"] and vs==["List","Ranked"]
          and "Initiatives" in await pg.inner_text("h1") and await pg.locator(".viewswitch a.on").inner_text()=="Ranked", str(kids)+str(curk)+str(vs))
    await pg.click(".viewswitch a:has-text('List')"); await pg.wait_for_timeout(600)
    check("initiatives: List goes back to the list", "/plan/initiatives" in pg.url)
    # one scenario across Initiatives, Ranking, the capital plan and the General Fund, chosen at the top
    opts=await pg.evaluate("[...document.querySelector('[data-top-scenario]').options].map(o=>[o.value,o.text])")
    other=[o for o in opts if "board version" not in o[1]][0]
    await pg.select_option("select[data-top-scenario]", other[0]); await pg.wait_for_timeout(700)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/money/capital"); await pg.wait_for_timeout(900)
    capv=await pg.evaluate("(()=>{const s=document.querySelector('[data-cap-scenario]');return s?s.value:''})()")
    topv=await pg.evaluate("document.querySelector('[data-top-scenario]').value")
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pg.wait_for_timeout(700)
    hl=await pg.inner_text("#scenTop")
    check("scenario picker: the choice carries to the capital plan; screens that always use the board version say so", capv==other[0] and topv==other[0] and "Board version" in hl and await pg.locator("#scenTop select").count()==0, f"{capv} {topv} {other} {hl}")
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/money/capital"); await pg.wait_for_timeout(800)
    bv=[o[0] for o in opts if "board version" in o[1]][0]
    await pg.select_option("select[data-top-scenario]", bv); await pg.wait_for_timeout(800)
    # the menu fits a laptop screen without scrolling, and search sits at the right, between the demo tag and Help
    await pg.set_viewport_size({"width":1366,"height":768}); await pg.goto("http://localhost:8765/#/d/ironwood-valley/plan/initiatives"); await pg.wait_for_timeout(700)
    fit=await pg.evaluate("(()=>{const r=document.querySelector('.rail');return [r.scrollHeight,r.clientHeight]})()")
    order=await pg.evaluate("[...document.querySelector('.frame>.topbar').children].map(e=>e.dataset.action||e.className.split(' ')[0])")
    check("menu: fits a 1366×768 laptop screen with Plan open, without scrolling", fit[0]<=fit[1], str(fit))
    check("top bar: search is a magnifying-glass button on the right, between the demo tag and Help", order.index("openSearch")>order.index("demo-bar") and order.index("openSearch")<order.index("openHelp") and await pg.locator(".frame>.topbar [data-action=openSearch]").evaluate("e=>e.getBoundingClientRect().width")<=40, str(order))
    await pg.screenshot(path=SHOTS+"/laptop-initiatives.png")
    # presenting at a board meeting: no menus, bigger type, the board's view, a bar to move between screens
    await pg.set_viewport_size({"width":1360,"height":900}); await pg.goto("http://localhost:8765/#/d/ironwood-valley/money/capital"); await pg.wait_for_timeout(800)
    await pg.click("button[data-action=startPresent]"); await pg.wait_for_timeout(900)
    pres=await pg.locator(".frame.presenting").count()==1 and await pg.locator(".present-bar").is_visible() and not await pg.locator(".frame>.topbar").is_visible() and not await pg.locator(".rail").is_visible()
    board_only=await pg.locator("button[data-action=capCopy], button[data-action=publishBoard]").count()==0
    await pg.screenshot(path=SHOTS+"/presenting.png")
    u1=pg.url; await pg.keyboard.press("ArrowRight"); await pg.wait_for_timeout(800); u2=pg.url
    check("presenting: the menus go, the board's view shows, and → moves to the next screen", pres and board_only and u1!=u2 and "/money/general" in u2 and await pg.locator(".frame.presenting").count()==1, f"{pres} {board_only} {u1} {u2}")
    await pg.keyboard.press("Escape"); await pg.wait_for_timeout(700)
    check("presenting: Esc ends it and brings back your own view", await pg.locator(".frame.presenting").count()==0 and await pg.locator(".preview-bar").count()==0 and await pg.locator(".rail").is_visible())
    # dates: balances from before this fiscal year start the plan in a year that has ended
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/setup"); await pg.wait_for_timeout(900)
    await pg.fill("form[data-form=saveSetup] [name=as_of]","2025-06-30"); await pg.dispatch_event("form[data-form=saveSetup] [name=as_of]","input"); await pg.wait_for_timeout(300)
    ck=await pg.inner_text("[data-setup-checks]")
    check("starting numbers: balances dated before July 1 of last year are flagged (the plan would start in a year that has ended)", "a year that has already ended" in ck and "FY2026" in ck, ck[:300])
    # the menu: hover an area for its screens; collapse to icons with [ (remembered); Ctrl K finds screens and features
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/money/capital"); await pg.wait_for_timeout(700)
    await pg.hover("[data-area-head=share]"); await pg.wait_for_timeout(250)
    fl=await pg.inner_text(".fly") if await pg.locator(".fly").count() else ""
    check("menu: pointing at an area shows its screens beside the menu", "Board reports" in fl and "District plan" in fl and "Community page" in fl, fl)
    await pg.screenshot(path=SHOTS+"/menu-flyout.png")
    await pg.locator(".fly a", has_text="District plan").click(); await pg.wait_for_timeout(600)
    check("menu: picking a screen from the flyout goes there", "/share/plans" in pg.url and await pg.locator(".fly").count()==0, pg.url)
    await pg.mouse.move(800,400); await pg.keyboard.press("["); await pg.wait_for_timeout(300)
    w=await pg.evaluate("document.querySelector('.rail').getBoundingClientRect().width")
    await pg.reload(); await pg.wait_for_timeout(700)
    kept=await pg.locator(".frame.rail-collapsed").count()==1
    await pg.hover("[data-area-head=money]"); await pg.wait_for_timeout(250)
    rf=await pg.inner_text(".fly") if await pg.locator(".fly").count() else ""
    await pg.screenshot(path=SHOTS+"/menu-collapsed.png")
    check("menu: [ collapses it to icons, it stays collapsed after a reload, and icons still show their screens", w<=80 and kept and "Capital plan" in rf, f"{w} {kept} {rf}")
    await pg.mouse.move(800,400); await pg.click("button[data-action=toggleRail]"); await pg.wait_for_timeout(300)
    check("menu: the panel button opens it again", await pg.locator(".frame.rail-collapsed").count()==0 and await pg.evaluate("document.querySelector('.rail').getBoundingClientRect().width")>200)
    await pg.keyboard.press("Control+k"); await pg.wait_for_timeout(600)
    await pg.fill("[data-search]","what-if"); await pg.dispatch_event("[data-search]","input"); await pg.wait_for_timeout(200)
    sr=await pg.inner_text("[data-search-results]")
    check("search: Ctrl K opens it, and it finds features inside screens", "What-if levers" in sr and "Money / Capital plan" in sr, sr[:200])
    await pg.locator("[data-search-results] a", has_text="What-if levers").click(); await pg.wait_for_timeout(700)
    check("search: a feature takes you to its screen", "/money/capital" in pg.url, pg.url)
    # an admin can preview the board member's view, and come back
    await pg.click("button[data-action=userMenu]"); await pg.wait_for_timeout(200)
    um=await pg.inner_text(".umenu")
    await pg.click(".umenu [data-action=togglePreview]"); await pg.wait_for_timeout(800)
    ph=[x.strip() for x in await pg.locator("nav.rail .grp .kids a").all_inner_texts()]
    pv=await pg.locator(".preview-bar").count()==1 and "/share/board" in pg.url and ph==["Priorities","Capital plan","General fund","All funds","Measures","Board reports","District plan"]
    await pg.screenshot(path=SHOTS+"/board-preview.png")
    await pg.click(".preview-bar [data-action=togglePreview]"); await pg.wait_for_timeout(700)
    check("account menu: Your account, Preview as board member and Sign out; the preview shows the board's menu with a banner, and Back to my view ends it",
          "Your account" in um and "Preview as board member" in um and "Sign out" in um and pv and await pg.locator(".preview-bar").count()==0 and await pg.locator("nav.rail .grp[data-area=track]").count()==1, um+str(ph))
    # the district's logo in place of its name at the top of the menu
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/district"); await pg.wait_for_timeout(600)
    tl=await pg.evaluate("""(async()=>{const mk=async(clear)=>{const c=document.createElement('canvas');c.width=1000;c.height=1000;const g=c.getContext('2d');if(!clear){g.fillStyle='#fff';g.fillRect(0,0,1000,1000);}g.fillStyle='#123';g.fillRect(400,450,300,100);return new Promise(r=>c.toBlob(b=>r(new File([b],'l.png',{type:'image/png'})),'image/png'));};
      const a=await HG.trimLogo(await mk(false)),b=await HG.trimLogo(await mk(true));return [a.w,a.h,a.trimmed,b.w,b.h,b.trimmed]})()""")
    check("logo: empty white or clear margins are cropped away on upload, leaving a little space", tl[0] in range(310,330) and tl[1] in range(110,130) and tl[2] and tl[3] in range(310,330) and tl[4] in range(110,130) and tl[5], str(tl))
    check("district: “Logo only” waits for a logo", await pg.locator("input[data-logo-mode][value=logo]").is_disabled() and await pg.locator("input[data-logo-mode][value=name]").is_checked())
    D1["logo_path"]="d1/logo.png"; D1["menu_logo_only"]=True
    await pg.reload(); await pg.wait_for_timeout(800)
    lg=await pg.locator(".org .org-logo").count()==1 and await pg.locator(".org .org-tx").count()==0
    await pg.click("button[data-action=toggleRail]"); await pg.wait_for_timeout(300); await pg.mouse.move(800,400)
    cx=await pg.evaluate("""(()=>{const c=e=>{const r=e.getBoundingClientRect();return r.left+r.width/2};const m=document.querySelector('.org .rm'),i=[...document.querySelectorAll('.rail .it .sq')].filter(e=>e.offsetParent);return [m&&m.offsetParent?c(m):-99, i.length?c(i[0]):-1]})()""")
    await pg.screenshot(path=SHOTS+"/menu-collapsed-logo.png", clip={"x":0,"y":0,"width":200,"height":400})
    await pg.click("button[data-action=toggleRail]"); await pg.wait_for_timeout(300)
    check("menu: collapsed, the district's mark sits centered over the menu icons", abs(cx[0]-cx[1])<=1.5, str(cx))
    n0=len(calls); await pg.check("input[data-logo-mode][value=name]"); await pg.wait_for_timeout(500)
    lp=[json.loads(c[2]) for c in calls[n0:] if c[0]=="PATCH" and "/rest/v1/district?" in c[1]]
    check("district: with “Logo only”, the logo replaces the initials and name in the menu; switching back saves", lg and lp and lp[0]=={"menu_logo_only":False}, str(lp))
    D1.pop("logo_path",None); D1.pop("menu_logo_only",None)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/district"); await pg.wait_for_timeout(500)
    # step A: search is never stale
    await pg.click("button[data-action=openSearch]"); await pg.wait_for_timeout(400)
    await pg.click("[data-modal] button[data-action=closeModal]"); await pg.wait_for_timeout(100)
    TABLES["initiative"].append({"district_id":"d1","id":"ini-new","name":"Greenhouse for the ag program","type":"capital","status":"idea"})
    await pg.click("button[data-action=openSearch]"); await pg.wait_for_timeout(500)
    await pg.fill("[data-search]","greenhouse"); await pg.wait_for_timeout(200)
    check("search: finds something added a moment ago", "Greenhouse for the ag program" in await pg.inner_text("[data-search-results]"))
    await pg.click("[data-modal] button[data-action=closeModal]"); await pg.wait_for_timeout(100)
    TABLES["initiative"]=[i for i in TABLES["initiative"] if i.get("id")!="ini-new"]
    # Phase 6: the General Fund
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/money/general"); await pg.wait_for_timeout(800)
    t=await pg.inner_text("#view")
    check("general fund: caution, solvency, unspent ratio, ending balance", "Check the starting figures and results with the business manager" in t and "Solvency ratio" in t and "16.8%" in t and "Unspent balance ratio" in t and "Ending fund balance" in t, t[:600])
    check("general fund: new money vs a negotiated raise (total package), and the year table", "New money vs. a negotiated raise" in t and "Each 1% of total package costs" in t and "Spending authority" in t and "Staff share of spending" in t and "$8,148" in t)
    await pg.screenshot(path=SHOTS+"/generalfund.png", full_page=True)
    box=await pg.locator(".gflevers").bounding_box()
    check("general fund: what-if levers laid out compactly, not one tall column", box and box["height"]<260, str(box))
    before=await pg.inner_text("#gf-results")
    await pg.fill("input[data-gf-lever=settle]","4.5"); await pg.wait_for_timeout(300)
    after=await pg.inner_text("#gf-results")
    check("general fund: a what-if total package increase changes the forecast at once", before!=after and "A 4.5% total package costs" in after)
    await pg.fill("input[data-gf-lever=settle]","9"); await pg.fill("input[data-gf-lever=health]","12"); await pg.wait_for_timeout(300)
    check("general fund: warns when spending passes spending authority", "passes spending authority" in await pg.inner_text("#gf-results"))
    await pg.click("button[data-action=gfReset]"); await pg.wait_for_timeout(600)
    await pg.click("button[data-action=gfEdit]"); await pg.wait_for_timeout(500)
    await pg.fill("[data-modal] input[name=enrollment]","1240")
    n0=len(calls); await pg.click("[data-modal] button[type=submit]"); await pg.wait_for_timeout(600)
    gp=[json.loads(c[2]) for c in calls[n0:] if c[0]=="PATCH" and "/rest/v1/district_settings?" in c[1]]
    check("general fund: starting figures saved, staff groups and default assumptions with them", gp and gp[0]["gf_inputs"]["enrollment"]==1240 and len(gp[0]["gf_inputs"]["staff"])==3 and gp[0]["gf_inputs"]["assume"]["settle"]==0.025, str(gp)[:300])
    saved_gf=TABLES["district_settings"][0].get("gf_inputs")
    TABLES["district_settings"][0]["gf_inputs"]=None
    await pg.reload(); await pg.wait_for_timeout(700)
    check("general fund: not set up yet, with a way in", "Set up the General Fund" in await pg.inner_text("#view"))
    TABLES["district_settings"][0]["gf_inputs"]=saved_gf
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/money/summary"); await pg.wait_for_timeout(700)
    check("summary: the General Fund half", "General Fund" in await pg.inner_text("#view") and "Staff share of spending" in await pg.inner_text("#view"))
    # adopted budget upload
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/uploads"); await pg.wait_for_timeout(500)
    TABLES["gl_account"]=[]
    await pg.select_option("select[data-upload-kind]","budget")
    await pg.set_input_files("input[data-upload-file]",files=[{"name":"budget-fy2027.csv","mimeType":"text/csv","buffer":GLCSV}]); await pg.wait_for_timeout(600)
    rv=await pg.inner_text("#upload-review")
    check("budget upload: reviewed with the General Fund totals it gives", "What this budget says" in rv and "revenue $12.30M" in rv and "staff spending (objects 1xx–2xx) $5.28M" in rv, rv[-400:])
    n0=len(calls); await pg.click("button[data-action=applyUpload]"); await pg.wait_for_timeout(1000)
    bb=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].split("?")[0]=="/rest/v1/import_batch"]
    bl=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and c[1].startswith("/rest/v1/budget_line")]
    check("budget upload: saved as the adopted budget for the year", bb and bb[0]["kind"]=="budget" and bb[0]["period_end"] is None and bl and all(x["version"]=="adopted" and x["fiscal_year"]==bb[0]["fiscal_year"] for x in bl[0]) and len(bl[0])==16, str(bb)[:150]+str(len(bl[0]) if bl else 0))
    check("budget upload: no balances touched", not any(c[1].startswith("/rest/v1/gl_amount") for c in calls[n0:]) and "The adopted FY" in await pg.inner_text("#toasts"))
    TABLES["gl_account"]=[]
    # milestone 6: summary, all funds, exports, activity
    EXP=json.loads(subprocess.check_output(["node","-e","""
      const C=require('./capital.js'),E=require('./engine.js'),D=require('./demo_data.js');let i=0;
      const R=C.demoRows(D['ironwood-valley'],'d1',()=>'00000000-0000-4000-8000-'+String(++i).padStart(12,'0'));
      const rows={district:{name:'x'},settings:R.district_settings[0],balances:R.fund_balance,debts:R.debt_obligation,scenarios:R.scenario,initiatives:R.initiative,phases:R.phase,funding:R.phase_funding,financing:R.financing};
      const sc=R.scenario.find(s=>s.is_board_version);const inp=C.buildInputs(rows,sc.id);const r=E.compute(inp.projects,inp.levers,inp.cfg);
      const P=C.fundPaths(r,inp.cfg);process.stdout.write(JSON.stringify({save27:P.save.years[0].end,saveLow:P.save.low,saveLowFY:P.save.lowFY,phases:R.phase.length}));"""],cwd=os.path.dirname(os.path.abspath(__file__))))
    money=lambda v: "$"+format(round(v),",")
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/money/summary"); await pg.wait_for_timeout(700); await pg.evaluate("document.querySelectorAll('details.fold').forEach(d=>d.open=true)"); await pg.evaluate("document.querySelectorAll('details.fold').forEach(d=>d.open=true)")
    t=await pg.inner_text("#view")
    check("summary: board version totals", "$14.79M" in t and "$5.35M" in t and "District baseline" in t)
    check("summary: each fund's low point, rounded on the summary", fmtK(EXP["saveLow"])+" (FY"+str(EXP["saveLowFY"])+")" in t, fmtK(EXP["saveLow"]))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/resources/funds"); await pg.wait_for_timeout(900); await pg.evaluate("document.querySelectorAll('details.fold').forEach(d=>d.open=true)")
    t=await pg.inner_text("#view")
    kfmt=lambda v: "$%.2fM"%(v/1e6) if v>=1e6 else "$%dk"%round(v/1000)
    check("all funds: year-by-year from the engine (forecast tables in compact dollars)", "SAVE" in t and kfmt(EXP["save27"]) in t and "Borrowing room" in t and "$4.59M" in t, kfmt(EXP["save27"]))
    check("all funds: what each fund may pay for, with a caution", "423F" in t and "298.3" in t and "not legal advice" in t)
    await pg.screenshot(path=SHOTS+"/funds.png",full_page=True)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/reports/exports"); await pg.wait_for_timeout(500)
    import csv, io as _io, openpyxl as _ox
    def xl2csv(path):   # the downloads are Excel now: turn the first sheet back into CSV text for the checks
      out=_io.StringIO(); w=csv.writer(out, lineterminator="\r\n")
      for r in _ox.load_workbook(_io.BytesIO(open(path,"rb").read())).active.iter_rows(values_only=True): w.writerow(["" if v is None else v for v in r])
      return out.getvalue()
    async with pg.expect_download() as dl: await pg.click("button[data-action=exportProjects]")
    f=await dl.value; text=xl2csv(await f.path())
    names={r["Project"] for r in csv.DictReader(_io.StringIO(text))}
    reparsed=json.loads(subprocess.run(["node","-e","const U=require('./uploads.js');let s='';process.stdin.on('data',d=>s+=d).on('end',()=>{const p=U.parseProjects(U.parseCSV(s),2027,10);process.stdout.write(JSON.stringify({n:p.projects.length,e:p.issues.filter(i=>i.l==='e').length}))})"],input=text,capture_output=True,text=True,cwd=os.path.dirname(os.path.abspath(__file__))).stdout)
    check("exports: projects file is a valid upload with every project", len(names)==16 and reparsed=={"n":16,"e":0} and f.suggested_filename.startswith("ironwood-valley-"), str(reparsed))
    async with pg.expect_download() as dl: await pg.click("button[data-action=exportPhases]")
    f=await dl.value; prow=list(csv.DictReader(_io.StringIO(xl2csv(await f.path()))))
    check("exports: every scenario's phases in one sheet", len(prow)==EXP["phases"] and {r["Scenario"] for r in prow}=={"District baseline","Addition phased, bond in FY2030"}, str(len(prow)))
    async with pg.expect_download() as dl: await pg.click("button[data-action=exportBackup]")
    f=await dl.value; bk=json.load(open(await f.path()))
    BACKUP_PATH=await f.path()
    check("exports: full backup has the plan, not people", bk["kind"]=="HighGround district backup" and len(bk["tables"]["phase"])==EXP["phases"] and "district_member" not in bk["tables"] and "audit_log" not in bk["tables"])
    check("exports: the backup includes the monthly ledger and adopted budget", all(t in bk["tables"] for t in ["gl_account","gl_amount","budget_line","report_snapshot"]))
    check("exports: the backup message is plain", "The whole plan, as one file" in await pg.inner_text("#toasts"))
    # restore from a backup
    import shutil; shutil.copy(BACKUP_PATH,"/tmp/hg-backup.json")
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/reports/exports"); await pg.wait_for_timeout(500)
    check("restore: offered to admins, with the safety note", "Restore from a backup" in await pg.inner_text("#view") and "downloads a fresh backup" in await pg.inner_text("#view"))
    other=json.load(open("/tmp/hg-backup.json")); other["district"]["id"]="d-somewhere-else"; other["district"]["name"]="Another CSD"; json.dump(other,open("/tmp/hg-other.json","w"))
    await pg.set_input_files("input[data-restore-file]","/tmp/hg-other.json"); await pg.wait_for_timeout(400)
    check("restore: a backup from another district is refused", "That backup is from “Another CSD”" in await pg.inner_text("#toasts"))
    await pg.set_input_files("input[data-restore-file]","/tmp/hg-backup.json"); await pg.wait_for_timeout(400)
    rv=await pg.inner_text("[data-restore-review]")
    check("restore: shows what the backup holds before anything changes", "Backup from" in rv and "2 scenarios" in rv and ("%d initiatives"%len(TABLES["initiative"])) in rv and "1 yearly cost ·" in rv, rv[:200])
    PROMPT["text"]="wrong-district"; n0=len(calls); await pg.click("button[data-action=restoreRun]"); await pg.wait_for_timeout(400)
    check("restore: a mistyped link id changes nothing", "Not restored" in await pg.inner_text("#toasts") and not any(c[0] in ("DELETE","POST") for c in calls[n0:]))
    PROMPT["text"]="ironwood-valley"; n0=len(calls)
    async with pg.expect_download() as dl2: await pg.click("button[data-action=restoreRun]")
    safety=await dl2.value; await pg.wait_for_timeout(1500); PROMPT["text"]=None
    new=calls[n0:]
    dels=[c[1].split("?")[0].split("/")[-1] for c in new if c[0]=="DELETE"]
    posts=[c[1].split("?")[0].split("/")[-1] for c in new if c[0]=="POST"]
    locks=[c for c in new if c[0]=="PATCH" and "/rest/v1/scenario?" in c[1] and '"is_locked": true' in (c[2] or "").replace('":true','": true')]
    first_del=next(k for k,c in enumerate(new) if c[0]=="DELETE")
    check("restore: a safety backup downloads before anything is erased", safety.suggested_filename.endswith("-backup.json") and first_del>0)
    check("restore: erases the plan, then puts every table back in order", dels==["publication","report_snapshot","project_request","budget_line","gl_amount","gl_account","scenario","measure","outcome","survey","initiative","priority","assumption_set","debt_obligation","fund_balance","import_batch","district_settings"]
          and posts[:3]==["district_settings","initiative","scenario"] and posts.index("phase")>posts.index("scenario") and posts.index("fund_balance")>posts.index("phase"), str(posts))
    scen=[json.loads(c[2]) for c in new if c[0]=="POST" and c[1].split("?")[0].endswith("/scenario")][0]
    check("restore: scenarios go back unlocked, then the locked ones are locked again", all(x["is_locked"] is False for x in scen) and len(locks)==sum(1 for sc in TABLES["scenario"] if sc.get("is_locked")), f"{len(locks)} locks")
    check("restore: says it's done", "Plan restored" in await pg.inner_text("#toasts"))
    # district logo
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/district"); await pg.wait_for_timeout(500)
    import base64 as _b64
    open("/tmp/logo.png","wb").write(_b64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="))
    n0=len(calls); await pg.set_input_files("input[data-logo-file]","/tmp/logo.png"); await pg.wait_for_timeout(700)
    up=[c for c in calls[n0:] if c[0]=="POST" and "/storage/v1/object/district-public/d1/logo-" in c[1]]
    lp=[json.loads(c[2]) for c in calls[n0:] if c[0]=="PATCH" and "/rest/v1/district?" in c[1]]
    check("logo: uploads to the district's public folder and is saved on the district", up and lp and lp[0]["logo_path"].startswith("d1/logo-") and lp[0]["logo_path"].endswith(".png"), str(lp))
    open("/tmp/notimage.txt","w").write("hello")
    await pg.set_input_files("input[data-logo-file]",{"name":"logo.txt","mimeType":"text/plain","buffer":b"hello"}); await pg.wait_for_timeout(300)
    check("logo: only images are accepted", "Use a PNG, JPEG or WebP image" in await pg.inner_text("#toasts"))
    D1["logo_path"]="d1/logo-1.png"
    await pg.goto("http://localhost:8765/#/p/ironwood-valley"); await pg.wait_for_timeout(700)
    check("logo: shown on the public board page", await pg.locator("img.pub-dlogo-lg[src*='/storage/v1/object/public/district-public/d1/logo-1.png']").count()==1)
    del D1["logo_path"]
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/activity"); await pg.wait_for_timeout(500)
    t=await pg.inner_text("#view")
    check("activity: who changed what, before and after", "Pat Admin" in t and "Project: Gym floor" in t and "focus area: Facilities → Activities" in t and "updated at" not in t)
    check("activity: someone outside the district shows as staff", "Willow Holler staff" in t and "Scenario: Plan B" in t)
    await pg.goto("http://localhost:8765/#/d/cottonwood-ridge/settings/activity"); await pg.wait_for_timeout(400)
    check("activity: admins only (a viewer can't open it)", "Not part of your view" in await pg.inner_text("#content"))
    # help map
    await signout(pg); await pg.wait_for_timeout(400)
    check("sign out", "Sign in" in await pg.inner_text("h1") and any("/auth/v1/logout" in c[1] for c in calls))
    # no-district user
    await pg.fill("input[name=email]","new@example.test"); await pg.fill("input[name=password]",PW); await pg.click("button[type=submit]"); await pg.wait_for_timeout(600)
    check("no-district welcome", "don’t have access to a district yet" in await pg.inner_text("body"))
    check("no Search button outside a district", await pg.locator("[data-action=openSearch]").count()==0)
    await pg.fill("form[data-form=requestAccess] input[name=slug]","ironwood-valley"); await pg.fill("form[data-form=requestAccess] textarea","I'm the new principal")
    n0=len(calls); await pg.click("form[data-form=requestAccess] button"); await pg.wait_for_timeout(500)
    rq=[json.loads(c[2]) for c in calls[n0:] if "rpc/request_access" in c[1]]
    check("request access: sent by link id", rq==[{"p_slug":"ironwood-valley","p_message":"I'm the new principal"}] and "Request sent" in await pg.inner_text("#toasts"), str(rq))
    await signout(pg); await pg.wait_for_timeout(300)
    # two-step sign-in: someone who has it on is asked for a code
    await pg.fill("input[name=email]","mfa@example.test"); await pg.fill("input[name=password]",PW); await pg.click("button[type=submit]"); await pg.wait_for_timeout(600)
    check("two-step: code asked for after the password", "Enter the 6-digit code" in await pg.inner_text("body"))
    await pg.goto("http://localhost:8765/#/staff"); await pg.wait_for_timeout(500)
    check("two-step: nothing else opens until the code is entered", "Enter the 6-digit code" in await pg.inner_text("body"))
    await pg.fill("input[name=code]","123456"); await pg.click("button[type=submit]"); await pg.wait_for_timeout(700)
    check("two-step: right code lets them in", "Welcome to HighGround" in await pg.inner_text("body"))
    await signout(pg); await pg.wait_for_timeout(300)
    # staff with no districts yet
    await pg.fill("input[name=email]","empty.staff@example.test"); await pg.fill("input[name=password]",PW); await pg.click("button[type=submit]"); await pg.wait_for_timeout(600)
    t=await pg.inner_text("body")
    check("staff with no districts lands on Willow Holler page", "#/staff" in pg.url and "Add a district" in t and "don’t have access" not in t, pg.url)
    await signout(pg); await pg.wait_for_timeout(300)
    # staff
    await pg.evaluate("localStorage.removeItem('highground-last-district')")
    await pg.fill("input[name=email]","staff@example.test"); await pg.fill("input[name=password]",PW); await pg.click("button[type=submit]"); await pg.wait_for_timeout(600)
    check("staff land on the Willow Holler page", pg.url.endswith("#/staff"), pg.url)
    import datetime as _dt
    _now=_dt.datetime.now(_dt.timezone.utc)
    TABLES["hg_health_run"]=[{"ran_at":(_now-_dt.timedelta(hours=2)).isoformat(),"status":"fail","n_ok":30,"n_warn":1,"n_fail":1,"emailed":True,
      "checks":[{"area":"data","check":"Tax rates are current","status":"fail","detail":"latest FY2026"},{"area":"connections","check":"Census Bureau key works","status":"warn","detail":"no key"},{"area":"code","check":"Every calculation test passes","status":"ok","detail":"19 test files"}]},
      {"ran_at":(_now-_dt.timedelta(days=1,hours=2)).isoformat(),"status":"ok","n_ok":32,"n_warn":0,"n_fail":0,"emailed":False,"checks":[]}]
    await pg.goto("http://localhost:8765/#/staff?x=1"); await pg.goto("http://localhost:8765/#/staff"); await pg.wait_for_timeout(500)
    ht=await pg.inner_text("[data-health]")
    check("staff: System health shows the latest daily check, what failed, and that it was emailed", "Something is broken" in ht and "Tax rates are current" in ht and "latest FY2026" in ht and "emailed" in ht and await pg.locator("[data-health] .hdot").count()==2, ht[:300])
    TABLES["hg_health_run"]=[]
    await pg.goto("http://localhost:8765/#/staff"); await pg.wait_for_timeout(400)
    await pg.fill("input[name=name]","Harvest Plains Community School District"); await pg.fill("input[name=slug]","harvest-plains"); await pg.check("input[name=is_demo]"); await pg.fill("input[name=admin_email]","Boss@Example.test")
    await pg.click("form[data-form=createDistrict] button"); await pg.wait_for_timeout(500)
    dp=[c for c in calls if c[0]=="POST" and c[1].startswith("/rest/v1/district?")]
    ip=[c for c in calls if c[0]=="POST" and "/rest/v1/invitation" in c[1]]
    check("staff creates district and admin invite", dp and json.loads(dp[-1][2])["slug"]=="harvest-plains" and json.loads(ip[-1][2])=={"district_id":"d9","email":"boss@example.test","role":"admin"})
    check("staff sees Willow Holler page", "Added Harvest Plains" in await pg.inner_text("body"))
    # deleting a district: a confirmation that names what goes, needs the link id typed, then removes it
    await pg.click("tr:has-text('Cottonwood Ridge') button[data-action=deleteDistrictAsk]"); await pg.wait_for_timeout(500)
    dm=await pg.inner_text("[data-modal]")
    off0=await pg.locator("[data-delete-go]").is_disabled()
    await pg.fill("[data-delete-confirm]","cottonwood"); off1=await pg.locator("[data-delete-go]").is_disabled()
    await pg.fill("[data-delete-confirm]","cottonwood-ridge"); on=await pg.locator("[data-delete-go]").is_enabled()
    await pg.screenshot(path=SHOTS+"/delete-district.png")
    n0=len(calls); await pg.click("[data-delete-go]"); await pg.wait_for_timeout(800)
    dels=[c[1] for c in calls[n0:] if c[0]=="DELETE"]
    check("staff: delete a district only after typing its link id; it says what goes and offers the backup first",
          "Delete Cottonwood Ridge" in dm and "can’t be undone" in dm and "Download its backup" in dm and off0 and off1 and on
          and any(d.startswith("/rest/v1/district?id=eq.d2") for d in dels) and "was deleted" in await pg.inner_text("body"), str(dels)+dm[:200])
    await pg.select_option("select[data-demo-district]","d2"); await pg.select_option("select[data-demo-set]","harvest-plains")
    n0=len(calls); await pg.click("button[data-action=loadDemo]"); await pg.wait_for_timeout(900)
    posts=[c[1].split("?")[0].split("/")[-1] for c in calls[n0:] if c[0]=="POST"]
    locks=[c for c in calls[n0:] if c[0]=="PATCH" and "/rest/v1/scenario?" in c[1] and "is_locked" in (c[2] or "")]
    want=["district_settings","fund_balance","debt_obligation","initiative","scenario","scenario_initiative","phase","phase_funding"]
    check("staff loads demo data in dependency order, then locks the baseline", posts==want and len(locks)==1, str(posts))
    # reset a demo district to Bridger Hollow
    await pg.goto("http://localhost:8765/#/staff"); await pg.wait_for_timeout(600)
    check("staff: Bridger Hollow is offered first", (await pg.locator("select[data-demo-set] option").first.inner_text()).startswith("Bridger Hollow Community School District (925 students)"))
    await pg.select_option("select[data-demo-district]","d2"); await pg.select_option("select[data-demo-set]","bridger-hollow")
    n0=len(calls); await pg.click("button[data-action=resetDemo]"); await pg.wait_for_timeout(1200)
    dels=[c[1].split("?")[0].split("/")[-1] for c in calls[n0:] if c[0]=="DELETE"]
    posts=[c[1].split("?")[0].split("/")[-1] for c in calls[n0:] if c[0]=="POST"]
    check("reset: erases the demo district's plan, then reloads it", dels==["publication","report_snapshot","project_request","budget_line","gl_amount","gl_account","scenario","measure","outcome","survey","initiative","priority","assumption_set","debt_obligation","fund_balance","import_batch","district_settings"]
          and posts[:16]==["district_settings","fund_balance","debt_obligation","priority","outcome","measure","measure_value","initiative","scenario","scenario_initiative","phase","phase_funding","recurring_cost","financing","survey","survey_result"], str(dels)+str(posts))
    check("reset: only that district is touched", all("district_id=eq.d2" in c[1] for c in calls[n0:] if c[0]=="DELETE"))
    check("reset: says it's done", "Demo reset" in await pg.inner_text("#toasts"))
    await pg.screenshot(path=SHOTS+"/staff.png",full_page=True)
    await signout(pg); await pg.wait_for_timeout(300)
    # sign up
    await pg.goto("http://localhost:8765/#/signup"); await pg.wait_for_timeout(200)
    await pg.fill("input[name=full_name]","Nia New"); await pg.fill("input[name=email]","nia@example.test"); await pg.fill("input[name=password]","short1"); await pg.fill("input[name=again]","short1")
    await pg.evaluate("document.querySelectorAll('input[minlength]').forEach(i=>i.removeAttribute('minlength'))")
    await pg.click("button[type=submit]"); await pg.wait_for_timeout(300)
    check("short password refused", "8 characters" in await pg.inner_text("#toasts"))
    await pg.fill("input[name=password]","longenough"); await pg.fill("input[name=again]","longenough"); await pg.click("button[type=submit]"); await pg.wait_for_timeout(300)
    check("password without capital, number, symbol refused", "one capital" in await pg.inner_text("#toasts") and not any("/auth/v1/signup" in c[1] for c in calls))
    await pg.fill("input[name=password]","Longer pass 9!"); await pg.fill("input[name=again]","Longer pass 9!"); await pg.click("button[type=submit]"); await pg.wait_for_timeout(400)
    su=[c for c in calls if "/auth/v1/signup" in c[1]]
    check("sign up asks to confirm email", su and "redirect_to=" in su[-1][1] and "confirmation link" in await pg.inner_text("body"))
    # recovery link
    await pg.goto("http://localhost:8765/#access_token=tok-u-admin&refresh_token=r&expires_in=3600&type=recovery"); await pg.wait_for_timeout(600)
    check("reset link opens new-password screen", "Choose a new password" in await pg.inner_text("body") and "access_token" not in pg.url, pg.url)
    await pg.fill("input[name=password]","Another pass 7#"); await pg.fill("input[name=again]","Another pass 7#"); await pg.click("button[type=submit]"); await pg.wait_for_timeout(600)
    check("password saved", any(c[0]=="PUT" and "/auth/v1/user" in c[1] for c in calls) and "Password changed" in await pg.inner_text("body"))
    await pg.goto("http://localhost:8765/#error=access_denied&error_description=Email+link+is+invalid+or+has+expired"); await pg.wait_for_timeout(500)
    check("expired link explained", "expired" in await pg.inner_text("body"))
    # first-session wizard: number → starting numbers → General Fund → done
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/welcome"); await pg.wait_for_timeout(800)
    w=await pg.inner_text(".wiz")
    check("wizard: full screen, three steps, starting with the Iowa number", "Find Ironwood Valley in the state’s data" in w and "1. Find the district" in w and "Exit setup" in w and await pg.locator(".rail").count()==0, w[:300])
    await pg.click("form[data-form=wizNumber] button[type=submit]"); await pg.wait_for_timeout(900)
    w=await pg.inner_text(".wiz")
    check("wizard: step 2 has the state's figures filled in", pg.url.endswith("/welcome/start") and "Filled from the state’s annual report (FY2025)" in w and "Also filled in from the state’s data" in w and "General-obligation bonds of $295,000" in w and await pg.input_value("form[data-form=wizStart] [name=save_receipts]")!="" , pg.url+" "+w[:200])
    n0=len(calls); await pg.click("form[data-form=wizStart] button[type=submit]"); await pg.wait_for_timeout(900)
    saved=[c for c in calls[n0:] if c[0]=="POST" and "district_settings" in c[1]]
    check("wizard: step 2 saves the starting numbers (same rules as the full screen) and moves on", saved and pg.url.endswith("/welcome/gf") and not [c for c in calls[n0:] if c[0]=="DELETE"], pg.url)
    wb=json.loads(saved[0][2])[0] if saved else {}
    check("wizard: step 2 keeps what's already saved, carries the rest from the state's data, and doesn't blank fields it doesn't show", wb.get("taxable_valuation")==727273000 and wb.get("ppel_growth")==0.035 and wb.get("save_ongoing")==250000 and "go_outstanding" not in wb and "save_receipts" in wb, str(wb)[:600])
    await pg.wait_for_timeout(500)
    check("wizard: step 3 is the General Fund form, state figures filled in", await pg.locator("form[data-form=wizGf]").count()==1 and await pg.input_value("form[data-form=wizGf] [name=fund_balance]")=="2,250,000")
    n0=len(calls); await pg.click("form[data-form=wizGf] button[type=submit]"); await pg.wait_for_timeout(900)
    check("wizard: step 3 saves and finishes with what's next", any(c[0]=="PATCH" and "district_settings" in c[1] and "gf_inputs" in (c[2] or "") for c in calls[n0:]) and pg.url.endswith("/welcome/done") and "Upload projects" in await pg.inner_text(".wiz"), pg.url)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pg.wait_for_timeout(600)
    # peer comparisons from the state's annual reports, and the state prefill
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/money/general"); await pg.wait_for_timeout(700)
    t=await pg.inner_text("#view")
    check("peers: General Fund shows its own callout only", "Compared with districts your size" in t and "General Fund · Student Transportation spending" in t and "1.8× the average of 41" in t and "SAVE · ending" not in t and "FY2025" in t, t[-600:])
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/money/summary"); await pg.wait_for_timeout(700)
    t=await pg.inner_text("#view"); check("peers: Funds shows the all-funds callout", "All funds · total revenue" in t and "Student Transportation" not in t)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/money/capital"); await pg.wait_for_timeout(700)
    t=await pg.inner_text("#view"); check("peers: Capital plan shows SAVE and PPEL callouts", "SAVE · ending balance" in t and "Low vs. peers" in t)
    await pg.click("[data-peers] a[data-action=peerDetail]"); await pg.wait_for_timeout(300)
    md=await pg.inner_text("[data-modal]")
    check("peers: Details shows the district against its peers", "SAVE · ending balance" in md and "Typical peer (median)" in md and "Middle half of peers" in md, md[:300])
    await pg.click("[data-modal] button[data-action=closeModal]"); await pg.wait_for_timeout(200)
    # what-if: leave a project out (nothing saved), then save it as a new scenario
    await pg.evaluate("document.querySelectorAll('details.fold').forEach(d=>d.open=true)")
    g0=await pg.inner_text("#cap-lead"); n0=len(calls)
    await pg.locator("#cap-out input[data-cap-out]").first.check(); await pg.wait_for_timeout(300)
    co=await pg.inner_text("#cap-out"); g1=await pg.inner_text("#cap-lead")
    check("what-if: leaving a project out recomputes the plan, saves nothing", "the gap goes from" in co and g1!=g0 and await pg.locator("#cap-years li.leftout").count()>0 and not [c for c in calls[n0:] if c[0] in ("POST","PATCH","DELETE")], co[:300])
    PROMPT["text"]="Without the first project"; n0=len(calls)
    await pg.click("#cap-out button[data-action=capOutSave]"); await pg.wait_for_timeout(700); PROMPT["text"]=None
    cp=[c for c in calls[n0:] if "/rpc/copy_scenario" in c[1]]; dl=[c for c in calls[n0:] if c[0]=="DELETE"]
    check("what-if: saved as a new scenario without it; the original untouched", cp and json.loads(cp[0][2])["p_name"]=="Without the first project" and len(dl)==3 and all("copied-scenario-id" in c[1] for c in dl), str(dl)[:300])
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/money/summary"); await pg.wait_for_timeout(700)
    t=await pg.inner_text("#view")
    key=await pg.evaluate("[...document.querySelectorAll('.cbar-row')].map(r=>[[...r.querySelectorAll('.cbar span')].map(x=>x.style.background),[...r.querySelectorAll('.cbar-key i')].map(x=>x.style.background),r.querySelector('.cbar-key').innerText])")
    check("funds: funding vs. committed, per fund, each bar with a key in the same colours and amounts", "Funding vs. committed" in t and "available" in t and key and all(k[0]==k[1] and "$" in k[2] for k in key), str(key)[:300])
    await pg.locator(".cbars").screenshot(path=SHOTS+"/funds-committed.png")
    sb=await pg.locator("svg[aria-label*='balance at the end of each year']").bounding_box(); cb=await pg.locator("svg[aria-label*='balance at the end of each year']").locator("xpath=ancestor::div[contains(@class,'card')][1]").bounding_box(); w=[sb["width"],cb["width"]]
    check("funds: the balance chart fills the card's width", w[0] > w[1]*0.9, str(w))
    bench=[c for c in calls if "/rpc/ia_benchmark" in c[1]]
    check("peers: asks for the latest year, districts its size, this district's rules", bench and json.loads(bench[-1][2])=={"p_de":"7777","p_fy":2025,"p_status":"Actual","p_peer":"size","p_district":"d1"}, str(bench[-1:]))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pg.wait_for_timeout(700)
    t=await pg.inner_text("#view")
    check("overview: peer numbers and open register questions need attention", "3 numbers stand out against similar Iowa districts" in t and "check-register questions are waiting" in t, t[:900])
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/district"); await pg.wait_for_timeout(600)
    check("district settings: the Iowa district number and its name", await pg.input_value("input[name=state_district_id]")=="7777" and "Ironwood Valley, AEA 267" in await pg.inner_text("[data-ia-linked]"))
    await pg.fill("input[name=state_district_id]",""); await pg.fill("input[data-ia-q]","iron"); await pg.click("button[data-action=iaSearch]"); await pg.wait_for_timeout(300)
    check("district settings: find the number by name", "Use 7777" in await pg.inner_text("[data-ia-results]"))
    await pg.click("button[data-action=iaPick]"); n0=len(calls); await pg.click("form[data-form=saveDistrict] button[type=submit]"); await pg.wait_for_timeout(600)
    pt=[c for c in calls[n0:] if c[0]=="PATCH" and "/rest/v1/district?" in c[1]]
    check("district settings: picking a match saves the number", pt and json.loads(pt[0][2]).get("state_district_id")=="7777", str(pt[:1]))
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/setup"); await pg.wait_for_timeout(700)
    folded=await pg.evaluate("(()=>{const d=document.querySelector('[data-setup-state]').closest('details');return d?!d.open:false})()")
    check("starting numbers: once set up, the state's figures fold into a comparison", folded and "Compare with the state’s annual report" in await pg.inner_text("#view"))
    await pg.evaluate("document.querySelectorAll('details.fold').forEach(d=>d.open=true)")
    stt=await pg.inner_text("[data-setup-state]")
    check("starting numbers: the state's figures offered, with the PPEL rates and a voted PPEL", "From the state’s data" in stt and "voter-approved $1.340" in stt and "V-PPEL revenue, FY2025, voted share (estimated)" in stt, stt)
    n1=len(calls); await pg.click("button[data-action=setupPrefill]"); await pg.wait_for_timeout(200)
    vals=[await pg.input_value(f"form[data-form=saveSetup] [name={n}]") for n in ["as_of","bal_save","bal_ppel","enrollment","enrollment_year","save_receipts","save_receipts_fy","ppel_receipts","ppel_rate","vppel_status","vppel_annual"]]
    check("starting numbers: Fill in sets the form, nothing saved (PPEL balance split by the two rates)", vals==["2025-06-30","1,900,000","126,467","1188","2025-26","1,450,000","2025","102,754","0.33","active","417,246"] and not [c for c in calls[n1:] if c[0] in("POST","PATCH")], str(vals))
    more=[await pg.input_value(f"form[data-form=saveSetup] [name={n}]") for n in ["bal_vppel","taxable_valuation","actual_valuation","ppel_growth","save_trend","save_ongoing","ppel_ongoing","grants_avg","vppel_first_fy","vppel_last_fy"]]
    check("starting numbers: Fill in also sets valuations, growth, SAVE trend, ongoing spending, grants (5 years) and the V-PPEL years", more==["513,533","727,000,000","1,250,000,000","3.1","2.77","180,000","95,000","46,000","2022","2031"], str(more))
    dfunds=await pg.evaluate("[...document.querySelectorAll('form[data-form=saveSetup] [data-debt-body] [data-debt-row]')].map(tr=>tr.querySelector('[name=debt_fund]').value+':'+tr.querySelector('[name=debt_annual]').value)")
    check("starting numbers: debt payments from the annual report fill funds that have no row yet", dfunds.count("debt_levy:295,000")==1 and sum(1 for x in dfunds if x.startswith("save:"))>=1, str(dfunds))
    stt2=await pg.inner_text("[data-setup-state]")
    await pg.locator("[data-setup-state]").screenshot(path=SHOTS+"/setup-state.png")
    check("starting numbers: the comparison lists the new state figures", "Taxable valuation, FY2027" in stt2 and "Ongoing SAVE spending a year" in stt2 and "Gifts and grants to the capital funds" in stt2 and "Voter-approved PPEL levied since" in stt2, stt2[-900:])
    await pg.select_option("select[data-grant-years]","3"); await pg.click("button[data-action=setupPrefill]"); await pg.wait_for_timeout(200)
    check("starting numbers: a 3-year grants average when chosen", await pg.input_value("form[data-form=saveSetup] [name=grants_avg]")=="60,000")
    hv=[await pg.input_value(f"form[data-form=saveSetup] [name={n}]") for n in ["tax_home_value","construction_inflation"]]
    check("starting numbers: the example home is the district's median home value (Census) and construction inflation comes from school construction prices (BLS)", hv==["187,000","4.94"], str(hv))
    await pg.click("[data-setup-state] button.defn[data-term=sources]"); await pg.wait_for_timeout(300)
    tt=await pg.inner_text("#toasts")
    check("the “?” by the state's figures names the public sources", "Where these figures come from" in tt and "U.S. Census Bureau" in tt and "Iowa Department of Management" in tt, tt[-500:])
    await pg.wait_for_timeout(9000)
    check("a definition stays until it's closed (no timing out mid-read)", "Where these figures come from" in await pg.inner_text("#toasts"))
    await pg.click("#toasts .toast.term .toast-x"); await pg.wait_for_timeout(200)
    check("…and closes with its ×", await pg.locator("#toasts .toast.term").count()==0)
    await pg.click("button[data-action=openHelp]"); await pg.wait_for_timeout(300)
    hm=await pg.inner_text("[data-modal]")
    check("Help on Starting numbers lists where the starting figures come from, with the Census Bureau's required notice", "Where the starting figures come from" in hm and "U.S. Bureau of Labor Statistics" in hm and "Certified Annual Report" in hm and "This product uses the Census Bureau Data API but is not endorsed or certified by the Census Bureau." in hm, hm[:400])
    check("the state's figures card carries the Census Bureau notice when it shows a home value", "not endorsed or certified by the Census Bureau" in await pg.inner_text("[data-setup-state]"))
    await pg.click("[data-modal] button[data-action=closeModal]"); await pg.wait_for_timeout(200)
    await pg.select_option("form[data-form=saveSetup] select[name=vppel_status]","none"); await pg.wait_for_timeout(200)
    check("starting numbers: warns when the state lists a voted PPEL but V-PPEL is None", "voter-approved PPEL of $1.340" in await pg.inner_text("[data-setup-checks]"))
    await pg.click("button[data-action=gfFromState]"); await pg.wait_for_timeout(1200)
    check("starting numbers: the General Fund figures go to the General Fund setup", "/money/general" in pg.url and await pg.input_value("form[data-form=saveGf] [name=fund_balance]")=="2,250,000", pg.url)
    await pg.click("[data-modal] button[data-action=closeModal]"); await pg.wait_for_timeout(200)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/money/general"); await pg.wait_for_timeout(600)
    await pg.click("button[data-action=gfEdit]"); await pg.wait_for_timeout(700)
    gst=await pg.inner_text("[data-gf-state]")
    check("General Fund setup: the state's figures offered (budget-year formula, unspent, salaries)", "From the state’s data" in gst and "budget enrollment 1,190.2 for FY2027" in gst and "district cost per pupil $8,148" in gst and "should add up to about" in gst, gst)
    await pg.click("button[data-action=gfPrefill]"); await pg.wait_for_timeout(200)
    gv=[await pg.input_value(f"form[data-form=saveGf] [name={n}]") for n in ["fund_balance","aea_flowthrough","enrollment","dcpp","other_formula","misc_income","misc_growth","unspent","nonstaff"]]
    check("General Fund setup: Fill these in", gv==["2,250,000","410,000","1190.2","8,148","3,802,250","1,900,000","2.1","2,100,000","3,300,000"] , str(gv))
    bh=await pg.evaluate("[...document.querySelectorAll('form[data-form=saveGf] [data-gf-staff]')].map(tr=>[tr.querySelector('[name=s_benefits]').value,tr.querySelector('[name=s_health]').value])")
    check("General Fund setup: the state's benefits % (which includes health) goes only where health per FTE is blank, so health isn't counted twice", bh and all((b!="30") if h.strip() not in ("","0") else (b=="30") for b,h in bh) and ("already includes health insurance" in await pg.inner_text("#toasts") if any(h.strip() not in ("","0") for b,h in bh) else True), str(bh))
    roster=b"Position,Group,FTE,Annual salary,District health insurance contribution (annual)\r\nA,Teachers,1,52000,13200\r\nB,teachers,0.5,26000,0\r\nC,Paraeducators,1,27000,9600\r\nD,,1,1,1\r\n"
    await pg.set_input_files("input[data-gf-roster]",files=[{"name":"staff.csv","mimeType":"text/csv","buffer":roster}]); await pg.wait_for_timeout(500)
    srows=await pg.evaluate("[...document.querySelectorAll('form[data-form=saveGf] [data-gf-staff]')].map(tr=>['s_name','s_fte','s_salary','s_benefits','s_health'].map(n=>tr.querySelector('[name='+n+']').value).join('|'))")
    check("General Fund setup: a staff list fills the groups, averaged per FTE", srows==["Teachers|1.5|52,000|17.09|8,800","Paraeducators|1|27,000|17.09|9,600"] and "Read 3 people in 2 groups" in await pg.inner_text("[data-gf-roster-note]") and "no group" in await pg.inner_text("[data-gf-roster-note]"), str(srows))
    await pg.locator("[data-modal] .modal").screenshot(path=SHOTS+"/gf-filled.png")
    async with pg.expect_download() as dl: await pg.click("a[data-action=gfRosterTemplate]")
    check("General Fund setup: the staff list template is an Excel file", (await dl.value).suggested_filename=="highground-staff-list-template.xlsx")
    await pg.click("[data-modal] button[data-action=closeModal]"); await pg.wait_for_timeout(200)
    # check registers: upload two months, review, apply, then the questions
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/settings/uploads"); await pg.wait_for_timeout(600)
    await pg.select_option("select[data-upload-kind]","check_register")
    reg=("Ironwood Valley CSD - Check Register\n\nCheck Date,Check #,Vendor #,Vendor Name,Invoice #,Description,Account Number,Amount\n"
         "09/03/2026,10421,V0012,Midwest Bus Parts,INV-5531,Brake pads,10-2700-000-000-1100-000-618,\"1,240.50\"\n"
         "09/15/2026,10422,V0044,Alliant Energy,88231,Electric,10-2600-000-000-1100-000-622,8411.00\n"
         "10/01/2026,10424,V0012,Midwest Bus Parts,INV-5602,Filters,10-2700-000-000-1100-000-618,99\n,,,Total,,,,9750.50\n")
    await pg.set_input_files("input[data-upload-file]",files=[{"name":"register.csv","mimeType":"text/csv","buffer":reg.encode()}]); await pg.wait_for_timeout(600)
    t=await pg.inner_text("#upload-review")
    check("register upload: reviewed by month before anything is saved", "3 payments (1 total or blank row left out), in 2 months" in t and "September 2026" in t and "October 2026" in t and "first two digits = fund" in t, t[:400])
    n0=len(calls); await pg.click("#upload-review button[data-action=applyUpload]"); await pg.wait_for_timeout(1200)
    new=calls[n0:]
    bt=[json.loads(c[2]) for c in new if c[0]=="POST" and c[1].startswith("/rest/v1/import_batch")]
    rl=[json.loads(c[2]) for c in new if c[0]=="POST" and c[1].startswith("/rest/v1/register_line")]
    check("register upload: one upload per month, its payments, applied and checked", [b["period_end"] for b in bt]==["2026-09-30","2026-10-31"] and all(b["kind"]=="check_register" for b in bt)
          and sum(len(x) for x in rl)==3 and rl[0][0]["batch_id"]==bt[0]["id"] and rl[0][0]["fund"]=="10" and len([c for c in new if "/rpc/apply_import" in c[1]])==2 and len([c for c in new if "/rpc/register_check" in c[1]])==2, str(bt)[:300])
    check("register upload: opens the check register afterwards", "/track/registers" in pg.url, pg.url)
    await pg.goto("http://localhost:8765/#/d/ironwood-valley/track/registers"); await pg.wait_for_timeout(700)
    folds=await pg.evaluate("[...document.querySelectorAll('details.fold')].map(d=>d.open)")
    await pg.evaluate("document.querySelectorAll('details.fold').forEach(d=>d.open=true)"); t=await pg.inner_text("#view")
    await pg.evaluate("document.querySelectorAll('details.fold').forEach(d=>d.open=false)")
    check("check register: questions first, most serious first; information and answers folded", t.index("Vendor name changed")<t.index("New vendor") and "For information (1)" in t and "Answered (1)" in t and "Second check voided" in t and folds and not any(folds), t[:600])
    await pg.screenshot(path=SHOTS+"/registers.png",full_page=True)
    n0=len(calls); await pg.click("[data-reg-flag='1'] button[data-status=explained]"); await pg.wait_for_timeout(300)
    check("check register: an answer is required to explain", "Write the answer first" in await pg.inner_text("#toasts") and not [c for c in calls[n0:] if c[0]=="PATCH"])
    await pg.fill("[data-reg-flag='1'] textarea[name=response]","W-9 on file; the vendor merged with its parent."); await pg.click("[data-reg-flag='1'] button[data-status=explained]"); await pg.wait_for_timeout(500)
    pt=[json.loads(c[2]) for c in calls[n0:] if c[0]=="PATCH" and "register_flag" in c[1]]
    check("check register: the answer is saved with who and when", pt and pt[0]["status"]=="explained" and pt[0]["response"].startswith("W-9") and pt[0]["resolved_by"]=="u-admin", str(pt))
    n0=len(calls); await pg.click("[data-reg-flag='2'] button[data-action=regVendorOk]"); await pg.wait_for_timeout(400)
    vn=[c for c in calls[n0:] if c[0]=="POST" and "vendor_note" in c[1]]
    check("check register: a new vendor can be marked as expected", vn and json.loads(vn[0][2])[0]["vendor_key"]=="hawkeye roofing" and "on_conflict=district_id,vendor_key" in urllib.parse.unquote(vn[0][1]), str(vn))
    await pg.click("summary:has-text(\"Settings for these checks\")"); await pg.fill("input[data-reg-threshold]","25000"); n0=len(calls); await pg.click("button[data-action=regThreshold]"); await pg.wait_for_timeout(500)
    rr=[json.loads(c[2]) for c in calls[n0:] if c[0]=="POST" and "register_rule" in c[1]]
    check("check register: the bid threshold turns on two checks and re-checks every month", len(rr)==2 and all(x["params"]["threshold"]==25000 and x["district_id"]=="d1" for x in rr) and any("/rpc/register_check_all" in c[1] for c in calls[n0:]), str(rr))
    # public page, opened by someone with no account (fresh browser, no session)
    # step A: menus by role, landing on the board report, the guide, definitions
    TABLES["report_snapshot"]=[{"id":"rep-sep","district_id":"d1","kind":"board_monthly","title":"Board report, September 2026","period_end":"2026-09-30","created_at":"2026-10-01T15:00:00Z",
      "payload":{"version":1,"periodEnd":"2026-09-30","fiscalYear":2027,"district":"Ironwood Valley Community School District","scenario":{"id":"s1","name":"District baseline"},"ledgerThrough":None,
        "plan":{"need":14790000,"levyFunded":9190000,"financed":0,"gap":5350000,"overflow":0,"start":2027,"years":10},"balances":{},"budget":[],"progress":[],"initiatives":{},"done":[],"pending":[],"pendingCount":0,"tax":None,"gf":None}}]
    async def session(email):
      cx=await b.new_context(viewport={"width":1360,"height":900}); await cx.route("**/config.js",cfg); await cx.route(SB+"/**",handler); await cx.route("**/fonts.g*/**",lambda r:r.abort())
      pp=await cx.new_page(); pp.on("pageerror",lambda e:errs.append(str(e)))
      await pp.goto("http://localhost:8765/"); await pp.evaluate("localStorage.clear()"); await pp.goto("http://localhost:8765/"); await pp.wait_for_timeout(400)
      await pp.fill("input[name=email]",email); await pp.fill("input[name=password]",PW); await pp.click("button[type=submit]"); await pp.wait_for_timeout(900)
      return cx,pp
    menu=lambda pp: pp.locator("nav.rail .nav:not(.nav-foot) .it .tx").all_inner_texts()
    cx,pp=await session("board@example.test")
    check("board member: lands on the latest board report", "/share/board" in pp.url and "Board report · Sep 30, 2026" in await pp.inner_text("#view"), pp.url)
    m=[x.strip() for x in await menu(pp)]
    check("board member: a short menu", m==["Home","Plan","Money","Track","Share"], str(m))
    bsub=[x.strip() for x in await pp.locator("nav.rail .grp .kids a").all_inner_texts()]
    check("board member: each area lists only the board's screens", bsub==["Priorities","Capital plan","General fund","All funds","Measures","Board reports","District plan"], str(bsub))
    check("board member: no scenario picker at the top, just which board version they're seeing", await pp.locator("#scenTop select").count()==0 and "Board version" in await pp.inner_text("#scenTop"))
    check("board member: no Settings in the menu; the account menu has Your account and Sign out, but no preview", await pp.locator(".nav-foot a[data-area=settings]").count()==0 and await pp.locator(".umenu a", has_text="Your account").count()==1 and await pp.locator("[data-action=togglePreview]").count()==0)
    check("board member: Reports has no Community page tab; the public link sits on Board reports", await pp.locator("nav.tabs a", has_text="Community page").count()==0)
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pp.wait_for_timeout(700); bt=await pp.inner_text("#view")
    check("board member: Overview tiles are the board's (no ledger or people counts; no links to staff screens)", "People with access" not in bt and "Ledger" not in bt and "Awaiting a decision" in bt and "add measures" not in bt and "(decisions)" not in bt, bt[:500])
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/money/summary"); await pp.wait_for_timeout(500)
    check("board member: Money shows Capital plan, General fund, All funds", [x.strip() for x in await pp.locator("nav.tabs a").all_inner_texts()]==["Capital plan","General fund","All funds"])
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/money/capital"); await pp.wait_for_timeout(800)
    folds=await pp.evaluate("[...document.querySelectorAll('details.fold')].map(d=>d.open)")
    await pp.screenshot(path=SHOTS+"/board-capital.png", full_page=True)
    check("board member: no editing notices", "Copy it to try changes" not in await pp.inner_text("#view"))
    check("board member: the capital plan leads with the answer, details folded away", "of projects over 10 years" in await pp.inner_text("#cap-lead") and folds and not any(folds), str(folds))
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/settings/account"); await pp.wait_for_timeout(500)
    check("board member: no “show every screen” option", await pp.locator("input[data-show-all]").count()==0)
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/plan/initiatives"); await pp.wait_for_timeout(500)
    check("board member: typing a staff screen's address doesn't open it", "Not part of your view" in await pp.inner_text("#content") and await pp.locator("#view").count()==0)
    # every screen a board member can open: no links to screens they can't
    bad=[]
    for path in ["home/today","plan/priorities","track/measures","plan/community","money/summary","money/general","money/capital","share/board","share/plans","settings/account","help/guide"]:
      await pp.goto("http://localhost:8765/#/d/ironwood-valley/share/board"); await pp.wait_for_timeout(150)
      await pp.goto("http://localhost:8765/#/d/ironwood-valley/"+path); await pp.wait_for_timeout(700)
      hrefs=await pp.evaluate("[...document.querySelectorAll('#view a[href^=\"#/d/\"]')].map(a=>a.getAttribute('href'))")
      ok=("home/today","plan/priorities","track/measures","money/summary","money/general","money/capital","share/board","share/plans","settings/account","help/guide")
      bad+=[path+" → "+h for h in hrefs if "/".join(h.split("?")[0].split("/")[3:5]) not in ok]
    check("board member: no links to staff screens anywhere they can go", not bad, str(bad))
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/plan/priorities"); await pp.wait_for_timeout(700)
    check("board member: Priorities doesn't show the planners' to-do about unlinked initiatives", "linked to a priority yet" not in await pp.inner_text("#view"))
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/help/guide"); await pp.wait_for_timeout(500)
    g=await pp.inner_text("#view")
    check("guide: where to start for a board member, how-to, the glossary, and the data sources with the Census notice", "open the latest board report" in g and "How do I" in g and "Solvency ratio" in g and "Encumbered" in g and "support@willowholler.com" in g and "Where the starting figures come from" in g and "not endorsed or certified by the Census Bureau" in g, g[:300])
    check("guide: no “What’s built” tab for districts", "What’s built" not in await pp.inner_text("nav.tabs") if await pp.locator("nav.tabs").count() else True)
    await cx.close()
    cx,pp=await session("bm@example.test")
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/settings/account"); await pp.wait_for_timeout(500)
    await pp.check("input[data-show-all]"); await pp.wait_for_timeout(600)
    m2=[x.strip() for x in await menu(pp)]
    check("show every screen (staff roles): the full menu comes back", "Plan" in m2 and "Track" in m2, str(m2))
    await pp.uncheck("input[data-show-all]"); await pp.wait_for_timeout(500)
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/plan/initiatives"); await pp.wait_for_timeout(500)
    check("staff roles: a direct link to a screen outside the menu still opens", "Initiatives" in await pp.inner_text("h1"))
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pp.wait_for_timeout(700)
    m=[x.strip() for x in await menu(pp)]
    check("business manager: lands on Home; menu for the monthly close", "/home/today" in pp.url and m==["Home","Money","Track","Share"], str(m)+pp.url)
    tsub=[x.strip() for x in await pp.locator("nav.rail .grp[data-area=track] .kids a").all_inner_texts()]
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/settings/uploads"); await pp.wait_for_timeout(500)
    check("business manager: Check register in Track, and as an upload", "Check register" in tsub and "Check register (bills paid)" in await pp.inner_text("select[data-upload-kind]"), str(tsub))
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pp.wait_for_timeout(700); bmt=await pp.inner_text("#view")
    check("business manager: Overview leads with the monthly close", bmt.find("Ledger") < bmt.find("Gap to close") and "Check register" in bmt and "Next month-end" in bmt, bmt[:400])
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/settings/uploads"); await pp.wait_for_timeout(500)
    check("business manager: Uploads first in Settings, after the overview", [x.strip() for x in await pp.locator("nav.tabs a").all_inner_texts()][:2]==["Settings","Uploads"])
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pp.wait_for_timeout(600)
    await pp.click("button.defn[data-term=gap]"); await pp.wait_for_timeout(200)
    dw=await pp.evaluate("(()=>{const t=document.querySelector('#toasts .toast.term');return t?[t.getBoundingClientRect().width,t.scrollWidth<=t.clientWidth+1]:null})()")
    check("definitions: the explanation is a full-size card (not squeezed into the “?” circle)", dw and dw[0]>=280 and dw[1], str(dw))
    check("definitions: tapping “?” explains the term", "Capital costs in the plan that SAVE, PPEL, V-PPEL and grants can’t cover" in await pp.inner_text("#toasts"))
    await cx.close()
    cx,pp=await session("viewer@example.test")
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/track/registers"); await pp.wait_for_timeout(700)
    t=await pp.inner_text("#content")
    check("viewer: the check register is a staff screen (doesn't open)", "Not part of your view" in t and await pp.locator("[data-action=regAnswer], [data-reg-threshold]").count()==0, t[:300])
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/settings/district"); await pp.wait_for_timeout(500)
    check("viewer: cannot change the Iowa district number (district settings don't open)", await pp.locator("input[name=state_district_id]").count()==0 and await pp.locator("[data-ia-q]").count()==0 and "Not part of your view" in await pp.inner_text("#content"))
    await cx.close()
    cx,pp=await session("sup@example.test")
    m=[x.strip() for x in await menu(pp)]
    check("superintendent: the planning menu", m==["Home","Plan","Money","Track","Share"], str(m))
    await cx.close()
    TABLES["report_snapshot"]=[]
    anonctx=await b.new_context(viewport={"width":1360,"height":900}); await anonctx.route("**/config.js",cfg); await anonctx.route(SB+"/**",handler); await anonctx.route("**/fonts.g*/**",lambda r:r.abort())
    pub=await anonctx.new_page(); pub.on("pageerror",lambda e:errs.append(str(e)))
    n0=len(calls); await pub.goto("http://localhost:8765/#/p/ironwood-valley"); await pub.wait_for_timeout(700)
    t=await pub.inner_text("body")
    authed=[c for c in calls[n0:] if "/rest/v1/" in c[1] and "rpc/public_publication" not in c[1]]
    check("public link: plan shows with no account ($5.35M)", "Ironwood Valley" in t and "$5.35M" in t and "Sign in" in t and "Sign out" not in t, t[:200])
    check("public link: only the public function is called", not authed, str(authed[:2]))
    check("public link: branded, with the community note", "The board adopts the plan in October" in await pub.inner_text("body") and await pub.locator(".pubhead .pub-mono, .pubhead .pub-dlogo-lg").count()==1)
    check("public link: nothing to edit or publish", await pub.locator("[data-notbuilt], [data-action=publishBoard], form").count()==0)
    check("public link: no what-if tag before a lever moves", await pub.locator(".whatif-tag").count()==0)
    await pub.click("input[data-lever=sf]"); await pub.wait_for_timeout(200)
    check("public link: moved levers tag the figures as a what-if", "not the published plan" in await pub.inner_text("#cap-results"))
    check("public link: shows what it means for taxpayers", "What it means for taxpayers" in await pub.inner_text("body"))
    check("public link: visitors can move levers", fmtK(gold("orig",{"sf":False})["gap"]) in await pub.inner_text("#cap-results"))
    await pub.screenshot(path=SHOTS+"/public.png",full_page=True)
    await pub.goto("http://localhost:8765/#/p/nowhere"); await pub.wait_for_timeout(400)
    check("unpublished link", "Nothing published here yet" in await pub.inner_text("body"))
    # phone
    ph=await b.new_context(viewport={"width":375,"height":812},is_mobile=True,has_touch=True,device_scale_factor=3); await ph.route("**/config.js",cfg); await ph.route(SB+"/**",handler); await ph.route("**/fonts.g*/**",lambda r:r.abort())
    pp=await ph.new_page(); pp.on("pageerror",lambda e:errs.append(str(e)))
    await pp.goto("http://localhost:8765/#/signin"); await pp.fill("input[name=email]","admin@example.test"); await pp.fill("input[name=password]",PW); await pp.click("button[type=submit]"); await pp.wait_for_timeout(700)
    sw=await pp.evaluate("document.documentElement.scrollWidth"); check("phone: no sideways scroll", sw<=375, str(sw))
    WIDE_JS="""(()=>{const W=innerWidth,out=[];document.querySelectorAll('body *').forEach(e=>{const r=e.getBoundingClientRect();if(r.width>0&&r.right>W+1){let p=e.parentElement,clipped=false;while(p&&p!==document.body){const cs=getComputedStyle(p);if(/auto|scroll/.test(cs.overflowX)&&p.getBoundingClientRect().right<=W+1){clipped=true;break}p=p.parentElement}if(!clipped)out.push((e.className&&e.className.baseVal===undefined?e.tagName+'.'+String(e.className).split(' ')[0]:e.tagName)+':'+Math.round(r.right))}});return [document.documentElement.scrollWidth,[...new Set(out)].slice(0,8)]})()"""
    wide=[]
    for path in ["home/today","plan/priorities","track/measures","plan/initiatives","plan/ranking","plan/scenarios","money/summary","money/general","money/capital","track/initiatives","track/actuals","settings/uploads","track/registers","share/board","share/plans","share/community","settings/district","settings/setup","settings/people","settings/activity","settings/assumptions","settings/exports","settings/account","help/guide"]:
      await pp.goto("http://localhost:8765/#/d/ironwood-valley/"+path); await pp.wait_for_timeout(700)
      await pp.evaluate("document.querySelectorAll('details.fold').forEach(d=>d.open=true)"); await pp.wait_for_timeout(150)
      sw2,els=await pp.evaluate(WIDE_JS)
      if sw2>375 or els: wide.append(f"{path}:{sw2}:{els}")
    check("phone: nothing on any screen runs past the edge or is cut off (wide tables scroll in their own box)", not wide, "\n".join(wide))
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pp.wait_for_timeout(700)
    await pp.click("button.defn >> nth=0"); await pp.wait_for_timeout(300)
    tb=await pp.evaluate("(()=>{const t=document.querySelector('#toasts .toast'),r=t.getBoundingClientRect(),cs=getComputedStyle(t);return [r.left,r.right,innerWidth,r.width,t.scrollWidth<=t.clientWidth+1,cs.backgroundColor,cs.borderRadius]})()")
    check("phone: a definition sits inside the screen, nearly full width, on a solid card with its text inside it", tb[0]>=0 and tb[1]<=tb[2] and tb[3]>=tb[2]*0.85 and tb[4] and tb[5] not in ("rgba(0, 0, 0, 0)","transparent") and tb[6]!="50%", str(tb))
    await pp.click("#toasts .toast-x")
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/settings/setup"); await pp.wait_for_timeout(800)
    await pp.click("button[data-action=openHelp]"); await pp.wait_for_timeout(400)
    mb=await pp.evaluate("(()=>{const m=document.querySelector('[data-modal] .modal'),b=m.querySelector('.modal-body');b.scrollTop=b.scrollHeight;const r=m.getBoundingClientRect();const h=m.querySelector('#modal-title').getBoundingClientRect();const c=m.querySelector('[data-action=closeModal]').getBoundingClientRect();return [r.top,r.bottom,innerHeight,h.top>=r.top-1&&h.bottom<=r.bottom&&c.bottom<=r.bottom&&h.bottom<=b.getBoundingClientRect().top+1,b.scrollTop>0]})()")
    check("phone: Help slides up inside the screen, scrolls inside itself, and keeps its title and Close in view", mb[0]>=0 and mb[1]<=mb[2]+1 and mb[3] and mb[4], str(mb))
    await pp.screenshot(path=SHOTS+"/phone-help.png")
    await pp.click("[data-modal] button[data-action=closeModal]"); await pp.wait_for_timeout(200)
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/settings/assumptions"); await pp.wait_for_timeout(800)
    tabs=await pp.evaluate("(()=>{const n=document.querySelector('nav.tabs');return [...n.querySelectorAll('a')].every(a=>a.getBoundingClientRect().right<=innerWidth)})()")
    check("phone: every sub-menu tab is visible (they wrap instead of scrolling)", tabs)
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/settings/setup"); await pp.wait_for_timeout(900)
    await pp.evaluate("document.querySelectorAll('details.fold').forEach(d=>d.open=true)"); await pp.wait_for_timeout(300)
    sc=await pp.evaluate("[...document.querySelectorAll('[data-setup-state] .scroll')].map(b=>b.scrollWidth-b.clientWidth)")
    check("phone: the state comparison fits without sideways scrolling (it becomes a list)", sc and max(sc)<=2 and await pp.locator("[data-setup-state] table.stacked").count()>=1, str(sc))
    await pp.screenshot(path=SHOTS+"/phone-setup-state.png", full_page=True)
    small=[]
    for path in ["money/summary","money/general"]:
      await pp.goto("http://localhost:8765/#/d/ironwood-valley/"+path); await pp.wait_for_timeout(900)
      hs=await pp.evaluate("[...document.querySelectorAll('#view svg text')].filter(t=>t.getBoundingClientRect().width>0).map(t=>t.getBoundingClientRect().height)")
      if hs and min(hs)<9: small.append(f"{path}: {min(hs):.1f}px")
    check("phone: chart labels stay readable (at least 9px tall on screen)", not small, str(small))
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/plan/initiatives"); await pp.wait_for_timeout(900)
    ov=await pp.evaluate("""(()=>{const els=[...document.querySelectorAll('#view .row > .chip, #view .row > .btn')].map(e=>e.getBoundingClientRect());for(let i=0;i<els.length;i++)for(let j=i+1;j<els.length;j++){const a=els[i],b=els[j];if(a.left<b.right-1&&b.left<a.right-1&&a.top<b.bottom-1&&b.top<a.bottom-1)return true}return false})()""")
    check("phone: the initiative pickers and buttons don't overlap", not ov)
    await pp.screenshot(path=SHOTS+"/phone-initiatives.png")
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pp.wait_for_timeout(900)
    await pp.screenshot(path=SHOTS+"/phone-home.png", full_page=True)
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/money/capital"); await pp.wait_for_timeout(1000)
    await pp.evaluate("document.querySelectorAll('details.fold').forEach(d=>d.open=true)"); await pp.wait_for_timeout(300)
    rows=await pp.evaluate("[...document.querySelectorAll('.outrow')].slice(0,6).map(l=>{const c=l.querySelector('input').getBoundingClientRect(),n=l.querySelector('span').getBoundingClientRect();return Math.abs(c.top-n.top)<12})")
    check("phone: Leave projects out keeps each checkbox beside its project", rows and all(rows), str(rows))
    await pp.screenshot(path=SHOTS+"/phone-capital.png", full_page=True)
    await pp.screenshot(path=SHOTS+"/phone-assumptions.png", full_page=True)
    await pp.goto("http://localhost:8765/#/d/ironwood-valley/home/today"); await pp.wait_for_timeout(600)
    inview=lambda: pp.evaluate("(()=>{const r=document.querySelector('.rail').getBoundingClientRect();return r.right>40&&r.left>=-1})()")
    navvis=await inview()
    await pp.click(".menu-btn"); await pp.wait_for_timeout(350)
    opened=await inview() and await pp.locator(".rail .user").is_visible()
    await pp.screenshot(path=SHOTS+"/phone-menu.png")
    await pp.click(".rail [data-area-head=money]"); await pp.wait_for_timeout(250)
    inline=await pp.locator(".rail .grp[data-area=money] .kids a", has_text="Capital plan").is_visible()
    await pp.click(".rail .grp[data-area=money] .kids a:has-text('Capital plan')"); await pp.wait_for_timeout(700)
    check("phone: a Menu button slides in the menu (with your account and Sign out); tapping an area opens its screens; picking one goes there and closes it", not navvis and opened and inline and "/money/capital" in pp.url and not await inview(), pp.url)
    await pp.click(".menu-btn"); await pp.wait_for_timeout(350); await pp.click(".rail button[data-action=userMenu]"); await pp.wait_for_timeout(200)
    check("phone: one background colour behind the page (no menu column showing through)", await pp.evaluate("getComputedStyle(document.querySelector('.frame')).backgroundImage")=="none")
    check("phone: Sign out is in the account menu at the bottom of the menu", await pp.locator(".umenu .menu-signout").is_visible())
    await pp.click(".menu-btn"); await pp.wait_for_timeout(300)
    await pp.set_viewport_size({"width":1360,"height":900}); await pp.wait_for_timeout(300)
    check("desktop: the side menu shows every area and no Menu button", await pp.locator(".menu-btn").is_hidden() and await inview() and await pp.locator(".rail .user").is_visible())
    await pp.set_viewport_size({"width":375,"height":812})
    await pp.screenshot(path=SHOTS+"/phone.png")
    # phone sign-in: fields at 16px (iPhone Safari zooms into smaller ones and the page drifts); card pinned to the top
    ph2=await b.new_context(viewport={"width":390,"height":844},is_mobile=True,has_touch=True); await ph2.route("**/config.js",cfg); await ph2.route(SB+"/**",handler); await ph2.route("**/fonts.g*/**",lambda r:r.abort())
    sp=await ph2.new_page(); sp.on("pageerror",lambda e:errs.append(str(e)))
    await sp.goto("http://localhost:8765/#/signin"); await sp.wait_for_timeout(500)
    fs=await sp.evaluate("[...document.querySelectorAll('.auth input')].map(i=>parseFloat(getComputedStyle(i).fontSize))")
    top0=(await sp.locator(".auth-card").bounding_box())["y"]
    await sp.tap("input[name=email]"); await sp.wait_for_timeout(200); await sp.tap("input[name=password]"); await sp.wait_for_timeout(200)
    top1=(await sp.locator(".auth-card").bounding_box())["y"]
    check("phone sign-in: fields are 16px so iPhone doesn't zoom; the card stays put when fields are tapped", fs and min(fs)>=16 and abs(top1-top0)<1 and await sp.evaluate("getComputedStyle(document.querySelector('.auth-main')).alignItems")=="flex-start", f"{fs} {top0}->{top1}")
    check("phone sign-in: the logo shows above the form", await sp.locator(".auth-mobile-brand img").is_visible())
    await sp.screenshot(path=SHOTS+"/phone-signin.png")
    await b.close()
  srv.terminate()
  for n,okk,d in R: print(("PASS " if okk else "FAIL ")+n+(("  ["+d+"]") if d and not okk else ""))
  print("page errors:",errs)
asyncio.run(main())
